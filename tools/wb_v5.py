"""v5 sheets: 18 Computer Vision, 19 Squad & Mission, 20 Audio & HUD.

Numbers come from build/scene_export.json (schema 4, key "v5": squad / mission / tracker / cv / armour constants and the
exporter's own roundFlightTime(d) values) and from build/test_results.json (observed values of T130-T181). Pale-yellow
cells (NOTE_FILL) are hand-maintained text from the design notes (src header comments, docs/v5_notes/v5-a5.md,
v5-h5.md, docs/PLAN_V5.md). Blue = imported, white = Excel formula.
"""
import os
import re

from openpyxl.formatting.rule import CellIsRule

from wb_common import (BOLD, FAIL_FILL, INPUT_FILL, NOTE_FILL, PASS_FILL, PENDING_FILL, ROOT, S_AUDIO, S_COMBAT, S_CV,
                       S_SQUAD, S_TEST, TOTAL_FILL, col, define_name, finish, header, image, path, put, section, title,
                       widths)

# ---- helpers -----------------------------------------------------------------------------------------------


def _consts(ctx, key):
    return {c["key"]: c for c in ctx["scene"].get("v5", {}).get(key, [])}


def _const_table(ws, r, rows, ref_map, head=("Constant", "Value", "Unit", "Source file")):
    """Writes constants (blue) and records their cell addresses in ref_map. Returns next row."""
    header(ws, r, list(head))
    r += 1
    for c in rows:
        put(ws, r, 1, c["key"], font=BOLD)
        put(ws, r, 2, c["value"], "0.###", INPUT_FILL)
        put(ws, r, 3, c["unit"])
        put(ws, r, 4, c["source"])
        ref_map[c["key"]] = f"$B${r}"
        r += 1
    return r


def _tests_table(ws, r, ctx, ids, label):
    """Imported test rows (id, description, expected, observed) + a live status formula over sheet 10."""
    by_id = {t["id"]: t for t in ctx["tests"]}
    t_first, t_last = ctx["test_rows"]
    section(ws, r, 1, label)
    header(ws, r + 1, ["Test", "What is checked", "Expected", "Observed (from build/test_results.json)", "Status (sheet 10)"])
    r += 2
    first = r
    for tid in ids:
        t = by_id.get(tid)
        if not t:
            continue
        put(ws, r, 1, tid, font=BOLD)
        put(ws, r, 2, t["description"])
        put(ws, r, 3, t["expected"])
        put(ws, r, 4, t["observed"], fill=INPUT_FILL)
        put(ws, r, 5, f"=INDEX('{S_TEST}'!$F${t_first}:$F${t_last},MATCH(A{r},'{S_TEST}'!$A${t_first}:$A${t_last},0))", font=BOLD)
        ws.row_dimensions[r].height = 45
        r += 1
    if r > first:
        ws.conditional_formatting.add(f"E{first}:E{r - 1}", CellIsRule(operator="equal", formula=['"PASS"'], fill=PASS_FILL))
        ws.conditional_formatting.add(f"E{first}:E{r - 1}", CellIsRule(operator="equal", formula=['"FAIL"'], fill=FAIL_FILL))
        put(ws, r, 1, "Passing", fill=TOTAL_FILL, font=BOLD)
        put(ws, r, 2, f'=COUNTIF(E{first}:E{r - 1},"PASS")&" / "&COUNTA(A{first}:A{r - 1})', fill=TOTAL_FILL, font=BOLD)
        pass_cell = f"B{r}"
        r += 1
    else:
        pass_cell = None
    return r + 2, pass_cell


def _fig(ws, r, anchor_col, rel, caption, width=520):
    p = path(rel)
    put(ws, r, anchor_col, caption + ("" if os.path.exists(p) else f" - pending capture {rel}"), border=False)
    return image(ws, f"{col(anchor_col)}{r + 1}", p, width=width)


# ---- 18 Computer Vision ------------------------------------------------------------------------------------
STAGES = [
    ("1", "Adaptive threshold", "t = max(floor, Otsu threshold of the frame histogram)",
     "Otsu: the level that maximises the between-class variance of the 256-bin histogram; floor keeps cold frames from "
     "producing targets", "CvPipeline.cpp cv::otsuThreshold"),
    ("2", "Morphological despeckle", "drop mask pixels with fewer than 2 hot 8-neighbours",
     "removes isolated hot pixels (sparks, grass glints, sensor noise)", "cv::despeckle"),
    ("3", "Connected components", "two-pass labelling, 8-connectivity, union-find with path compression",
     "turns the hot mask into blobs", "cv::labelComponents"),
    ("4", "Blob features", "area, centroid, box, mean / peak heat, central moments -> orientation, elongation",
     "elongation = sqrt(lambda_max / lambda_min) of the second central moments", "cv::detectTargets"),
    ("5", "Target filter", "area in [minArea, maxArea], a pixel >= core, elongation <= maxElongation",
     "rejects fires (too big), tracers / wires (thin streaks) and cold warm-ups", "cv::detectTargets"),
    ("6", "Pixel ray -> ground point", "ray through the blob centroid (gimbal camera matrix), intersected with the terrain",
     "gives one ground measurement with sigma growing with range", "DroneController::imagePointToGround"),
    ("7", "Gating", "squared Mahalanobis distance of the innovation < chi-square (2 dof, 99 %)",
     "a measurement may only update a track it is statistically close to", "MultiTracker::mahalanobis2"),
    ("8", "Assignment", "Hungarian algorithm (Kuhn-Munkres) on the gated costs, one sensor at a time",
     "optimal one-to-one track <-> measurement matching (tested against brute force, T133)", "MultiTracker::hungarian"),
    ("9", "Two-view triangulation", "midpoint of the shortest segment between the two UAVs' pixel rays",
     "when both UAVs see the same rover in one frame: smaller error than one ground ray (T134)", "MultiTracker::triangulate"),
    ("10", "Kalman filter", "constant-velocity model, state (x, z, vx, vz), white-noise acceleration",
     "smooths positions, estimates velocity used for the strike's lead prediction (T135)", "MultiTracker::kalmanUpdate"),
    ("11", "Track life cycle", "tentative -> confirmed (confirmHits) -> coasting (coastAfter s) -> deleted (deleteAfter s)",
     "one-frame clutter never becomes a track (T136); close tracks with similar velocity are merged", "MultiTracker::update"),
]


def build_cv(wb, ctx):
    ws = wb[S_CV]
    v5 = ctx["scene"].get("v5")
    if not v5:
        put(ws, 4, 1, "PENDING - scene_export.json has no 'v5' section (re-run `make docs`)")
        return
    title(ws, "18  Computer Vision",
          "v5 vision: every UAV turns its 320 x 240 thermal image into blobs (CvPipeline), the blobs become ground "
          "measurements, and one shared multi-target tracker (Kalman + Hungarian + two-view triangulation) builds the "
          "squad's track picture. Parameters are exported from the code (blue); the stage descriptions are from the source "
          "header comments (pale yellow); the derived figures are Excel formulas; the test rows are the program's own "
          "observed values (T130-T137).", span=9)
    r = 4
    section(ws, r, 1, "A. Pipeline stages")
    header(ws, r + 1, ["#", "Stage", "Operation", "Purpose", "Code"])
    r += 2
    for sid, name, op, why, code in STAGES:
        put(ws, r, 1, sid, font=BOLD)
        put(ws, r, 2, name, font=BOLD)
        put(ws, r, 3, op, fill=NOTE_FILL)
        put(ws, r, 4, why, fill=NOTE_FILL)
        put(ws, r, 5, code)
        ws.row_dimensions[r].height = 48
        r += 1
    r += 2
    refs = {}
    section(ws, r, 1, "B. Parameters (CvConfig, camera constants, MultiTracker defaults) - exported")
    r = _const_table(ws, r + 1, v5["cv"], refs)
    r += 1
    r = _const_table(ws, r, v5["tracker"], refs)
    r += 2
    ctx["refs"]["cv_refs"] = refs
    # ---- derived ----
    section(ws, r, 1, "C. Derived figures (formulas over B)")
    header(ws, r + 1, ["Quantity", "Value", "Formula meaning"])
    r += 2
    d = [
        ("Pixels per frame", f"={refs['kCvWidth']}*{refs['kCvHeight']}", "0", "width x height"),
        ("Pixels per degree (vertical)", f"={refs['kCvHeight']}/{refs['kCvFovYDeg']}", "0.00", "rows / vertical FOV"),
        ("Horizontal FOV (deg)", f"=2*DEGREES(ATAN(TAN(RADIANS({refs['kCvFovYDeg']}/2))*{refs['kCvWidth']}/{refs['kCvHeight']}))", "0.0",
         "square pixels: tan(h/2) = tan(v/2) x w/h"),
        ("Threshold floor as heat (0..1)", f"={refs['floor']}/255", "0.00", "8-bit level / 255"),
        ("Core pixel as heat (0..1)", f"={refs['core']}/255", "0.00", "8-bit level / 255"),
        ("Largest accepted blob, share of the frame", f"={refs['maxArea']}/({refs['kCvWidth']}*{refs['kCvHeight']})", "0.0%", "maxArea / pixels"),
        ("Frames until a coasting track is deleted (60 Hz)", f"={refs['deleteAfter']}*60", "0", "deleteAfter x 60"),
        ("Frames until a track starts to coast (60 Hz)", f"={refs['coastAfter']}*60", "0", "coastAfter x 60"),
        ("Position gate radius for sigma = 0.3 m (m)", f"=SQRT({refs['gateChi2']})*0.3", "0.00", "sqrt(chi2) x sigma (sigma = typical single-ray error)"),
    ]
    for lab, f, fmt, note in d:
        put(ws, r, 1, lab)
        put(ws, r, 2, f, fmt)
        put(ws, r, 3, note, fill=NOTE_FILL)
        r += 1
    r += 2
    section(ws, r, 1, "D. Ground footprint of the camera and measurement gate vs altitude / sigma (formulas)")
    header(ws, r + 1, ["UAV altitude (m, example input)", "Footprint height at nadir (m) = 2 h tan(fov/2)",
                       "Metres per pixel at nadir = footprint / rows", "sigma (m, example input)", "Gate radius (m) = sqrt(chi2) x sigma"])
    r += 2
    alts = [5.0, 7.5, 10.0, 15.0, 20.0]
    sigmas = [0.1, 0.2, 0.3, 0.5, 1.0]
    for a, sg in zip(alts, sigmas):
        put(ws, r, 1, a, "0.0", NOTE_FILL)
        put(ws, r, 2, f"=2*A{r}*TAN(RADIANS({refs['kCvFovYDeg']}/2))", "0.00")
        put(ws, r, 3, f"=B{r}/{refs['kCvHeight']}", "0.000")
        put(ws, r, 4, sg, "0.00", NOTE_FILL)
        put(ws, r, 5, f"=SQRT({refs['gateChi2']})*D{r}", "0.00")
        r += 1
    put(ws, r, 1, "Altitudes include the squad's patrol (kSquadPatrolAlt) and engage altitudes; sigmas are example inputs.",
        border=False)
    r += 3
    r, passed = _tests_table(ws, r, ctx, [f"T{n}" for n in range(130, 138)], "E. Tests T130-T137 (observed values)")
    if passed:
        ctx["refs"]["cv_tests"] = f"'{S_CV}'!${passed[0]}${passed[1:]}"
    _fig(ws, r, 1, "report/figures_v5/thermal_fpv_tracks.png", "UAV thermal view with all tracks (HUD boxes)")
    _fig(ws, r, 6, "report/figures_v5/tactical_map.png", "Tactical map")
    widths(ws, [44, 22, 46, 46, 34, 4, 4, 4])
    finish(ws, 3, freeze="A4")
    ws.print_title_rows = None


# ---- 19 Squad & Mission --------------------------------------------------------------------------------------
ROLES = [
    ("PATROL", "No confirmed target", "Both UAVs fly the compound 180 deg apart on a breathing orbit (opposite phases); the cameras sweep outward."),
    ("LEAD", "UAV with the best view of the engaged track", "Designator: keeps its gimbal on the predicted target position and launches against it."),
    ("WING", "The other UAV", "Flies about 90 deg round the target (good triangulation geometry, second attack axis); a detour orbit when it has no clear shot."),
    ("SEARCH", "Track lost / target behind cover", "Circles the last known position for kSquadReacquireTime."),
    ("MANUAL", "Possessed by the user", "Only its camera contributes to the tracker; the squad plans around it."),
]
PHASES = [
    ("None", "no target"),
    ("Tracking", "a confirmed track is being held; locks accumulate"),
    ("Planned", "both rounds' flight times are predicted, the farther UAV launches first"),
    ("Inbound", "at least one round is in flight; arrival spread logged"),
]


def build_squad(wb, ctx):
    ws = wb[S_SQUAD]
    v5 = ctx["scene"].get("v5")
    if not v5:
        put(ws, 4, 1, "PENDING - scene_export.json has no 'v5' section (re-run `make docs`)")
        return
    title(ws, "19  Squad & Mission",
          "The two-UAV squad (roles, strike rules, time-on-target) and the night-watch mission (waves, scoring, base "
          "integrity). Constants are exported (blue); rule text is from the Squad.h / Mission.h header comments "
          "(pale yellow); time-on-target, waves, scoring and armour are Excel formulas; the measured results are the "
          "program's test observations (T138-T145, T180, T181, T105) plus one run reported by the supervisor.", span=10)
    refs = {}
    r = 4
    section(ws, r, 1, "A. Roles and strike phases (Squad.h)")
    header(ws, r + 1, ["Role", "When", "Behaviour"])
    r += 2
    for a, b, c in ROLES:
        put(ws, r, 1, a, font=BOLD)
        put(ws, r, 2, b, fill=NOTE_FILL)
        put(ws, r, 3, c, fill=NOTE_FILL)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        ws.row_dimensions[r].height = 34
        r += 1
    r += 1
    header(ws, r, ["Strike phase", "Meaning"])
    for a, b in PHASES:
        r += 1
        put(ws, r, 1, a, font=BOLD)
        put(ws, r, 2, b, fill=NOTE_FILL)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    r += 3
    section(ws, r, 1, "B. Squad constants (exported)")
    r = _const_table(ws, r + 1, v5["squad"], refs)
    ctx["refs"]["uavs"] = f"'{S_SQUAD}'!{refs['Squad::kCount']}"
    r += 1
    section(ws, r, 1, "Strike rules (text from the Squad.h header; thresholds above)")
    rules = [
        ("Combined strike", "when both UAVs have held the target (lead lock >= kSquadLeadLock, wing lock >= kSquadWingLock) "
                            "each round's flight time is predicted; the farther UAV launches first, the nearer later, so both rounds arrive together."),
        ("Solo fallback", "if the wing has no lock after kSquadWingWait, the lead strikes alone after kSquadSoloLock."),
        ("Lead prediction", "each round homes on the target's position predicted at its arrival time (Kalman velocity x flight time)."),
        ("Threat ranking", "the confirmed track closest to a camp asset is engaged first; the squad only sees what the cameras see."),
        ("Line of fire", "clearShot() checks the straight path against static solids; a UAV without a clear shot detours (orbit angle / height)."),
        ("Safety rules", "no rounds at targets inside the bunker-ramp keep-out box, nor within kFriendlyRadius of a friendly vehicle; hot wrecks are ignored (sheet 16)."),
    ]
    for i, (a, b) in enumerate(rules):
        put(ws, r + 1 + i, 1, a, font=BOLD)
        put(ws, r + 1 + i, 2, b, fill=NOTE_FILL)
        ws.merge_cells(start_row=r + 1 + i, start_column=2, end_row=r + 1 + i, end_column=8)
        ws.row_dimensions[r + 1 + i].height = 32
    r += len(rules) + 3

    # ---- time on target ----
    fm = v5["flight_model"]
    section(ws, r, 1, "C. Time-on-target: round flight time over a distance (formula) vs the code's roundFlightTime(d)")
    header(ws, r + 1, ["Model parameter", "Value", "Unit", "Source"])
    r += 2
    mp = {}
    for key, val, unit, src in (("v0 (launch speed)", fm["launch_speed"], "m/s", "Squad.cpp comment"),
                                ("a (acceleration during ramp)", fm["accel"], "m/s2", "Squad.cpp comment"),
                                ("t_ramp", fm["ramp_time"], "s", "Squad.cpp comment (= kMissileRampTime)"),
                                ("v_max (cruise)", None, "m/s", "Combat.h kMissileMaxSpeed (exported, sheet 16)"),
                                ("fuse", None, "m", "Combat.h kMissileFuse (exported, sheet 16)")):
        put(ws, r, 1, key, font=BOLD)
        if val is None:
            cr = ctx["refs"].get("combat_rows", {})
            name = "kMissileMaxSpeed" if "v_max" in key else "kMissileFuse"
            put(ws, r, 2, f"='{S_COMBAT}'!{cr[name]}" if name in cr else 0, "0.0")
        else:
            put(ws, r, 2, val, "0.0", INPUT_FILL)
        put(ws, r, 3, unit)
        put(ws, r, 4, src)
        mp[key[:2]] = f"$B${r}"
        r += 1
    v0, a, tr, vm, fuse = (mp[k] for k in ("v0", "a ", "t_", "v_", "fu"))
    put(ws, r, 1, "Ramp distance = (v0 + v_max)/2 x t_ramp", font=BOLD)
    put(ws, r, 2, f"=({v0}+{vm})/2*{tr}", "0.00", font=BOLD)
    put(ws, r, 3, "m")
    ramp = f"$B${r}"
    r += 2
    header(ws, r, ["Distance d (m, exported list)", "d' = max(0, d - fuse) (m)", "Flight time t (s) - formula",
                   "roundFlightTime(d) from the code (s)", "|difference| (s)", "Average speed d / t (m/s)"], height=46)
    r += 1
    t_first = r
    for pt in v5["flight_time"]:
        put(ws, r, 1, pt["d"], "0.0", INPUT_FILL)
        put(ws, r, 2, f"=MAX(0,A{r}-{fuse})", "0.00")
        put(ws, r, 3, f"=IF(B{r}<={ramp},(-{v0}+SQRT({v0}^2+2*{a}*B{r}))/{a},{tr}+(B{r}-{ramp})/{vm})", "0.0000")
        put(ws, r, 4, pt["t"], "0.0000", INPUT_FILL)
        put(ws, r, 5, f"=ABS(C{r}-D{r})", "0.0E+00")
        put(ws, r, 6, f"=A{r}/C{r}", "0.0")
        r += 1
    t_last = r - 1
    put(ws, r, 1, "max difference", fill=TOTAL_FILL, font=BOLD)
    put(ws, r, 5, f"=MAX(E{t_first}:E{t_last})", "0.0E+00", TOTAL_FILL, BOLD)
    put(ws, r, 6, f'=IF(E{r}<0.001,"formula = code","CHECK")', fill=TOTAL_FILL, font=BOLD)
    ws.conditional_formatting.add(f"F{r}", CellIsRule(operator="equal", formula=['"formula = code"'], fill=PASS_FILL))
    ctx["refs"]["tot_check"] = f"'{S_SQUAD}'!$F${r}"
    r += 3
    header(ws, r, ["Combined strike example", "Lead distance (m, example input)", "Wing distance (m, example input)",
                   "Lead flight time (s)", "Wing flight time (s)", "Launch delay of the nearer UAV (s)", "Arrival spread after planning (s)"],
           height=46)
    r += 1
    for i, (d1, d2) in enumerate(((8, 12), (10, 10), (6, 20), (15, 25), (12, 30))):
        put(ws, r, 1, f"example {i + 1}", font=BOLD)
        put(ws, r, 2, d1, "0.0", NOTE_FILL)
        put(ws, r, 3, d2, "0.0", NOTE_FILL)
        for c_i, src in ((4, "B"), (5, "C")):
            dd = f"MAX(0,{src}{r}-{fuse})"
            put(ws, r, c_i, f"=IF({dd}<={ramp},(-{v0}+SQRT({v0}^2+2*{a}*{dd}))/{a},{tr}+({dd}-{ramp})/{vm})", "0.000")
        put(ws, r, 6, f"=ABS(D{r}-E{r})", "0.000")
        put(ws, r, 7, f"=ABS((D{r}+IF(D{r}<E{r},F{r},0))-(E{r}+IF(E{r}<D{r},F{r},0)))", "0.000")
        r += 1
    put(ws, r, 1, "The farther UAV launches first; the nearer one waits |T_lead - T_wing|, so both rounds arrive together "
                  "(the in-game spread is measured by T141: closest blast pair 0.00 s).", border=False)
    r += 3

    # ---- waves ----
    mi = _consts(ctx, "mission")
    mr = {}
    section(ws, r, 1, "D. Mission constants (exported)")
    r = _const_table(ws, r + 1, v5["mission"], mr)
    r += 1
    section(ws, r, 1, "Mission waves (formulas): wave n sends MIN(2 + n, kMaxRovers) rovers, one every kSpawnInterval s")
    header(ws, r + 1, ["Wave n", "Rovers = MIN(2 + n, max)", "Spawn time of the wave (s) = (rovers - 1) x interval",
                       "Max score (all combined kills)", "Min score (all solo kills)", "Integrity lost if none is stopped (%)",
                       "Cumulative rovers", "Cumulative max score"], height=58)
    r += 2
    w_first = r
    for n in range(1, 11):
        put(ws, r, 1, n, "0", NOTE_FILL)
        put(ws, r, 2, f"=MIN(2+A{r},{mr['kMaxRovers']})", "0")
        put(ws, r, 3, f"=(B{r}-1)*{mr['kSpawnInterval']}", "0.0")
        put(ws, r, 4, f"=B{r}*({mr['kKillScore']}+{mr['kCombinedBonus']})", "0")
        put(ws, r, 5, f"=B{r}*{mr['kKillScore']}", "0")
        put(ws, r, 6, f"=B{r}*{mr['kBreachDamage']}", "0")
        put(ws, r, 7, f"=SUM(B${w_first}:B{r})", "0")
        put(ws, r, 8, f"=SUM(D${w_first}:D{r})", "0")
        r += 1
    w_last = r - 1
    put(ws, r, 1, "Breaches that destroy the base", font=BOLD)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    put(ws, r, 4, f"=100/{mr['kBreachDamage']}", "0")
    ctx["refs"]["breaches_to_lose"] = f"'{S_SQUAD}'!$D${r}"
    r += 3

    # ---- armour ----
    ar = {}
    section(ws, r, 1, "E. Armour and the rule '4-5 shots = 1 kill' (exported constants; formulas)")
    r = _const_table(ws, r + 1, v5["armour"], ar)
    r += 1
    header(ws, r, ["Blast falloff (1 = centre / direct hit, 0.5 = edge of the radius; example input)", "Damage per round (HP)",
                   "Rounds to kill = ROUNDUP(health / damage)", "Rounds with 2 UAVs firing together (pairs)"], height=58)
    r += 1
    a_first = r
    for fo in (1.0, 0.9, 0.8, 0.7, 0.6, 0.5):
        put(ws, r, 1, fo, "0.00", NOTE_FILL)
        put(ws, r, 2, f"={ar['kRoundDamage']}*1*A{r}", "0.0")
        put(ws, r, 3, f"=ROUNDUP({ar['kRoverHealth']}/B{r},0)", "0")
        put(ws, r, 4, f"=ROUNDUP(C{r}/2,0)", "0")
        r += 1
    put(ws, r, 1, "Rounds to kill at the centre; the user's rule is 4-5 rounds per kill (test T105: 5 hits observed on a real engagement).",
        border=False)
    ctx["refs"]["rounds_to_kill"] = f"'{S_SQUAD}'!$C${a_first}"
    r += 1
    put(ws, r, 1, "Reload of one used rail round (s)", font=BOLD)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    put(ws, r, 4, f"={ar['kUavReloadTime']}", "0.0")
    r += 3

    # ---- measured ----
    section(ws, r, 1, "F. Measured results")
    r, p1 = _tests_table(ws, r + 1, ctx, ["T138", "T139", "T140", "T141", "T142", "T143", "T144", "T145", "T180", "T181", "T105"],
                         "Program tests (observed values)")
    section(ws, r, 1, "Headless real-render run reported by the supervisor (90 s, not exported by this tool)")
    header(ws, r + 1, ["Quantity", "Value", "Source"])
    run = [("Rail rounds launched", 21), ("Hits", 12), ("Kills", 3), ("Combined kills", 2), ("Wave reached", 2),
           ("Fused (two-view) tracker updates", 689)]
    r += 2
    run_first = r
    for lab, v in run:
        put(ws, r, 1, lab, font=BOLD)
        put(ws, r, 2, v, "0", NOTE_FILL)
        put(ws, r, 3, "supervisor message (docs/PLAN_V5.md status): headless real render, 90 s", fill=NOTE_FILL)
        r += 1
    put(ws, r, 1, "Hits per kill", font=BOLD)
    put(ws, r, 2, f"=B{run_first + 1}/B{run_first + 2}", "0.0")
    put(ws, r, 3, "formula: hits / kills (the 4-5 shots rule; unarmoured misses are rounds minus hits)")
    r += 1
    put(ws, r, 1, "Hit rate", font=BOLD)
    put(ws, r, 2, f"=B{run_first + 1}/B{run_first}", "0.0%")
    put(ws, r, 3, "formula: hits / rounds launched")
    r += 3
    _fig(ws, r, 1, "report/figures_v5/combined_strike.png", "Combined strike (both rounds in flight)")
    _fig(ws, r, 6, "report/figures_v5/kill_banner.png", "Kill banner and score")
    widths(ws, [46, 26, 30, 22, 22, 24, 22, 22, 4])
    finish(ws, 3, freeze="A4")
    ws.print_title_rows = None


# ---- 20 Audio & HUD -------------------------------------------------------------------------------------------
SYNTHS = [
    # (sound, trigger, technique, key parameters) - docs/v5_notes/v5-a5.md
    ("UAV drone hum", "continuous per UAV", "4 detuned rotors x 6 blade-pass harmonics (amplitude h^-1.3) + band-passed noise 1.8 kHz, slow wobble",
     "f = 105 + 7 v Hz (v = speed m/s), speed smoothed 50 ms; reference distance 4 m"),
    ("Helicopter", "continuous (rotor %)", "blade-slap train (55 Hz decaying sine + LP-noise click) + 2.4 kHz turbine whine + LP noise",
     "slap rate 25 Hz x rotor fraction; ref distance 8 m"),
    ("Tank engine / tracks", "continuous (speed)", "firing-pulse train through 300 Hz LP + 130 Hz tremolo noise + track clatter (HF impulses)",
     "16 Hz + 3 Hz per m/s; ref distance 6 m"),
    ("Rounds in flight", "projectiles (max 6)", "missile motor roar / shell whistle: band-passed noise, Doppler shifted",
     "positions matched between frames; Doppler = 0.5 x physical, clamped 0.9..1.1"),
    ("Launch whoosh", "RoundLaunched", "SVF band-pass noise (Q 1.4) centre 350 -> 2750 Hz then fall + decaying crackle",
     "rise 30 %; ref distance 5 m"),
    ("Explosion", "Explosion / EnemyDestroyed", "HP crack (6 ms) + sine boom 100 -> 30 Hz + noise body through a closing SVF LP (4 kHz -> 250 Hz) + 90 Hz rumble tail",
     "length 2 + 1.5 p s (max 7), peak 0.3 + 0.45 sqrt(p); ref distance 25 m"),
    ("Tank cannon", "CannonFire", "crack + 95 -> 42 Hz boom + body + tail", "ref distance 12 m"),
    ("Gatling burst", "GatlingBurst", "1.4 s, 55 shots/s with spin-up and jitter: HP click + 150 Hz thump + 310 Hz motor whirr", "ref distance 8 m"),
    ("Blast doors / barrier", "DoorMoving", "70 Hz saw-harmonic motor + whine + noise ramp, then clunk (55-70 Hz boom, noise tick, 430 / 790 Hz ring)",
     "ref distance 4 m"),
    ("Wave siren", "WaveStart", "650-1000 Hz triangle sweep, 2.4 s period, 6 s", "ref distance 12 m"),
    ("Breach klaxon", "Breach", "820 / 610 Hz alternating tones", "-"),
    ("Radio lock beeps", "LockOn", "2 x 1250 Hz beeps (not spatialised)", "-"),
    ("Radio confirm chirp", "EnemySpotted", "900 -> 1900 Hz chirp (not spatialised)", "-"),
    ("Kill notes", "EnemyDestroyed", "880 / 1175 / 1568 Hz notes (not spatialised)", "-"),
    ("Ambient night bed", "always on", "gusting wind (3 slow LFOs drive an SVF centre 250-600 Hz) + 7 crickets (4.3-5.2 kHz, 3-pulse chirps at 29 ms) + 50 Hz generator hum",
     "fixed pans; quiet bed about -32 dB RMS"),
]
MIXER = [
    ("Sample rate", "48 kHz float", "all synthesis; miniaudio resamples if the device differs"),
    ("Voices", "24", "one-shots + continuous sources; the quietest is stolen (a new sound quieter than all is dropped)"),
    ("Panning", "equal-power cos / sin law", "pan = dot(direction, listener right), faded to centre inside 1 m"),
    ("Distance gain", "ref / max(d, ref)", "per-source reference distance (see table above)"),
    ("Air absorption", "one-pole LP, cutoff 16 kHz x exp(-d / 90), min 1.2 kHz", "distance darkens the sound"),
    ("Doppler", "0.5 x physical, clamped 0.9..1.1", "listener velocity ignored above 40 m/s (camera cuts)"),
    ("Master chain", "master volume then tanh soft limiter, ceiling 0.88 (-1.1 dBFS)", "NaN-safe"),
    ("Live playback", "miniaudio device (Device.cpp), only when a device exists", "headless = silent"),
    ("Offline mix", "renderOffline() per video frame -> 16-bit WAV, muxed into the MP4 with ffmpeg (AAC 192 kbit/s)", "no sound files: everything is synthesised, like no image files"),
    ("Demo check", "RMS -22.0 dBFS, peak -4.46 dBFS (10 s demo mix)", "per-second RMS: ambient -32, launch -23, explosion -17, siren -21 dB"),
]
HUD = [
    ("Enemy markers", "red brackets sized by distance, 'ID n  d M', health bar, pulsing LOCK ring; amber diamonds on predicted lead points while a strike is inbound; edge arrows for off-screen tracks (coasting amber with '?')"),
    ("UAV tags", "'UAV-1 LEAD' / 'UAV-2 WING' / 'PATROL' / 'SEARCH' / 'MANUAL' in role colours, bracket round the drone, line drone -> target"),
    ("Top bar", "WAVE n, SCORE, KILLS (COMBINED c), ENEMIES alive, BASE INTEGRITY bar + %, strike indicator (PLANNED / INBOUND blinking), TRACKING ID n, NEXT WAVE IN n S"),
    ("Banners", "WAVE n INCOMING (siren), ENEMY SPOTTED, COMBINED STRIKE, +100 / COMBINED STRIKE +150 rising at the kill, white X hit marker (0.35 s), BREACH -10 % + red edge flash, BASE LOST dim then RESTARTING"),
    ("Tactical mini-map", "compound ring, 20 m ring, 5 camp assets, UAV triangles with heading, track dots (red confirmed, grey tentative, amber hollow coasting), target ring, strike lines, sweep line, viewer arrow; north-up, 180 px at 900p bottom-right"),
    ("Thermal FPV HUD", "boxes of all confirmed tracks (TRK n / LOCK n / LOST n, range, '2-VIEW' when triangulated), role + lock time box per UAV"),
    ("Possession HUD", "possessed panel, tank sight, AA lead ring, helicopter tapes, launcher panel (SAM-1 / MLRS-1), ammo / reload bars, combat feed"),
]


def _event_types():
    """GameEventType enumerators with their comments, parsed from src/scene/GameEvents.h."""
    try:
        text = open(path("src/scene/GameEvents.h"), encoding="utf-8").read()
    except OSError:
        return []
    block = re.search(r"enum class GameEventType \{(.*?)\};", text, re.S)
    out = []
    for line in (block.group(1) if block else "").splitlines():
        m = re.match(r"\s*(\w+),\s*(?://\s*(.*))?", line)
        if m:
            out.append((m.group(1), (m.group(2) or "").strip()))
    return out


def build_audio_hud(wb, ctx):
    ws = wb[S_AUDIO]
    title(ws, "20  Audio & HUD",
          "Procedural sound (no sound files, miniaudio only for the device) and the game HUD. Synth and mixer rows are "
          "hand-maintained text from docs/v5_notes/v5-a5.md and v5-h5.md (pale yellow); the game-event table is parsed "
          "from src/scene/GameEvents.h; test rows are the program's observed values (T160-T169 audio, T170-T178 HUD).",
          span=8)
    r = 4
    section(ws, r, 1, "A. Synthesised sounds (all 48 kHz float, seeded, deterministic)")
    header(ws, r + 1, ["Sound", "Trigger", "Technique", "Key parameters"])
    r += 2
    s_first = r
    for a, b, c, d in SYNTHS:
        put(ws, r, 1, a, font=BOLD)
        put(ws, r, 2, b, fill=NOTE_FILL)
        put(ws, r, 3, c, fill=NOTE_FILL)
        put(ws, r, 4, d, fill=NOTE_FILL)
        ws.row_dimensions[r].height = 46
        r += 1
    put(ws, r, 1, "Sounds", fill=TOTAL_FILL, font=BOLD)
    put(ws, r, 2, f"=COUNTA(A{s_first}:A{r - 1})", "0", TOTAL_FILL, BOLD)
    ctx["refs"]["synths"] = f"'{S_AUDIO}'!$B${r}"
    define_name(wb, "SynthCount", S_AUDIO, f"B{r}")
    r += 3
    section(ws, r, 1, "B. Mixer, spatialisation and delivery")
    header(ws, r + 1, ["Item", "Setting", "Note"])
    r += 2
    for a, b, c in MIXER:
        put(ws, r, 1, a, font=BOLD)
        put(ws, r, 2, b, fill=NOTE_FILL)
        put(ws, r, 3, c, fill=NOTE_FILL)
        ws.row_dimensions[r].height = 34
        r += 1
    r += 2
    section(ws, r, 1, "C. Game events that drive HUD banners and sounds (parsed from GameEvents.h)")
    header(ws, r + 1, ["Event", "Meaning (header comment)"])
    r += 2
    ev_first = r
    for name, note in _event_types():
        put(ws, r, 1, name, font=BOLD)
        put(ws, r, 2, note, fill=INPUT_FILL)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        r += 1
    put(ws, r, 1, "Event types", fill=TOTAL_FILL, font=BOLD)
    put(ws, r, 2, f"=COUNTA(A{ev_first}:A{r - 1})", "0", TOTAL_FILL, BOLD)
    ctx["refs"]["event_types"] = f"'{S_AUDIO}'!$B${r}"
    r += 3
    r, pa = _tests_table(ws, r, ctx, [t["id"] for t in ctx["tests"] if t["id"].startswith("T16")], "D. Audio tests (T160-T169)")
    section(ws, r, 1, "E. HUD elements (docs/v5_notes/v5-h5.md)")
    header(ws, r + 1, ["Element", "Content"])
    r += 2
    h_first = r
    for a, b in HUD:
        put(ws, r, 1, a, font=BOLD)
        put(ws, r, 2, b, fill=NOTE_FILL)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 60
        r += 1
    put(ws, r, 1, "HUD elements", fill=TOTAL_FILL, font=BOLD)
    put(ws, r, 2, f"=COUNTA(A{h_first}:A{r - 1})", "0", TOTAL_FILL, BOLD)
    ctx["refs"]["hud_elements"] = f"'{S_AUDIO}'!$B${r}"
    r += 3
    r, ph = _tests_table(ws, r, ctx, [f"T{n}" for n in range(170, 179)], "F. HUD tests (T170-T178)")
    _fig(ws, r, 1, "report/figures_v5/kill_banner.png", "Kill banner (+100 / COMBINED STRIKE +150)")
    _fig(ws, r, 3, "report/figures_v5/tactical_map.png", "Tactical mini-map")
    widths(ws, [30, 42, 60, 54, 4])
    finish(ws, 3, freeze="A4")
    ws.print_title_rows = None
