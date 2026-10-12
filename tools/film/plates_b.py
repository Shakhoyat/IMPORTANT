"""Plates of Act 2 (viewing) and Act 3 (light)."""
import math
import os
import re

import numpy as np
from matplotlib.patches import Circle, Polygon, Rectangle

from plates import (AMB, CYAN, DIM, FG, GREEN, LINE, MONO, PANEL, RED, ROOT, VIOLET, arrow, box, canvas, line, reveal, txt)  # noqa: F401


def head(title, sub=""):
    def draw(ax, now):
        txt(ax, 96, 110, title, size=50, weight="bold")
        if sub:
            txt(ax, 96, 168, sub, size=28, color=DIM)
    return draw


def workbook(sheet):
    from openpyxl import load_workbook
    return load_workbook(os.path.join(ROOT, "report", "workbooks", "cg_workbook_values.xlsx"), data_only=True)[sheet]


# ------------------------------------------------------------------------------------------------------------ viewing
def p_lookat(c):
    E, T = np.array([430.0, 700.0]), np.array([900.0, 470.0])
    f = (T - E) / np.linalg.norm(T - E)
    r = np.array([-f[1], f[0]])

    def given(ax, now):
        ax.add_patch(Circle(E, 16, fc=AMB, ec="none", zorder=6))
        txt(ax, E[0] - 30, E[1] + 48, "eye", size=30, color=AMB, ha="right")
        ax.add_patch(Circle(T, 12, fc=CYAN, ec="none", zorder=6))
        txt(ax, T[0] + 26, T[1] - 26, "target", size=30, color=CYAN)
        line(ax, [E[0], T[0]], [E[1], T[1]], color=DIM, lw=2, ls="--")
        arrow(ax, 180, 620, 180, 380, color=DIM, lw=3)
        txt(ax, 198, 380, "world up", size=24, color=DIM)

    def axes(ax, now):
        arrow(ax, E[0], E[1], E[0] + f[0] * 230, E[1] + f[1] * 230, color=CYAN, lw=5)
        txt(ax, E[0] + f[0] * 245, E[1] + f[1] * 245 - 8, "forward", size=27, color=CYAN)
        arrow(ax, E[0], E[1], E[0] + r[0] * 190, E[1] + r[1] * 190, color=RED, lw=5)
        txt(ax, E[0] + r[0] * 205, E[1] + r[1] * 205 + 22, "right", size=27, color=RED)
        arrow(ax, E[0], E[1], E[0] - 30, E[1] - 215, color=GREEN, lw=5)
        txt(ax, E[0] - 44, E[1] - 240, "up", size=27, color=GREEN, ha="right")

    def maths(ax, now):
        box(ax, 1060, 250, 770, 560, ec=AMB if now else LINE)
        rows = [(r"$\mathbf{f}=\mathrm{normalise}(\mathbf{target}-\mathbf{eye})$", CYAN),
                (r"$\mathbf{r}=\mathrm{normalise}(\mathbf{f}\times\mathbf{up})$", RED),
                (r"$\mathbf{u}=\mathbf{r}\times\mathbf{f}$", GREEN)]
        for i, (s, col) in enumerate(rows):
            txt(ax, 1100, 320 + i * 74, s, size=36, color=col)
        txt(ax, 1100, 570, r"$V=\begin{pmatrix}r_x&r_y&r_z&-\mathbf{r}\cdot\mathbf{e}\\u_x&u_y&u_z&-\mathbf{u}\cdot\mathbf{e}\\-f_x&-f_y&-f_z&\ \mathbf{f}\cdot\mathbf{e}\\0&0&0&1\end{pmatrix}$"
            if False else r"$V=[\ \mathbf{r}\ |\ \mathbf{u}\ |\ -\mathbf{f}\ ]^{T}$  then move the eye to the origin", size=30)
        txt(ax, 1100, 650, "rows of V:   r,  u,  -f", size=28, color=DIM)
        txt(ax, 1100, 705, "last column:   -r.e,  -u.e,  f.e", size=28, color=DIM)
        txt(ax, 1100, 765, "written in src/math/Math3D.cpp (lookAt)", size=22, color=DIM)
    return reveal(c, [(0, head("The camera is one matrix")), (c.t(1), given), (c.t(2), axes), (c.t(2) + 3.5, maths)])


def p_project(c):
    ws = workbook("Transform_Trace")
    rows = [(str(ws.cell(r, 7).value), [ws.cell(r, k).value for k in (12, 13, 14, 15)]) for r in range(38, 44)]
    ex, ey, n, far = 150.0, 560.0, 190.0, 620.0

    def side(ax, now):
        ax.add_patch(Polygon([(ex, ey), (ex + far, ey - 270), (ex + far, ey + 270)], fc="#13233a", ec=LINE, lw=2, zorder=1))
        line(ax, [ex + n, ex + n], [ey - 83, ey + 83], color=FG, lw=4)
        txt(ax, ex + n, ey + 118, "screen", size=22, color=DIM, ha="center")
        ax.add_patch(Circle((ex, ey), 13, fc=AMB, ec="none", zorder=6))
        txt(ax, ex, ey + 46, "eye", size=24, color=AMB, ha="center")
        line(ax, [ex, ex + far + 30], [ey, ey], color=DIM, lw=1.5, ls="--")

    def point(ax, now):
        for z, col in ((300.0, CYAN), (600.0, VIOLET)):
            y = 150.0
            ax.add_patch(Circle((ex + z, ey - y), 11, fc=col, ec="none", zorder=6))
            line(ax, [ex + z, ex + z], [ey, ey - y], color=col, lw=3)
            line(ax, [ex, ex + z], [ey, ey - y], color=col, lw=1.6, alpha=0.8)
            ax.add_patch(Circle((ex + n, ey - y * n / z), 7, fc=col, ec=FG, lw=1.5, zorder=7))
        txt(ax, 150, 880, r"$x_{screen}=\dfrac{n\,x}{-z}\qquad y_{screen}=\dfrac{n\,y}{-z}$", size=40)
        txt(ax, 150, 960, "the same height, twice as far: half as tall", size=26, color=DIM)

    def ndc(ax, now):
        box(ax, 880, 240, 240, 240, fc="#13233a", ec=AMB if now else LINE, r=4)
        txt(ax, 1000, 360, "-1 ... +1", size=30, ha="center")
        txt(ax, 1000, 520, "the visible pyramid\nbecomes a cube", size=22, color=DIM, ha="center", linespacing=1.4)

    def table(ax, now):
        x0, y0 = 1190, 235
        box(ax, x0 - 20, y0 - 40, 690, 500, ec=AMB if now else LINE)
        for j, h in enumerate(("x", "y", "z", "w")):
            txt(ax, x0 + 318 + j * 100, y0, h, size=24, color=DIM, ha="right", weight="bold")
        names = ["cube corner", "world", "eye space", "clip space", "after divide", "pixel"]
        for i, (_, v) in enumerate(rows):
            y = y0 + 62 + i * 64
            txt(ax, x0, y, names[i], size=24, color=AMB if i in (4, 5) else FG)
            for j, q in enumerate(v):
                s = "" if not isinstance(q, (int, float)) else ("%.1f" % q if i == 5 and j < 2 else "%.3f" % q)
                txt(ax, x0 + 318 + j * 100, y, s, size=21, ha="right", mono=True)
        txt(ax, x0, y0 + 440, "one corner of a drone arm (workbook: Transform_Trace)", size=20, color=DIM)
    return reveal(c, [(0, head("Projection")), (0.2, side), (c.t(1), point), (c.t(3), ndc), (c.t(4), table)])


def p_clipping(c):
    x0, y0, w, h = 380, 330, 560, 380

    def window(ax, now):
        ax.add_patch(Rectangle((x0, y0), w, h, fc="#13233a", ec=FG, lw=3, zorder=2))
        txt(ax, x0 + w / 2, y0 + h + 44, "the window", size=24, color=DIM, ha="center")
        line(ax, [250, 840], [820, 400], color=CYAN, lw=4, z=5)

    def codes(ax, now):
        for xx in (x0, x0 + w):
            line(ax, [xx, xx], [y0 - 150, y0 + h + 150], color=LINE, lw=1.5, ls="--", z=1)
        for yy in (y0, y0 + h):
            line(ax, [x0 - 250, x0 + w + 250], [yy, yy], color=LINE, lw=1.5, ls="--", z=1)
        grid = [["1001", "1000", "1010"], ["0001", "0000", "0010"], ["0101", "0100", "0110"]]
        for r in range(3):
            for q in range(3):
                txt(ax, x0 - 125 + q * (w / 2 + 125), y0 - 80 + r * (h / 2 + 80), grid[r][q], size=26, color=AMB if (r, q) != (1, 1) else GREEN, ha="center", mono=True)
        txt(ax, x0 - 240, y0 + h + 200, "bits: above, below, right, left", size=24, color=DIM)

    def accept(ax, now):
        line(ax, [470, 800], [420, 520], color=GREEN, lw=5, z=6)
        txt(ax, 1100, 330, "0000 and 0000", size=30, mono=True, color=GREEN)
        txt(ax, 1420, 330, "keep it all", size=28)

    def reject(ax, now):
        line(ax, [1000, 1090], [250, 640], color=RED, lw=5, z=6)
        txt(ax, 1100, 420, "1010 and 0010", size=30, mono=True, color=RED)
        txt(ax, 1420, 420, "share a bit: drop it", size=28)

    def liang(ax, now):
        line(ax, [380, 840], [723, 400], color=AMB, lw=8, z=7)
        for (px, py), s in (((380, 723), "$t_0$"), ((840, 400), "$t_1$")):
            ax.add_patch(Circle((px, py), 11, fc=AMB, ec=FG, lw=2, zorder=8))
            txt(ax, px + 16, py - 34, s, size=30, color=AMB)
        box(ax, 1080, 520, 740, 250, ec=AMB)
        txt(ax, 1110, 580, r"$\mathbf{p}(t)=\mathbf{p}_0+t\,(\mathbf{p}_1-\mathbf{p}_0)$", size=34)
        txt(ax, 1110, 650, "each edge of the window narrows the interval", size=24, color=DIM)
        txt(ax, 1110, 710, r"keep $t_0\leq t\leq t_1$", size=34, color=AMB)
    return reveal(c, [(0, head("Clipping", "what lies outside the view is cut away before it is drawn")), (c.t(1), window), (c.t(2), codes),
                      (c.t(3), accept), (c.t(4), reject), (c.t(5), liang)])


def _grid(ax, x0, y0, nx, ny, s, col=LINE):
    for i in range(nx + 1):
        line(ax, [x0 + i * s, x0 + i * s], [y0, y0 + ny * s], color=col, lw=1, z=2)
    for j in range(ny + 1):
        line(ax, [x0, x0 + nx * s], [y0 + j * s, y0 + j * s], color=col, lw=1, z=2)


def _inside(p, tri):
    a, b, cc = tri
    e = lambda u, v: (p[0] - u[0]) * (v[1] - u[1]) - (p[1] - u[1]) * (v[0] - u[0])  # noqa: E731
    w = (e(b, cc), e(cc, a), e(a, b))
    return (min(w) >= 0 or max(w) <= 0), w


def p_zbuffer(c):
    nx, ny, s, x0, y0 = 14, 9, 62, 130, 250
    far = [(1.2, 7.6), (9.5, 8.2), (5.0, 0.8)]
    near = [(6.0, 8.4), (13.2, 5.0), (7.2, 1.6)]

    def fill(ax, tri, col, z, only_if=None):
        for j in range(ny):
            for i in range(nx):
                ok, _ = _inside((i + .5, j + .5), tri)
                if ok and (only_if is None or only_if(i, j)):
                    ax.add_patch(Rectangle((x0 + i * s + 2, y0 + j * s + 2), s - 4, s - 4, fc=col, ec="none", zorder=3))
                    txt(ax, x0 + i * s + s / 2, y0 + j * s + s / 2, z, size=17, ha="center", color="#0c121c", weight="bold")

    def first(ax, now):
        _grid(ax, x0, y0, nx, ny, s)
        fill(ax, far, "#e8963a", "0.7")
        txt(ax, x0, y0 + ny * s + 46, "each pixel stores a colour and a depth", size=26, color=DIM)

    def second(ax, now):
        fill(ax, near, "#3fa7e0", "0.3")

    def rule(ax, now):
        box(ax, 1080, 300, 760, 420, ec=AMB)
        txt(ax, 1110, 370, "for each fragment (x, y, z):", size=28, mono=True)
        txt(ax, 1110, 440, "  if z < depth[x, y]:", size=28, mono=True, color=AMB)
        txt(ax, 1110, 500, "      colour[x, y] = fragment colour", size=28, mono=True)
        txt(ax, 1110, 560, "      depth[x, y]  = z", size=28, mono=True)
        txt(ax, 1110, 650, "the order of drawing no longer matters", size=26, color=DIM)
    return reveal(c, [(0, head("Which surface do you see?")), (0.3, first), (c.t(1), second), (c.t(2), rule)])


def p_raster(c):
    nx, ny, s, x0, y0 = 18, 11, 56, 110, 240
    tri = [(2.2, 9.6), (16.4, 7.2), (7.4, 1.0)]
    cols = [np.array([1.0, 0.25, 0.2]), np.array([0.2, 0.85, 0.3]), np.array([0.2, 0.45, 1.0])]
    P = [(x0 + x * s, y0 + y * s) for x, y in tri]

    def outline(ax, now):
        _grid(ax, x0, y0, nx, ny, s)
        ax.add_patch(Polygon(P, fc="none", ec=FG, lw=3, zorder=6))
        for (px, py), col in zip(P, cols):
            ax.add_patch(Circle((px, py), 12, fc=col, ec=FG, lw=2, zorder=7))

    def one_edge(ax, now):
        a, b = tri[0], tri[1]
        for j in range(ny):
            for i in range(nx):
                p = (i + .5, j + .5)
                if (p[0] - a[0]) * (b[1] - a[1]) - (p[1] - a[1]) * (b[0] - a[0]) > 0:
                    ax.add_patch(Circle((x0 + p[0] * s, y0 + p[1] * s), 5, fc=AMB, ec="none", zorder=4, alpha=0.7))
        txt(ax, 1180, 300, r"$E_{ab}(\mathbf{p})=(\mathbf{p}-\mathbf{a})\times(\mathbf{b}-\mathbf{a})$", size=34)
        txt(ax, 1180, 362, "positive on one side of the edge", size=24, color=AMB)

    def inside(ax, now, colour=False):
        area = _inside(tri[2], [tri[0], tri[1], tri[0]])[1]
        for j in range(ny):
            for i in range(nx):
                ok, w = _inside((i + .5, j + .5), tri)
                if ok:
                    tot = sum(w)
                    col = sum(cols[k] * (w[k] / tot) for k in range(3)) if colour else np.array([0.25, 0.5, 0.75])
                    ax.add_patch(Rectangle((x0 + i * s + 2, y0 + j * s + 2), s - 4, s - 4, fc=tuple(np.clip(col, 0, 1)), ec="none", zorder=3))
        del area

    def all3(ax, now):
        inside(ax, now, False)
        txt(ax, 1180, 470, "inside:  all three positive", size=30)

    def bary(ax, now):
        inside(ax, now, True)
        txt(ax, 1180, 580, r"$\lambda_i=E_i/(E_0+E_1+E_2)$", size=34)
        txt(ax, 1180, 650, r"colour $=\lambda_0c_0+\lambda_1c_1+\lambda_2c_2$", size=34, color=AMB)
        txt(ax, 1180, 720, "the same weights blend depth and normals", size=24, color=DIM)
    return reveal(c, [(0, head("Which pixels does a triangle cover?")), (0.3, outline), (c.t(1), one_edge), (c.t(2), all3), (c.t(3), bary)])


# -------------------------------------------------------------------------------------------------------------- light
def _ball(kind):
    n = 300
    y, x = np.mgrid[1:-1:n * 1j, -1:1:n * 1j]
    m = x * x + y * y <= 1
    z = np.sqrt(np.clip(1 - x * x - y * y, 0, 1))
    N = np.dstack([x, y, z])
    L = np.array([-0.55, 0.6, 0.58])
    L = L / np.linalg.norm(L)
    Hh = L + np.array([0, 0, 1.0])
    Hh = Hh / np.linalg.norm(Hh)
    base = np.array([0.30, 0.52, 0.26])
    amb = 0.13 * base
    dif = np.clip(N @ L, 0, 1)[..., None] * base * 0.9
    spec = (np.clip(N @ Hh, 0, 1) ** 48)[..., None] * np.array([1.0, 1.0, 0.95]) * 0.9
    img = {"ambient": amb + 0 * dif, "diffuse": dif, "specular": spec, "all": amb + dif + spec}[kind]
    out = np.zeros((n, n, 4))
    out[..., :3] = np.clip(img, 0, 1) ** (1 / 2.2)
    out[..., 3] = m
    return out


def p_phong(c):
    def vectors(k):
        def draw(ax, now):
            px, py = 330, 720
            line(ax, [90, 570], [py, py], color=FG, lw=4)
            ax.add_patch(Circle((px, py), 10, fc=FG, ec="none", zorder=6))
            arrow(ax, px, py, px, py - 300, color=GREEN, lw=5)
            txt(ax, px + 14, py - 305, "N  normal", size=26, color=GREEN)
            if k >= 1:
                arrow(ax, px, py, px - 215, py - 215, color=AMB, lw=5)
                txt(ax, px - 232, py - 250, "L  to the light", size=26, color=AMB, ha="center")
                arrow(ax, px, py, px + 250, py - 150, color=CYAN, lw=5)
                txt(ax, px + 262, py - 150, "V  to the eye", size=26, color=CYAN)
            if k >= 2:
                arrow(ax, px, py, px + 30, py - 270, color=VIOLET, lw=4)
                txt(ax, px + 44, py - 262, "H  half-way", size=24, color=VIOLET)
        return draw

    def ball(i, kind, label, formula):
        def draw(ax, now):
            cx = 905 + i * 262
            ax.imshow(_ball(kind), extent=(cx - 112, cx + 112, 507, 283), zorder=3)
            ax.set_xlim(0, 1920)
            ax.set_ylim(1080, 0)
            txt(ax, cx, 566, label, size=28, ha="center", weight="bold", color=AMB if now else FG)
            txt(ax, cx, 624, formula, size=25, ha="center")
            if i:
                txt(ax, cx - 131, 395, "+" if i < 3 else "=", size=46, ha="center", color=DIM)
        return draw

    def eq(ax, now):
        box(ax, 96, 890, 1728, 130, ec=AMB)
        txt(ax, 140, 955, r"$I=k_a\,I_a+k_d\,(\mathbf{N}\cdot\mathbf{L})\,I_l+k_s\,(\mathbf{N}\cdot\mathbf{H})^{n}\,I_l$", size=42)
        txt(ax, 1280, 955, "summed over every light", size=26, color=DIM)
    return reveal(c, [(0, head("What colour is a point?")), (c.t(1), vectors(0)), (c.t(2), vectors(1)), (c.t(3), ball(0, "ambient", "ambient", r"$k_a I_a$")),
                      (c.t(4), ball(1, "diffuse", "diffuse", r"$k_d(\mathbf{N}\cdot\mathbf{L})$")), (c.t(5), vectors(2)),
                      (c.t(5) + 0.4, ball(2, "specular", "specular", r"$k_s(\mathbf{N}\cdot\mathbf{H})^n$")), (c.t(6), ball(3, "all", "sum", "")), (c.t(6) + 0.4, eq)])


def p_tsmath(c):
    src = open(os.path.join(ROOT, "src", "rendering", "ShaderMath.cpp")).read()
    want = ["float tsDot(vec3", "float tsLength(vec3", "vec3 tsNormalize(vec3", "vec3 tsCross(", "vec3 tsReflect(", "float tsClamp(float", "vec3 tsMix(vec3 a, vec3 b, float", "float tsSmoothstep(float"]
    code = []
    for wnt in want:
        m = re.search(r"^.*" + re.escape(wnt) + r".*$", src, re.M)
        if m:
            code.append(m.group(0).strip()[:92])

    def listing(ax, now):
        box(ax, 96, 230, 1728, 70 + 62 * len(code), fc="#0f1826")
        for i, ln in enumerate(code):
            txt(ax, 130, 290 + i * 62, ln, size=23, mono=True, color="#cfe3ff")
        txt(ax, 130, 230 + 70 + 62 * len(code) + 40, "src/rendering/ShaderMath.cpp: placed in front of every vertex, fragment and compute shader", size=22, color=DIM)

    def refuse(ax, now):
        box(ax, 96, 860, 1728, 150, ec=RED)
        txt(ax, 140, 912, "vec3 n = normalize(v);", size=28, mono=True, color=RED)
        line(ax, [130, 560], [912, 912], color=RED, lw=3)
        txt(ax, 640, 912, "refused at compile time", size=28, color=RED)
        txt(ax, 140, 968, "905 calls of the written-out functions in 13 shader source files", size=26, color=FG)
    return reveal(c, [(0, head("No built-in vector functions")), (0.3, listing), (c.t(2), refuse)])


def p_shadowmap(c):
    Lx, Ly = 330.0, 250.0
    gy = 800.0

    def scene(ax, now):
        line(ax, [120, 1160], [gy, gy], color=FG, lw=4)
        ax.add_patch(Circle((Lx, Ly), 26, fc=AMB, ec="none", zorder=6))
        txt(ax, Lx - 40, Ly - 46, "light", size=26, color=AMB)
        ax.add_patch(Rectangle((560, 520), 190, 120, fc="#3a5a80", ec=FG, lw=2, zorder=5))
        for tx in np.linspace(420, 1150, 9):
            line(ax, [Lx, tx], [Ly, gy], color=AMB, lw=1, alpha=0.25)
        txt(ax, 120, 880, "1   draw the scene from the light, keep only depth", size=26, color=DIM)

    def compare(ax, now):
        for tx, name, lit in ((930.0, "P", False), (480.0, "Q", True)):
            col = GREEN if lit else RED
            ax.add_patch(Circle((tx, gy), 12, fc=col, ec=FG, lw=2, zorder=8))
            txt(ax, tx + 6, gy + 40, name, size=28, color=col, weight="bold")
            line(ax, [Lx, tx], [Ly, gy], color=col, lw=3)
        hx = Lx + (930 - Lx) * (520 - Ly) / (gy - Ly)
        ax.add_patch(Circle((hx, 520), 9, fc=FG, ec="none", zorder=9))
        txt(ax, hx + 18, 494, "stored depth", size=22)
        box(ax, 1240, 300, 590, 330, ec=AMB if now else LINE)
        txt(ax, 1270, 360, "2   for each pixel:", size=26, color=DIM)
        txt(ax, 1270, 430, "distance to the light", size=28)
        txt(ax, 1270, 490, ">  stored depth + bias ?", size=28, color=AMB)
        txt(ax, 1270, 570, "P: yes, in shadow     Q: no, lit", size=26)

    def soft(ax, now):
        ax.add_patch(Polygon([(905, gy + 4), (1010, gy + 4), (1010, gy + 22), (905, gy + 22)], fc=VIOLET, ec="none", zorder=7, alpha=0.9))
        txt(ax, 1030, gy + 46, "soft edge", size=22, color=VIOLET)
        box(ax, 1240, 660, 590, 190, ec=VIOLET)
        txt(ax, 1270, 720, r"$w=\mathrm{gap}\times 2\tan(\theta/2)$", size=34)
        txt(ax, 1270, 790, "average many tests over a disc of width w", size=22, color=DIM)
    return reveal(c, [(0, head("Can this point see the light?")), (c.t(1), scene), (c.t(2), compare), (c.t(3), soft)])


def p_tonemap(c):
    def rng(ax, now):
        x0, x1, y = 160, 1760, 330
        line(ax, [x0, x1], [y, y], color=FG, lw=3)
        marks = [(-2, "moonlit\nground"), (0, "a lamp"), (2, "overcast\nsky"), (4, "sunlit\nground"), (6.2, "the sun's\ndisc")]
        for e, s in marks:
            x = x0 + (e + 3) / 10.0 * (x1 - x0)
            line(ax, [x, x], [y - 16, y + 16], color=FG, lw=3)
            txt(ax, x, y + 70, s, size=22, ha="center", color=DIM, linespacing=1.3)
            txt(ax, x, y - 44, "10^%d" % e if e == int(e) else "", size=20, ha="center", color=DIM)
        txt(ax, x0, y - 100, "light in the scene (relative luminance, logarithmic)", size=26)

    def window(ax, now):
        ax.add_patch(Rectangle((690, 290), 420, 80, fc=AMB, ec="none", alpha=0.28, zorder=1))
        txt(ax, 900, 258, "exposure: the part the eye is adapted to", size=22, color=AMB, ha="center")

    def curve(ax, now):
        bx, by, bw, bh = 470, 500, 620, 420
        box(ax, bx - 60, by - 30, bw + 500, bh + 90, fc="#101a28")
        line(ax, [bx, bx + bw], [by + bh, by + bh], color=DIM, lw=2)
        line(ax, [bx, bx], [by, by + bh], color=DIM, lw=2)
        x = np.linspace(0, 6, 200)
        t = x * (2.51 * x + 0.03) / (x * (2.43 * x + 0.59) + 0.14)
        g = np.clip(t, 0, 1) ** (1 / 2.2)
        line(ax, bx + x / 6 * bw, by + bh - np.clip(t, 0, 1) * bh, color=CYAN, lw=4)
        line(ax, bx + x / 6 * bw, by + bh - g * bh, color=AMB, lw=4)
        line(ax, bx + x / 6 * bw, by + bh - np.clip(x, 0, 1) * bh, color=RED, lw=2, ls="--")
        txt(ax, bx + bw + 40, by + 60, "tone curve, then gamma", size=26, color=AMB)
        txt(ax, bx + bw + 40, by + 120, "tone curve", size=26, color=CYAN)
        txt(ax, bx + bw + 40, by + 180, "no curve: clipped to white", size=26, color=RED)
        txt(ax, bx + bw / 2, by + bh + 36, "exposed light", size=22, color=DIM, ha="center")
        txt(ax, bx - 30, by - 4, "pixel", size=22, color=DIM, ha="right")
    return reveal(c, [(0, head("From light to a pixel value")), (0.3, rng), (c.t(1) + 3.0, window), (c.t(1) + 6.5, curve)])


def p_rays(c):
    eye = (170.0, 520.0)
    S1, r1 = (820.0, 560.0), 120.0      # mirror ball
    S2, r2 = (1160.0, 330.0), 90.0      # glass ball
    Lp = (760.0, 190.0)
    hit = (S1[0] - r1 * math.cos(math.radians(12)), S1[1] - r1 * math.sin(math.radians(12)))

    def base(ax, now):
        ax.add_patch(Circle(eye, 15, fc=AMB, ec="none", zorder=6))
        txt(ax, eye[0], eye[1] + 48, "eye", size=24, color=AMB, ha="center")
        for j in range(7):
            ax.add_patch(Rectangle((330, 400 + j * 34), 12, 30, fc="#3a5a80", ec=FG, lw=1, zorder=4))
        txt(ax, 336, 668, "pixels", size=22, color=DIM, ha="center")
        ax.add_patch(Circle(S1, r1, fc="#55708f", ec=FG, lw=2, zorder=3))
        ax.add_patch(Circle(S2, r2, fc="#2b6f8a", ec=FG, lw=2, zorder=3, alpha=0.7))
        ax.add_patch(Circle(Lp, 20, fc="#fff3b0", ec="none", zorder=6))
        txt(ax, Lp[0] + 30, Lp[1] - 14, "light", size=22, color="#fff3b0")
        line(ax, [120, 1420], [770, 770], color=FG, lw=3)

    def primary(ax, now):
        arrow(ax, eye[0], eye[1], hit[0], hit[1], color=FG, lw=4)
        txt(ax, 470, 490, "primary ray", size=22)

    def shadow(ax, now):
        arrow(ax, hit[0], hit[1], Lp[0], Lp[1], color="#fff3b0", lw=3)
        txt(ax, 590, 330, "to the light", size=22, color="#fff3b0")

    def reflect(ax, now):
        arrow(ax, hit[0], hit[1], 520, 250, color=CYAN, lw=4)
        txt(ax, 396, 236, "reflected", size=22, color=CYAN)
        txt(ax, 96, 880, r"$\mathbf{r}=\mathbf{d}-2\,(\mathbf{d}\cdot\mathbf{n})\,\mathbf{n}$", size=36, color=CYAN)

    def refract(ax, now):
        arrow(ax, 520, 250, S2[0] - r2, S2[1], color=CYAN, lw=3)
        arrow(ax, S2[0] - r2, S2[1], S2[0] + r2 * 0.9, S2[1] + 40, color=VIOLET, lw=4)
        arrow(ax, S2[0] + r2 * 0.9, S2[1] + 40, 1400, 470, color=VIOLET, lw=3)
        txt(ax, 1300, 300, "refracted", size=22, color=VIOLET)
        txt(ax, 700, 880, r"Snell: $\eta_1\sin\theta_1=\eta_2\sin\theta_2$", size=34, color=VIOLET)

    def bvh(ax, now):
        for (x, y, w, h), col in (((680, 220, 600, 480), GREEN), ((690, 430, 260, 260), AMB), ((1060, 230, 200, 200), AMB)):
            ax.add_patch(Rectangle((x, y), w, h, fc="none", ec=col, lw=2.5, ls=(0, (7, 5)), zorder=7))
        box(ax, 1470, 230, 380, 420, ec=GREEN)
        txt(ax, 1500, 290, "tree of boxes", size=28, weight="bold", color=GREEN)
        txt(ax, 1500, 350, "about 5,100 parts", size=24)
        txt(ax, 1500, 400, "a ray tests a few\nboxes, not every part", size=24, color=DIM, va="top", linespacing=1.4)
        txt(ax, 1500, 530, "each hit solved exactly\nin the part's own space", size=24, color=DIM, va="top", linespacing=1.4)
        txt(ax, 1210, 950, "at most 5 bounces", size=26, color=DIM)
    return reveal(c, [(0, head("Follow the light backwards")), (0.3, base), (c.t(2), primary), (c.t(3), shadow), (c.t(3) + 2.6, reflect),
                      (c.t(3) + 5.0, refract), (c.t(4) + 3.0, bvh)])
