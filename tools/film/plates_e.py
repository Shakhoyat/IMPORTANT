"""Plates of the second cut: the workbook, the plan of the video, polygon clipping and culling, procedural texture,
temporal anti-aliasing, Otsu's threshold, and the fixed simulation tick."""
import math
import os

import numpy as np
from matplotlib.patches import Circle, Polygon, Rectangle

from plates import (AMB, CYAN, DIM, FG, GREEN, LINE, PANEL, RED, ROOT, VIOLET, arrow, box, canvas, export, line, reveal, txt)  # noqa: F401
from plates_b import _grid, head


# ---------------------------------------------------------------------------------------------------------- opening
def p_workbook(c):
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(ROOT, "report", "workbooks", "cg_workbook_values.xlsx"), data_only=True, read_only=True)
    sheets = list(wb.sheetnames)
    rows, hdr, count = [], None, 0
    for i, r in enumerate(wb["World_Coordinates"].iter_rows(values_only=True)):
        if r and r[0] == "id":
            hdr = list(r)
        elif hdr and r and r[0]:
            count += 1
            if r[4] != "Group" and count % 430 == 7 and len(rows) < 11:
                rows.append(r)
    col = {n: k for k, n in enumerate(hdr)}
    show = [("id", 130, "left"), ("object", 250, "left"), ("component", 520, "left"), ("primitive", 860, "left"), ("world x", 1180, "right"),
            ("world y", 1330, "right"), ("world z", 1480, "right"), ("S x", 1600, "right"), ("S y", 1700, "right"), ("S z", 1800, "right")]

    def table(n):
        def draw(ax, now):
            box(ax, 96, 250, 1728, 640, fc="#0f1826")
            for name, x, ha in show:
                txt(ax, x, 292, name, size=22, color=AMB, weight="bold", ha=ha)
            line(ax, [116, 1804], [320, 320], color=LINE, lw=2)
            for i, r in enumerate(rows[:n]):
                for name, x, ha in show:
                    v = r[col[name]]
                    s = ("%.2f" % v) if isinstance(v, float) else str(v)
                    txt(ax, x, 352 + i * 49, s[:24], size=21, ha=ha, color=FG if name in ("id", "component") else "#cfe3ff")
        return draw

    def total(ax, now):
        txt(ax, 96, 950, "%s components" % format(count, ","), size=44, weight="bold", color=AMB)
        txt(ax, 560, 950, "one row each, on the sheet World_Coordinates", size=28)
        txt(ax, 96, 1010, "%d sheets:  " % len(sheets) + "  ·  ".join(s.replace("_", " ") for s in sheets[:9]) + "  ·  ...", size=20, color=DIM)
    return reveal(c, [(0, head("The workbook behind the project", "every point of every object, recorded in Excel while the project was built")),
                      (c.t(1), table(4)), (c.t(2), table(11)), (c.t(3), total)])


def p_roadmap(c):
    parts = [("Part 1", "The graphics lab", "primitives  ·  transformations  ·  viewing and projection  ·  clipping  ·  hidden surfaces\n"
              "rasterisation  ·  illumination  ·  shading  ·  texture  ·  ray tracing"),
             ("Part 2", "Beyond the lab", "microfacet materials  ·  soft shadows  ·  exposure and tone  ·  anti-aliasing  ·  reflections\n"
              "instancing and level of detail  ·  day, night and weather  ·  sound synthesis"),
             ("Part 3", "Take control", "walk into buildings  ·  operate twenty things  ·  jump into any point of view"),
             ("Part 4", "Machine vision", "thermal sensor  ·  threshold  ·  regions  ·  classification  ·  Kalman tracking"),
             ("Part 5", "The agents", "drones: lead, wing, time on target  ·  intruders: flood fill, line follower, hiding  ·  soldiers: A*")]
    ys = [230, 420, 610, 745, 880]
    hs = [165, 165, 110, 110, 110]

    def state(upto):
        def draw(ax, now):
            txt(ax, 96, 110, "The plan of the video", size=50, weight="bold")
            for i, (tag, name, sub) in enumerate(parts):
                on = i < upto or (upto >= 4 and i == 4)
                new = (i == upto - 1) or (upto == 4 and i >= 3)
                box(ax, 96, ys[i], 1728, hs[i], fc=PANEL if on else "#111926", ec=AMB if on and new else LINE, lw=3 if on else 1.5)
                txt(ax, 136, ys[i] + hs[i] / 2, tag, size=26, weight="bold", color=AMB if on else "#4a5a70")
                txt(ax, 262, ys[i] + hs[i] / 2, name, size=36, weight="bold", color=FG if on else "#4a5a70")
                txt(ax, 640, ys[i] + hs[i] / 2, sub, size=22, color=DIM if on else "#3a485c", linespacing=1.6)
        return draw
    out = []
    for k in range(5):
        fig, ax = canvas()
        state(k)(ax, True)
        out.append((c.t(k) if k else 0.0, c.save(fig)))
    return out


# ---------------------------------------------------------------------------------------------------------- viewing
def _clip_poly(poly, inside, cut):
    out = []
    for i in range(len(poly)):
        s, e = poly[i - 1], poly[i]
        if inside(e):
            if not inside(s):
                out.append(cut(s, e))
            out.append(e)
        elif inside(s):
            out.append(cut(s, e))
    return out


def p_polyclip(c):
    x0, y0, w, h = 150, 330, 560, 400
    tri = [(260.0, 660.0), (860.0, 520.0), (520.0, 230.0)]

    def cut_x(xc):
        return lambda s, e: (xc, s[1] + (e[1] - s[1]) * (xc - s[0]) / (e[0] - s[0]))

    def cut_y(yc):
        return lambda s, e: (s[0] + (e[0] - s[0]) * (yc - s[1]) / (e[1] - s[1]), yc)
    p1 = _clip_poly(tri, lambda p: p[0] <= x0 + w, cut_x(x0 + w))
    p2 = _clip_poly(p1, lambda p: p[1] >= y0, cut_y(y0))

    def window(ax, now):
        ax.add_patch(Rectangle((x0, y0), w, h, fc="#13233a", ec=FG, lw=3, zorder=2))
        txt(ax, x0 + w / 2, y0 + h + 44, "the window", size=24, color=DIM, ha="center")
        ax.add_patch(Polygon(tri, fc="none", ec=CYAN, lw=3, zorder=5, ls="--" if not now else "-"))

    def clipped(ax, now):
        ax.add_patch(Polygon(p2, fc="#1f6f8f", ec=AMB, lw=4, zorder=6, alpha=0.9))
        for p in p2:
            new = p not in tri
            ax.add_patch(Circle(p, 10, fc=AMB if new else FG, ec="none", zorder=8))
        txt(ax, x0, y0 + h + 110, "Sutherland-Hodgman: one border at a time", size=27, color=AMB)
        txt(ax, x0, y0 + h + 158, "3 corners in, %d corners out; the yellow ones are new" % len(p2), size=23, color=DIM)

    def cull(ax, now):
        ex, ey = 1380, 850
        box(ax, 980, 230, 840, 720, fc="#101a28")
        txt(ax, 1010, 280, "before that: is the whole object in view?", size=27, color=AMB)
        ax.add_patch(Polygon([(ex, ey), (ex - 300, ey - 500), (ex + 300, ey - 500)], fc="#17304a", ec=FG, lw=2.5, zorder=3))
        ax.add_patch(Circle((ex, ey), 13, fc=AMB, ec="none", zorder=6))
        txt(ax, ex, ey + 44, "eye", size=22, color=AMB, ha="center")
        for (x, y, r), ok in (((1360, 560, 46), True), ((1500, 470, 34), True), ((1120, 640, 44), False), ((1680, 760, 40), False), ((1420, 905, 0), None)):
            if ok is None:
                continue
            ax.add_patch(Circle((x, y), r, fc="none", ec=GREEN if ok else RED, lw=3.5, zorder=7))
            txt(ax, x, y, "draw" if ok else "skip", size=20, color=GREEN if ok else RED, ha="center")
        txt(ax, 1010, 330, "a sphere round each object, tested against the 6 planes", size=22, color=DIM)
    return reveal(c, [(0, head("Clipping a polygon", "a triangle that crosses the border of the view")), (0.3, window), (c.t(1) + 2.5, clipped), (c.t(2), cull)])


# ---------------------------------------------------------------------------------------------------------- texture
def _noise(n, cells, seed):
    rng = np.random.default_rng(seed)
    g = rng.random((cells + 1, cells + 1))
    u = np.linspace(0, cells, n, endpoint=False)
    i = u.astype(int)
    f = u - i
    s = f * f * (3 - 2 * f)
    a = g[i][:, i] * (1 - s)[None, :] + g[i][:, i + 1] * s[None, :]
    b = g[i + 1][:, i] * (1 - s)[None, :] + g[i + 1][:, i + 1] * s[None, :]
    return a * (1 - s)[:, None] + b * s[:, None]


def p_texture(c):
    n = 220
    octs = [_noise(n, 4 * 2 ** k, 11 + k) for k in range(4)]
    amp = [0.5, 0.25, 0.125, 0.0625]
    total = sum(a * o for a, o in zip(amp, octs)) / sum(amp)

    def tile(k):
        def draw(ax, now):
            x = 96 + k * 300
            ax.imshow(octs[k], extent=(x, x + 250, 560, 310), cmap="gray", vmin=0, vmax=1, zorder=3, interpolation="bicubic")
            ax.set_xlim(0, 1920)
            ax.set_ylim(1080, 0)
            txt(ax, x + 125, 600, "layer %d" % (k + 1), size=26, ha="center", weight="bold", color=AMB if now else FG)
            txt(ax, x + 125, 644, "detail x%d, strength 1/%d" % (2 ** k, 2 ** (k + 1)), size=20, ha="center", color=DIM)
            if k:
                txt(ax, x - 25, 435, "+", size=44, ha="center", color=DIM)
        return draw

    def result(ax, now):
        txt(ax, 1296, 435, "=", size=48, ha="center", color=DIM)
        col = np.stack([0.38 + 0.36 * total, 0.36 + 0.33 * total, 0.31 + 0.27 * total], axis=2)
        ax.imshow(np.clip(col, 0, 1), extent=(1340, 1820, 690, 210), zorder=3, interpolation="bicubic")
        ax.set_xlim(0, 1920)
        ax.set_ylim(1080, 0)
        txt(ax, 1580, 730, "the sum: a surface with character", size=26, ha="center", color=AMB)

    def proj(ax, now):
        box(ax, 96, 800, 1728, 200, ec=AMB)
        txt(ax, 140, 860, r"$\mathrm{fbm}(\mathbf{q})=\sum_{o=0}^{3}\ 0.5^{\,o+1}\ \mathrm{noise}(2.03^{\,o}\,\mathbf{q})$", size=38)
        txt(ax, 140, 945, "projected along the axis the surface faces most, so no part needs stored (u, v) and no pattern stretches", size=24, color=DIM)
    return reveal(c, [(0, head("Detail from a formula", "a procedural texture is computed where an image would be looked up")), (c.t(2), tile(0)), (c.t(2) + 2.2, tile(1)),
                      (c.t(2) + 4.2, tile(2)), (c.t(2) + 5.2, tile(3)), (c.t(3), result), (c.t(3) + 3.5, proj)])


# -------------------------------------------------------------------------------------------------------- rendering
def _halton(i, b):
    f, r = 1.0, 0.0
    while i > 0:
        f /= b
        r += f * (i % b)
        i //= b
    return r


def p_taa(c):
    nx, ny, s, x0, y0 = 12, 8, 62, 110, 270
    edge = lambda x, y: y - (7.4 - 0.52 * x)   # noqa: E731  (a slanted edge: inside below the line)
    pts = [(_halton(k, 2), _halton(k, 3)) for k in range(1, 9)]

    def cover(i, j, samples):
        return sum(1.0 for a, b in samples if edge(i + a, j + b) > 0) / len(samples)

    def raster(samples, label, x_off):
        def draw(ax, now):
            _grid(ax, x0 + x_off, y0, nx, ny, s)
            for j in range(ny):
                for i in range(nx):
                    v = cover(i, j, samples)
                    if v > 0:
                        ax.add_patch(Rectangle((x0 + x_off + i * s + 1, y0 + j * s + 1), s - 2, s - 2, fc=(0.25 * v + 0.05, 0.62 * v + 0.07, 0.9 * v + 0.1), ec="none", zorder=3))
            line(ax, [x0 + x_off, x0 + x_off + nx * s], [y0 + 7.4 * s, y0 + (7.4 - 0.52 * nx) * s], color=FG, lw=2, z=5, alpha=0.7)
            txt(ax, x0 + x_off + nx * s / 2, y0 + ny * s + 44, label, size=26, ha="center", color=AMB if now else FG)
        return draw

    def pixel(ax, now):
        px, py, ps = 1700, 100, 150
        ax.add_patch(Rectangle((px, py), ps, ps, fc="#13233a", ec=FG, lw=2.5, zorder=3))
        for k, (a, b) in enumerate(pts):
            ax.add_patch(Circle((px + a * ps, py + b * ps), 8, fc=AMB, ec="none", zorder=5))
            txt(ax, px + a * ps + 12, py + b * ps - 8, str(k + 1), size=15, color=DIM)
        txt(ax, px - 20, py + ps / 2, "one pixel,\n8 positions", size=20, color=DIM, ha="right", linespacing=1.4)

    def rule(ax, now):
        box(ax, 96, 880, 1728, 130, ec=AMB)
        txt(ax, 140, 945, r"$\mathrm{history}\leftarrow 0.9\,\mathrm{history}+0.1\,\mathrm{new\ frame}$", size=38)
        txt(ax, 1060, 945, "each frame is shifted by a different sub-pixel offset", size=25, color=DIM)
    return reveal(c, [(0, head("Smooth edges from many frames")), (c.t(1), raster([(0.5, 0.5)], "one sample at the centre: a staircase", 0)),
                      (c.t(2), pixel), (c.t(2) + 3.0, rule), (c.t(3), raster(pts, "eight frames blended: a smooth edge", 880))])


# ----------------------------------------------------------------------------------------------------------- vision
def p_otsu(c):
    rng = np.random.default_rng(3)
    data = np.concatenate([rng.normal(0.24, 0.07, 9000), rng.normal(0.78, 0.06, 700)]).clip(0, 1)
    hist, edges = np.histogram(data, bins=64, range=(0, 1))
    p = hist / hist.sum()
    mids = 0.5 * (edges[:-1] + edges[1:])
    best, bt = -1.0, 0
    for t in range(1, 64):
        w0, w1 = p[:t].sum(), p[t:].sum()
        if w0 <= 0 or w1 <= 0:
            continue
        m0, m1 = (p[:t] * mids[:t]).sum() / w0, (p[t:] * mids[t:]).sum() / w1
        v = w0 * w1 * (m0 - m1) ** 2
        if v > best:
            best, bt = v, t
    x0, y0, w, h = 200, 300, 1100, 480

    def bars(thr=None):
        def draw(ax, now):
            line(ax, [x0, x0 + w], [y0 + h, y0 + h], color=DIM, lw=2)
            hh = np.sqrt(hist / hist.max())
            for k in range(64):
                colr = "#5b7fae" if thr is None else (CYAN if k < thr else AMB)
                ax.add_patch(Rectangle((x0 + k * w / 64, y0 + h - hh[k] * h), w / 64 - 2, hh[k] * h, fc=colr, ec="none", zorder=3))
            txt(ax, x0, y0 + h + 40, "cold", size=22, color=DIM)
            txt(ax, x0 + w, y0 + h + 40, "hot", size=22, color=DIM, ha="right")
            txt(ax, x0 + w / 2, y0 + h + 40, "brightness of a pixel", size=22, color=DIM, ha="center")
            if thr is None:
                txt(ax, x0 + 0.24 * w, y0 - 26, "background", size=26, ha="center")
                txt(ax, x0 + 0.78 * w, y0 + 150, "warm things", size=26, ha="center")
            else:
                X = x0 + thr * w / 64
                line(ax, [X, X], [y0 - 40, y0 + h], color=RED if thr != bt else GREEN, lw=4, z=6)
        return draw

    def note(s, colr):
        def draw(ax, now):
            box(ax, 1380, 330, 450, 300, ec=colr)
            txt(ax, 1410, 480, s, size=25, va="center", linespacing=1.5)
        return draw

    def final(ax, now):
        box(ax, 96, 880, 1728, 130, ec=AMB)
        txt(ax, 140, 945, r"choose $t$ that maximises $\ w_0\,w_1\,(\mu_0-\mu_1)^2$", size=36)
        txt(ax, 1040, 945, "then the detector uses  max(0.80, t)", size=28, color=AMB)
    out = reveal(c, [(0, head("Otsu's method", "the picture chooses its own threshold")), (c.t(1), bars())])
    cuts = [18, 44, 30, bt]
    t0, t1 = c.t(2), c.t(3)
    for k, thr in enumerate(cuts):
        fig, ax = canvas()
        head("Otsu's method", "the picture chooses its own threshold")(ax, False)
        bars(thr)(ax, True)
        last = k == len(cuts) - 1
        note("the best cut:\neach side is as\ntightly packed as it can be" if last else "try a cut:\nhow spread out is\neach side?", GREEN if last else RED)(ax, True)
        if last:
            final(ax, True)
        out.append((t0 + (t1 - t0) * k / (len(cuts) - 1) if not last else t1 - 1.5, c.save(fig)))
    return out


# ---------------------------------------------------------------------------------------------------------- control
def p_tick(c):
    x0, x1, yt, yf = 160, 1760, 430, 640

    def ticks(ax, now):
        line(ax, [x0, x1], [yt, yt], color=FG, lw=3)
        for k in range(11):
            x = x0 + k * (x1 - x0) / 10
            line(ax, [x, x], [yt - 22, yt + 22], color=AMB, lw=4)
        txt(ax, x0, yt - 70, "the simulation: one step every 1/60 s, always", size=28, color=AMB)

    def frames(ax, now):
        line(ax, [x0, x1], [yf, yf], color=FG, lw=3)
        for u in (0.04, 0.17, 0.26, 0.41, 0.49, 0.66, 0.74, 0.93):
            x = x0 + u * (x1 - x0)
            ax.add_patch(Circle((x, yf), 11, fc=CYAN, ec="none", zorder=6))
            k = int(u * 10)
            xa = x0 + k * (x1 - x0) / 10
            line(ax, [xa, x], [yt + 22, yf - 11], color=CYAN, lw=1.5, alpha=0.6)
            line(ax, [xa + (x1 - x0) / 10, x], [yt + 22, yf - 11], color=CYAN, lw=1.5, alpha=0.6)
        txt(ax, x0, yf + 60, "the frames: whenever the display is ready, blended between the last two steps", size=28, color=CYAN)

    def seed(ax, now):
        box(ax, 160, 800, 760, 170, ec=GREEN)
        txt(ax, 200, 850, "same seed + same input", size=30, weight="bold", color=GREEN)
        txt(ax, 200, 915, "the same night, event for event", size=26)
        box(ax, 1000, 800, 760, 170, ec=VIOLET)
        txt(ax, 1040, 850, "F9: record a session", size=30, weight="bold", color=VIOLET)
        txt(ax, 1040, 915, "replay it later, exactly", size=26)
    return reveal(c, [(0, head("A fixed tick under everything")), (c.t(1), ticks), (c.t(2), frames), (c.t(3), seed)])
