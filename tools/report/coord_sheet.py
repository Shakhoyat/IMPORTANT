"""Section 7 coordinate sheets: a spreadsheet-style table of every component of a design after its transformation chain,
in the design's own frame (used by gen_designs.py).

For a component with world matrix W_part under a design root with world matrix W_root, M = W_root^-1 W_part maps the
unit primitive into the design frame. The centre is p = M (0,0,0,1); the extent is the box of the unit primitive
([-1/2, 1/2]^3) under M, extent_i = sum_j |M_ij|, so the corners run from p - extent/2 to p + extent/2.
The centre is checked against the product of the local matrices from the root down to the part."""
import os
import sys

import numpy as np

SHEET_ROWS = 18   # rows per column of a sheet block (two columns side by side); short blocks let pages break between them

HEAD = (r"\rowcolor{sheethead} & \textbf{A} & \textbf{B} & \textbf{C} & \textbf{D} & \textbf{E} & \textbf{F} & \textbf{G}\\\hline" "\n"
        r"\cellcolor{sheethead}\textbf{1} & \textbf{Component} & $x$ & $y$ & $z$ & $\Delta x$ & $\Delta y$ & $\Delta z$\\\hline")
COLS = r"\begin{tabular}[t]{|>{\columncolor{sheethead}}r|>{\raggedright\arraybackslash}p{2.55cm}|r|r|r|r|r|r|}\hline"


def cnum(v):
    """A coordinate in cm: one decimal below 10 m, none above; a dash where no single value exists."""
    if v is None or v != v:
        return "--"
    v = 0.0 if abs(v) < 0.05 else v
    return ("%.0f" % v if abs(v) >= 1000 else "%.1f" % v).replace("-", "$-$")


BOOK = os.path.join(os.path.dirname(__file__), "..", "..", "report", "workbooks", "cg_workbook_values.xlsx")
_book = None


def workbook_rows():
    """{design: [(centre, extent)]} of the workbook's sheet Coordinates (columns K-P), or {} without a workbook."""
    global _book
    if _book is None:
        _book = {}
        try:
            from openpyxl import load_workbook
            ws = load_workbook(BOOK, data_only=True, read_only=True)["Coordinates"]
        except (ImportError, OSError, KeyError):
            return _book
        for r in ws.iter_rows(min_row=6, values_only=True):
            if r[0] and isinstance(r[1], (int, float)) and r[10] is not None:
                nan = lambda v: float("nan") if v is None or v == "" else float(v)
                _book.setdefault(r[0], []).append((np.array([nan(v) for v in r[10:13]]), np.array([nan(v) for v in r[13:16]])))
    return _book


def compute_rows(design, own, root_id, by_id, mat, terrain_radius=None):
    """[(part kind, centre (cm), extent (cm))] for every component of the first instance, from the export.
    A design whose root is itself the part (a shell casing, the moon, the far ground) is given in the root's frame
    without the root's size, so the extent is the part's size in use. The terrain disc is a height field of radius R
    (not a unit primitive): its extent is 2R in x and z, and no single height is given."""
    inv = np.linalg.inv(mat(by_id[root_id]["world_matrix"]))
    rows = []
    for c in own:
        if c["id"] == root_id:
            M = np.diag(list(c["S_rest"]) + [1.0])
        else:
            M = inv @ mat(c["world_matrix"])
            L = np.eye(4)
            node = c
            while node and node["id"] != root_id:
                L = mat(node["local_matrix"]) @ L
                node = by_id.get(node["parent_id"])
            if node is None or np.abs(L[:3, 3] - M[:3, 3]).max() > 1e-3 * max(1.0, np.abs(M[:3, 3]).max()):
                sys.exit("%s / %s: centre from the chain differs from the exported world matrix" % (design["name"], c["name"]))
        p, e = M[:3, 3] * 100.0, np.abs(M[:3, :3]).sum(axis=1) * 100.0
        if c["primitive"] == "Terrain disc" and terrain_radius:
            e = np.array([e[0] * 2 * terrain_radius, float("nan"), e[2] * 2 * terrain_radius])
        rows.append((c["part_kind"], p, e))
    return rows


def sheet_rows(design, own, root_id, by_id, mat, terrain_radius=None):
    """The rows printed: the cells of the workbook's sheet Coordinates; the export is used to check them (and alone,
    when there is no workbook)."""
    rows = compute_rows(design, own, root_id, by_id, mat, terrain_radius)
    book = workbook_rows().get(design["name"])
    if book is None:
        return rows
    if len(book) != len(rows):
        sys.exit("%s: the workbook lists %d components, the export %d (rebuild the workbook)" % (design["name"], len(book), len(rows)))
    out = []
    for (kind, p, e), (bp, be) in zip(rows, book):
        if not (np.allclose(bp, p, atol=1e-6, equal_nan=True) and np.allclose(be, e, atol=1e-6, equal_nan=True)):
            sys.exit("%s / %s: the workbook's coordinates differ from the export (rebuild the workbook)" % (design["name"], kind))
        out.append((kind, bp, be))
    return out


def column(chunk, first, tex):
    o = [COLS, HEAD]
    for k, (kind, p, e) in enumerate(chunk):
        nm = tex(kind if len(kind) <= 24 else kind[:23] + ".")
        o.append(r"%d & %s & %s & %s & %s & %s & %s & %s\\\hline"
                 % (first + k, nm, cnum(p[0]), cnum(p[1]), cnum(p[2]), cnum(e[0]), cnum(e[1]), cnum(e[2])))
    o.append(r"\end{tabular}")
    return "\n".join(o)


def coord_sheet(design, own, root_id, by_id, mat, tex, shown_name, terrain_radius=None):
    rows = sheet_rows(design, own, root_id, by_id, mat, terrain_radius)
    o = [r"\par\smallskip\noindent{\small\textbf{Coordinate sheet, %s} (cm, design frame): centre B--D, extent E--G; %d component%s.}"
         r"\par\nopagebreak\vspace{0.6mm}" % (tex(shown_name), len(rows), "" if len(rows) == 1 else "s")]
    per = SHEET_ROWS
    blocks = -(-len(rows) // (2 * per))
    if blocks > 1 and len(rows) - 2 * per * (blocks - 1) <= 6:   # a few rows left over: share them among the blocks
        blocks -= 1
        per = -(-len(rows) // (2 * blocks))
    elif blocks == 1 and len(rows) > 2 * per:
        per = -(-len(rows) // 2)
    for b in range(0, len(rows), 2 * per):
        left = rows[b:b + per]
        right = rows[b + per:b + 2 * per]
        if len(rows) - b <= 2 * per:   # last block: balance the two columns
            rest = rows[b:]
            h = (len(rest) + 1) // 2 if len(rest) > 8 else len(rest)
            left, right = rest[:h], rest[h:]
        blk = (r"\noindent{\fontsize{6.6}{7.6}\selectfont\setlength{\tabcolsep}{2.2pt}\renewcommand{\arraystretch}{0.98}"
               + column(left, b + 2, tex))
        if right:
            blk += r"\hfill " + column(right, b + 2 + len(left), tex)
        if b == 0:   # the title is set in one box with the first (short) block, so it cannot be left behind on a page
            blk = r"\noindent\begin{minipage}{\linewidth}" + o.pop() + "\n" + blk + r"}\end{minipage}\par\vspace{1.5mm}"
        else:
            blk += r"}\par\vspace{1.5mm}"
        o.append(blk)
    return "\n".join(o), len(rows)
