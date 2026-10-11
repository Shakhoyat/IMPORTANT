"""The report's number macros, written from the workbook.

Sheet Report_Numbers of report/workbooks/cg_workbook_values.xlsx (the recalculated copy) lists every number macro of
the report: file, name, value, the text printed. This tool puts the workbook's text into the macro files of
report/sections (gen_game.tex, gen_numbers.tex, gen_tables.tex, gen_designs.tex, gen_m3.tex), so the number a reader
finds in the report is the one in the workbook.

    python tools/report/numbers_from_workbook.py            write the macros from the workbook
    python tools/report/numbers_from_workbook.py --check    change nothing; exit code 1 if a macro differs

Run with: ~/.local/bin/uv run -q --with openpyxl python tools/report/numbers_from_workbook.py"""
import os
import re
import sys

from openpyxl import load_workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SEC = os.path.join(ROOT, "report", "sections")
BOOK = os.path.join(ROOT, "report", "workbooks", "cg_workbook_values.xlsx")


def register():
    """[(file, name, value, printed text)] of section A of the sheet."""
    ws = load_workbook(BOOK, data_only=True)["Report_Numbers"]
    rows, on = [], False
    for r in ws.iter_rows(values_only=True):
        if r[0] == "file":
            on = True
            continue
        if on and not r[0]:
            break
        if on:
            rows.append((r[0], r[1], r[2], "" if r[3] is None else str(r[3])))
    return rows


def pattern(fname, name):
    """Regular expression with the macro's body as group 2."""
    if fname == "gen_game.tex":
        return re.compile(r"(\\expandafter\\def\\csname gv@%s\\endcsname\{)(.*)(\}[ \t]*)$" % re.escape(name), re.M)
    return re.compile(r"(\\newcommand\{\\%s\}\{)((?:[^\\{}]|\{,\})*)(\})" % re.escape(name))


def plain(text):
    t = text.replace("{,}", "").replace("\\,", "").replace(",", "")
    try:
        return float(t)
    except ValueError:
        return None


def main():
    check = "--check" in sys.argv
    rows = register()
    texts, changed, bad = {}, 0, []
    for fname, name, value, printed in rows:
        path = os.path.join(SEC, fname)
        if fname not in texts:
            texts[fname] = open(path).read()
        rx = pattern(fname, name)
        m = rx.search(texts[fname])
        if not m:
            bad.append("%s: macro %s is in the workbook but not in the report" % (fname, name))
            continue
        if value is not None and plain(printed) is not None and abs(plain(printed) - float(value)) > 1e-9:
            bad.append("%s: %s prints %s but the workbook value is %s" % (fname, name, printed, value))
        if m.group(2) != printed:
            if check:
                bad.append("%s: %s is %s in the report, %s in the workbook" % (fname, name, m.group(2), printed))
            else:
                texts[fname] = texts[fname][:m.start(2)] + printed + texts[fname][m.end(2):]
                changed += 1
    if not check:
        for fname, t in texts.items():
            if t != open(os.path.join(SEC, fname)).read():
                open(os.path.join(SEC, fname), "w").write(t)
    print("%d number macros in the workbook; %s" % (len(rows), "%d differ" % len([b for b in bad if " in the report, " in b])
                                                     if check else "%d written to the report" % changed))
    for b in bad:
        print("  " + b)
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
