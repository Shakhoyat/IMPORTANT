"""The Excel workbook (report/workbooks/cg_workbook_values.xlsx, the recalculated copy) against the scene export, the
code and the report's macros. Exit code 1 on any mismatch.
Run: ~/.local/bin/uv run -q --with openpyxl python tools/report/check_workbook.py"""
import json
import os
import re
import sys

from openpyxl import load_workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
wb = load_workbook(os.path.join(ROOT, "report", "workbooks", "cg_workbook_values.xlsx"), data_only=True)
exp = json.load(open(os.path.join(ROOT, "report", "data", "scene_export_v10.json")))
bad, ok = [], 0


def note(cond, what):
    global ok
    if cond:
        ok += 1
    else:
        bad.append(what)


def rows(sheet):
    return [list(r) for r in wb[sheet].iter_rows(values_only=True)]


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


# ---- every formula cell has a value and no error
errors = 0
for ws in wb.worksheets:
    for r in ws.iter_rows(values_only=True):
        errors += sum(1 for v in r if v in ("#DIV/0!", "#N/A", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#NULL!"))
note(errors == 0, "%d cells hold an Excel error value" % errors)

# ---- Designs: instances, parts per instance and totals of every design; the scene total
des = rows("Designs")
head = next(i for i, r in enumerate(des) if r and r[0] == "Design")
by = {d["name"]: d for d in exp["designs"]}
seen = 0
end = next(i for i, r in enumerate(des) if r and r[0] == "Scene")
for r in des[head + 1:end + 1]:
    if r[0] in by:
        d = by[r[0]]
        seen += 1
        note(num(r[2]) == d["instances"], "Designs: %s has %s instances, export %d" % (r[0], r[2], d["instances"]))
        note(num(r[4]) == d["parts_total"], "Designs: %s has %s parts, export %d" % (r[0], r[4], d["parts_total"]))
    elif r[0] == "Scene":
        note(num(r[2]) == exp["leaf_count"] or num(r[3]) == exp["leaf_count"] or exp["leaf_count"] in [num(v) for v in r],
             "Designs: scene total differs from the export's %d parts" % exp["leaf_count"])
note(seen == len(exp["designs"]), "Designs: %d of %d designs listed" % (seen, len(exp["designs"])))

# ---- Decomposition: the per-primitive scene totals add up to the scene
dec = rows("Decomposition")
tot = next((r for r in dec if r and r[0] == "Scene total"), None)
note(tot is not None and abs(sum(v for v in (num(x) for x in tot[1:]) if v is not None) - exp["leaf_count"]) < 0.5
     or (tot is not None and exp["leaf_count"] in [num(x) for x in tot[1:]]),
     "Decomposition: the scene total row does not add up to %d" % exp["leaf_count"])
chk = next((r for r in dec if r and r[0] == "check"), None)
note(chk is not None and all(v in (None, "OK") for v in chk[1:]), "Decomposition: a check cell is not OK")

# ---- Primitives: vertex counts of the generators equal the exported meshes
pr = rows("Primitives")
h = next(i for i, r in enumerate(pr) if r and r[0] == "Primitive")
cv, ce = pr[h].index("V"), pr[h].index("V export")
n = 0
for r in pr[h + 1:]:
    if r[0] and num(r[cv]) is not None and num(r[ce]) is not None:
        n += 1
        note(num(r[cv]) == num(r[ce]), "Primitives: %s formula gives %s vertices, the mesh has %s" % (r[0], r[cv], r[ce]))
note(n >= 14, "Primitives: only %d rows with a vertex count" % n)

# ---- constants copied from the code
def const(path, name):
    m = re.search(r"\b%s\s*=\s*([-\d.]+)f?\b" % re.escape(name), open(os.path.join(ROOT, path)).read())
    return float(m.group(1)) if m else None


lf = rows("Line_Follower")
hl = next(i for i, r in enumerate(lf) if r and "Kp" in r)
vals = dict(zip(lf[hl], lf[hl + 1]))
note(num(vals.get("Kp")) == const("src/ai/HideController.cpp", "kKp"), "Line_Follower: Kp %s, code %s" % (vals.get("Kp"), const("src/ai/HideController.cpp", "kKp")))
note(num(vals.get("Kd")) == const("src/ai/HideController.cpp", "kKd"), "Line_Follower: Kd %s, code %s" % (vals.get("Kd"), const("src/ai/HideController.cpp", "kKd")))

# ---- the report's macro for the scene size
nums = open(os.path.join(ROOT, "report", "sections", "gen_numbers.tex")).read()
m = re.search(r"\\newcommand\{\\nLeaves\}\{(.*)\}\s*$", nums, re.M)
note(m and int(re.sub(r"[^\d]", "", m.group(1))) == exp["leaf_count"], "report macro nLeaves differs from the export")

# ---- Report_Numbers, A: every number macro of the report is the workbook's cell
sys.path.insert(0, os.path.dirname(__file__))
import numbers_from_workbook as nfw  # noqa: E402

SEC = os.path.join(ROOT, "report", "sections")
reg = nfw.register()
texts = {}
for fname, name, value, printed in reg:
    t = texts.setdefault(fname, open(os.path.join(SEC, fname)).read())
    m = nfw.pattern(fname, name).search(t)
    note(m is not None and m.group(2) == printed, "Report_Numbers: %s %s is %s in the report, %s in the workbook"
         % (fname, name, m.group(2) if m else "missing", printed))
    if value is not None and nfw.plain(printed) is not None:
        note(abs(nfw.plain(printed) - float(value)) < 1e-9, "Report_Numbers: %s value %s, printed %s" % (name, value, printed))
# ... and no number macro of the report is missing from the workbook
listed = {(f, n) for f, n, _, _ in reg}
for fname, rx in (("gen_game.tex", r"\\expandafter\\def\\csname gv@([^\\]+?)\\endcsname\{(.*)\}\s*$"),
                  ("gen_numbers.tex", r"\\newcommand\{\\(\w+)\}\{((?:[^\\{}]|\{,\})*)\}"),
                  ("gen_tables.tex", r"\\newcommand\{\\(\w+)\}\{((?:[^\\{}]|\{,\})*)\}"),
                  ("gen_designs.tex", r"\\newcommand\{\\(\w+)\}\{((?:[^\\{}]|\{,\})*)\}"),
                  ("gen_m3.tex", r"\\newcommand\{\\(\w+)\}\{((?:[^\\{}]|\{,\})*)\}")):
    if not os.path.exists(os.path.join(SEC, fname)):
        continue
    for m in re.finditer(rx, open(os.path.join(SEC, fname)).read(), re.M):
        if re.search(r"\d", m.group(2)) and len(m.group(2)) <= 48:
            note((fname, m.group(1)) in listed, "Report_Numbers: macro %s of %s is not in the workbook" % (m.group(1), fname))

# ---- Report_Numbers, B-D: the stated numbers against the code, read again now
rn = rows("Report_Numbers")
stated = {r[1]: num(r[2]) for r in rn if r and r[0] in ("text", "constant", "derived")}


def grab(path, rx, group=1):
    m = re.search(rx, open(os.path.join(ROOT, path)).read(), re.S)
    return float(m.group(group)) if m else None


for name, path, rx in (
        ("split distance (m)", "src/ai/Squad.h", r"kSquadSplitDist\s*=\s*([\d.]+)"),
        ("re-join distance (m)", "src/ai/Squad.h", r"kSquadRejoinDist\s*=\s*([\d.]+)"),
        ("split urgency (m to a protected point)", "src/ai/Squad.h", r"kSquadSplitThreat\s*=\s*([\d.]+)"),
        ("wait for a free wing (s)", "src/ai/Squad.h", r"kSquadWingWaitFree\s*=\s*([\d.]+)"),
        ("guards' rifle range (m)", "src/scene/Scene.h", r"kRifleRange\s*=\s*([\d.]+)"),
        ("helicopter alert idle (rotor fraction)", "src/objects/Helicopter.h", r"kAlertIdle\s*=\s*([\d.]+)"),
        ("air mass (Kasten-Young): coefficient", "src/scene/DayCycle.cpp", r"\+ ([\d.]+)f \* std::pow\(e \+"),
        ("fog density at night (1/m)", "src/rendering/Lighting.h", r"float fogDensity = ([\d.]+)f"),
        ("fog density by day (1/m)", "src/scene/DayCycle.cpp", r"kDayFogDensity = ([\d.]+)f"),
        ("near plane N (m)", "src/scene/Camera.h", r"float nearZ = ([\d.]+)f"),
        ("far plane F (m)", "src/scene/Camera.h", r"farZ = ([\d.]+)f"),
        ("tone curve, denominator: quadratic term", "src/math/ToneMap.cpp", r"const float b = x \* \(([\d.]+)f \* x")):
    note(name in stated and stated[name] == grab(path, rx), "Report_Numbers: %s is %s, %s has %s" % (name, stated.get(name), path, grab(path, rx)))
note(stated.get("scene parts") == exp["leaf_count"], "Report_Numbers: scene parts differ from the export")
note(stated.get("designs") == len(exp["designs"]), "Report_Numbers: designs differ from the export")
note(abs((stated.get("far / near ratio F/N") or 0) - 12000) < 1e-6, "Report_Numbers: F/N is %s" % stated.get("far / near ratio F/N"))

# ---- Coordinates: every row of every coordinate sheet of the report's Section 7 is a row of the workbook
import coord_sheet as cs  # noqa: E402

book = {}
for r in rows("Coordinates")[5:]:
    if r[0] and num(r[1]) is not None and r[10] is not None:
        book.setdefault(r[0], []).append("%d|" % r[1] + "|".join(cs.cnum(None if v in (None, "") else float(v)) for v in r[10:16]))
des_tex = re.sub(r"\\ab[pd]?\{\d+\}\{([^{}]*)\}", r"\1", open(os.path.join(SEC, "gen_designs.tex")).read())
sheets = re.split(r"\\textbf\{Coordinate sheet, [^}]*\}", des_tex)[1:]
pool = {d: "\n".join(v) for d, v in book.items()}
used = set()
n_rows = 0
for k, block in enumerate(sheets):
    want = int(re.search(r"(\d+) components?\.", block).group(1))
    got = []
    for m in re.finditer(r"^(\d+) & [^&]* & ([^&]+) & ([^&]+) & ([^&]+) & ([^&]+) & ([^&]+) & ([^&\\]+)\\\\\\hline", block, re.M):
        if len(got) < want:
            got.append("|".join(x.strip() for x in m.groups()))
    n_rows += len(got)
    key = "\n".join(got)
    match = [d for d, v in pool.items() if v == key and d not in used]
    note(len(got) == want and bool(match), "Coordinates: sheet %d of Section 7 (%d components) has no equal design in the workbook" % (k + 1, want))
    if match:
        used.add(match[0])
note(len(sheets) == len(book) == len(exp["designs"]), "Coordinates: %d sheets in the report, %d designs in the workbook" % (len(sheets), len(book)))
note(n_rows == sum(len(v) for v in book.values()), "Coordinates: %d rows in the report, %d in the workbook" % (n_rows, sum(len(v) for v in book.values())))

# ---- every number typed by hand in the report's text is a value of the workbook
import bisect  # noqa: E402
import glob  # noqa: E402

cells = sorted({float(v) for ws in wb.worksheets for r in ws.iter_rows(values_only=True) for v in r
                if isinstance(v, (int, float)) and not isinstance(v, bool)})


def in_book(x, dec):
    tol = 0.5 * 10 ** (-dec) + 1e-9
    i = bisect.bisect_left(cells, x - tol)
    return i < len(cells) and cells[i] <= x + tol


NOT_DATA = {"4102", "2107104", "4070"}     # course code, roll number, the graphics card's model number
typed = missing = 0
for f in sorted(glob.glob(os.path.join(SEC, "*.tex"))):
    if os.path.basename(f).startswith(("gen_", "sheet_")):
        continue
    t = re.sub(r"(?<!\\)%.*", "", open(f).read())
    t = re.sub(r"\\begin\{(lstlisting|verbatim)\}.*?\\end\{\1\}", "", t, flags=re.S)
    t = re.sub(r"\\begin\{tikzpicture\}.*?\\end\{tikzpicture\}",          # a diagram: keep the node texts only
               lambda m: " ".join(re.findall(r"\{([^{}]*)\}", re.sub(r"\bat\s*\((?:[^()]|\{[^{}]*\})*\)", "", m.group(0)))), t, flags=re.S)
    t = re.sub(r"\\(label|ref|cref|Cref|eqref|cite|input|includegraphics|pageref|autoref|code|texttt|url|path|lstinline|ab|abp|gv"
               r"|vspace|hspace|rule|setlength|fontsize|cellcolor|rowcolor|color|textcolor|begin|end|multicolumn|multirow|cline"
               r"|resizebox|scalebox|parbox|needspace|addvspace|raisebox|captionsetup)(\[[^\]]*\])*(\{[^{}]*\})*", "", t)
    t = re.sub(r"[0-9.]+\\?(cm|mm|pt|em|ex|linewidth|textwidth|columnwidth|baselineskip)\b", "", t)
    t = t.replace("{,}", "").replace("\\,", "")
    for m in re.finditer(r"(?<![A-Za-z_\d.\\{^])(\d+(?:\.\d+)?)(?![\d]*[A-Za-z_}])", t):
        s_, dec = m.group(1), len(m.group(1).split(".")[1]) if "." in m.group(1) else 0
        if (float(s_) <= 12 and dec == 0) or s_ in NOT_DATA:
            continue
        typed += 1
        if not in_book(float(s_), dec):
            missing += 1
            bad.append("text: %s in %s is not in the workbook (%s)" % (s_, os.path.basename(f), t[max(0, m.start() - 40):m.end() + 20].replace("\n", " ")))
ok += typed - missing
print("report text: %d typed numbers (above 12 or with decimals), %d not in the workbook" % (typed, missing))

# ---- register sheets: World_Coordinates, Layout, Algorithms, Film against the export, the source files and the film's data
wc = [r for r in rows("World_Coordinates")[8:] if r[0]]
note(len(wc) == len(exp["components"]), "World_Coordinates: %d rows, the export has %d components" % (len(wc), len(exp["components"])))
note(sum(1 for r in wc if r[4] != "Group") == exp["leaf_count"], "World_Coordinates: drawn parts differ from the export's leaf_count")
note(rows("World_Coordinates")[4][6] == "agree", "World_Coordinates: the sheet's own count check says DIFFER")
by = {c["id"]: c for c in exp["components"]}
worst = 0.0
for r in wc:
    c = by.get(r[0])
    if not c:
        bad.append("World_Coordinates: unknown id %s" % r[0])
        continue
    m = c["world_matrix"]
    want = list(c["S"]) + list(c["R_deg"]) + list(c["T"]) + [m[3], m[7], m[11]]
    worst = max(worst, max(abs(float(a) - float(b)) for a, b in zip(r[7:19], want)))
note(worst < 6e-4, "World_Coordinates: a coordinate differs from the export by %.4f" % worst)
lay = rows("Layout")
ctl = [r for r in lay if r[0] and any(r[0] == c["name"] for c in exp["controllables"])]
note(len(ctl) >= len(exp["controllables"]), "Layout: %d of %d controllables listed" % (len(ctl), len(exp["controllables"])))
for c in exp["controllables"]:
    hit = [r for r in ctl if r[0] == c["name"] and num(r[2]) is not None]
    note(bool(hit) and all(abs(num(hit[0][2 + j]) - c["position"][j]) < 6e-3 for j in range(3)), "Layout: position of %s" % c["name"])
lamps = [r for r in lay if r[0] and any(r[0] == pl["name"] for pl in exp["practical_lights"]) and num(r[8]) is not None]
note(len(lamps) == len(exp["practical_lights"]), "Layout: %d lamps listed, the export has %d" % (len(lamps), len(exp["practical_lights"])))
alg = [r for r in rows("Algorithms") if r[0] and str(r[0]).startswith("src/") and isinstance(r[1], (int, float))]
note(len(alg) > 80, "Algorithms: only %d source constants listed" % len(alg))
for f, line, name, val, _ in alg:
    src_line = open(os.path.join(ROOT, f), errors="replace").read().splitlines()[int(line) - 1]
    m = re.search(r"\b%s\s*=\s*(-?[\d.]+)" % re.escape(str(name)), src_line)
    note(bool(m) and abs(float(m.group(1)) - float(val)) < 1e-6, "Algorithms: %s in %s line %s is not %s" % (name, f, line, val))
film_path = os.path.join(ROOT, "report", "data", "film.json")
if os.path.exists(film_path):
    film = json.load(open(film_path))
    fr = [r for r in rows("Film")[9:] if isinstance(r[0], (int, float)) and r[3]]
    note(len(fr) == len(film["beats"]), "Film: %d beats in the sheet, %d in film.json" % (len(fr), len(film["beats"])))
    note(abs(sum(float(r[6]) for r in fr) - film["length"]) < 0.5, "Film: the lengths do not add up to the film's length")
    said = {r[2]: r[3] for r in rows("Film") if r[2] and r[5] and str(r[5]).startswith("scene export")}
    note(said.get("drawn parts") == exp["leaf_count"] and said.get("designs") == len(exp["designs"]), "Film: parts / designs stated in the film differ from the export")

print("workbook: %d sheets, %d checks passed, %d failed" % (len(wb.sheetnames), ok, len(bad)))
for b in bad:
    print("  " + b)
sys.exit(1 if bad else 0)
