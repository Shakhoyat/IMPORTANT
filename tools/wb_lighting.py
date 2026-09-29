"""Sheets 7 (Lighting) and 8 (Shading)."""
from openpyxl.chart import Reference, ScatterChart, Series

from wb_common import (BOLD, FMT_3, INPUT_FILL, S_LIGHT, S_SHADE,
                       TOTAL_FILL, finish, style_chart, header, image, path, put, section, title, widths)

CURVE_SAMPLES = 40  # chart resolution (number of distance / angle steps), not a measured value


# ---- 7 Lighting -----------------------------------------------------------------------------------------
def _light_rows(L):
    moon, beacon, spot = L["moon"], L["beacon"], L["spot"]
    blank3 = ["-", "-", "-"]
    return [
        ["L0", "Directional (moonlight)", *blank3, *moon["direction"], *moon["color"], moon["intensity"], "-", "-",
         "-", "-", "-", "Cool key light for every surface; direction also drives the cascaded shadow map"],
        ["L1", "Point (bunker beacon)", *beacon["position"], *blank3, *beacon["color"],
         beacon.get("intensity", beacon.get("intensity_min", "-")), beacon.get("intensity_max", "-"),
         beacon["linear"], beacon["quadratic"], "-", "-", "Warm pulsing local light; shows distance attenuation"],
        ["L2", "Spot (UAV searchlight)", *spot["position"], *spot["direction"], *spot["color"], spot["intensity"],
         "-", spot["linear"], spot["quadratic"], spot["inner_deg"], spot["outer_deg"],
         "Hierarchical light: position/direction = gimbal pitch node at t = 0 (moves with the UAV)"],
        ["AMB", "Ambient term", *blank3, *blank3, *L["ambient"], "-", "-", "-", "-", "-", "-",
         "Keeps unlit sides from pure black"],
    ]


def build_lighting(wb, ctx):
    ws = wb[S_LIGHT]
    scene = ctx["scene"]
    title(ws, "7  Lighting",
          "The three simultaneous lights of the proposal, exported from the program's LightRig at t = 0. The "
          "attenuation curves use the shader's own formulas: point/spot attenuation 1 / (1 + kl d + kq d^2); spot "
          "cone factor c = clamp((cos theta - cos outer) / (cos inner - cos outer), 0, 1), applied as c^2.", span=16)
    hdr = ["ID", "Type", "Pos x", "Pos y", "Pos z", "Dir x", "Dir y", "Dir z", "Colour r", "Colour g", "Colour b",
           "Intensity", "Intensity max (pulse)", "k linear", "k quadratic", "Inner cone (deg)", "Outer cone (deg)",
           "Role"]
    header(ws, 4, hdr)
    for i, row in enumerate(_light_rows(scene["lights"])):
        for k, v in enumerate(row):
            numeric = isinstance(v, (int, float))
            put(ws, 5 + i, 1 + k, v, FMT_3 if numeric else None, INPUT_FILL if numeric else None,
                BOLD if k == 0 else None)
    # parameters for the curves (row 10..12)
    section(ws, 10, 1, "Curve settings")
    put(ws, 11, 1, "Max distance d (m) = terrain radius", font=BOLD)
    ws.merge_cells("A11:C11")
    put(ws, 11, 4, scene.get("terrain", {}).get("radius", 10), FMT_3, INPUT_FILL)
    put(ws, 12, 1, "Samples (chart resolution)", font=BOLD)
    ws.merge_cells("A12:C12")
    put(ws, 12, 4, CURVE_SAMPLES)
    top = 15
    section(ws, top - 1, 1, "Distance attenuation (formula)")
    header(ws, top, ["d (m)", "Beacon 1/(1+kl d+kq d^2)", "Spot 1/(1+kl d+kq d^2)", "Beacon radiance at max "
                     "pulse", "Spot radiance x intensity"])
    for i in range(CURVE_SAMPLES + 1):
        r = top + 1 + i
        put(ws, r, 1, 0 if i == 0 else f"=A{r - 1}+$D$11/$D$12", FMT_3)
        put(ws, r, 2, f"=1/(1+$N$6*A{r}+$O$6*A{r}^2)", "0.0000")
        put(ws, r, 3, f"=1/(1+$N$7*A{r}+$O$7*A{r}^2)", "0.0000")
        put(ws, r, 4, f"=IF(ISNUMBER($M$6),$M$6,$L$6)*B{r}", "0.0000")
        put(ws, r, 5, f"=$L$7*C{r}", "0.0000")
    last = top + 1 + CURVE_SAMPLES
    section(ws, top - 1, 8, "Spotlight cone fall-off (formula)")
    header(ws, top, ["theta (deg)", "cos theta", "cone c", "intensity factor c^2"], start_col=8)
    for i in range(CURVE_SAMPLES + 1):
        r = top + 1 + i
        put(ws, r, 8, 0 if i == 0 else f"=H{r - 1}+(2*$Q$7-$P$7)/$D$12", "0.00")
        put(ws, r, 9, f"=COS(RADIANS(H{r}))", "0.00000")
        put(ws, r, 10, f"=MAX(0,MIN(1,(I{r}-COS(RADIANS($Q$7)))/(COS(RADIANS($P$7))-COS(RADIANS($Q$7)))))", "0.000")
        put(ws, r, 11, f"=J{r}^2", "0.000")
    _scatter(ws, "Attenuation vs distance", "d (m)", "factor", 1, [2, 3], top, last, "M15")
    _scatter(ws, "Spot cone fall-off vs angle from axis", "theta (deg)", "factor", 8, [10, 11], top, last, "M33")
    widths(ws, [9, 22, 12, 12, 12, 8, 8, 10, 9, 9, 11, 9, 11, 9, 10, 9, 9, 60])
    finish(ws, 4)


def _scatter(ws, name, xt, yt, xcol, ycols, top, last, anchor):
    ch = ScatterChart()
    ch.title, ch.style = name, 13
    ch.x_axis.title, ch.y_axis.title = xt, yt
    xs = Reference(ws, min_col=xcol, min_row=top + 1, max_row=last)
    for yc in ycols:
        s = Series(Reference(ws, min_col=yc, min_row=top, max_row=last), xs, title_from_data=True)
        s.marker.symbol = "none"
        s.smooth = True
        ch.series.append(s)
    ch.height, ch.width = 8.5, 17
    ws.add_chart(style_chart(ch), anchor)


# ---- 8 Shading ------------------------------------------------------------------------------------------
SHADING_TEXT = {
    "Flat": ("One face normal per triangle (geometry shader: cross product of two edges)",
             "Once per triangle", "Lowest", "Faceted look; every triangle has one constant colour",
             "Blocky concrete architecture of the command bunker (proposal: flat shading)"),
    "Gouraud": ("Vertex normal", "Per vertex, colour interpolated across the triangle", "Low",
                "Smooth, matte; highlights and the spotlight edge are blurred between vertices",
                "Large terrain and background props: cheap and smooth"),
    "Phong": ("Vertex normal interpolated per fragment", "Per fragment (Blinn-Phong)", "Medium",
              "Sharp moving specular highlights; crisp spotlight edge on small parts",
              "UAV and rover: small metallic parts that the searchlight sweeps over"),
    "Additive beam": ("None (unlit)", "Per fragment, additive blending, no depth writes", "Low",
                      "Visible searchlight cone that fades along its axis and at its silhouette",
                      "Makes the hierarchical searchlight readable at night"),
    "PBR (GGX / Cook-Torrance)": ("Interpolated normal + material roughness / metalness",
                                  "Per fragment: D (GGX) x G (Smith-Schlick) x F (Schlick) / (4 N.V N.L)", "Highest",
                                  "Energy-conserving highlights whose size follows roughness; metals tint reflections",
                                  "Optional lighting path of the fragment shader (uniform uPbr); uses the "
                                  "roughness / metalness columns of sheet 3"),
}
FIGURES = [("Flat (bunker)", "report/figures/bunker_flat_visible.png"),
           ("Gouraud override (bunker)", "report/figures/bunker_all_gouraud.png"),
           ("Phong override (bunker)", "report/figures/bunker_all_phong.png")]
BUNKER_CROP = (470, 110, 1170, 830)


def build_shading(wb, ctx):
    ws = wb[S_SHADE]
    comps = ctx["scene"]["components"]
    title(ws, "8  Shading models",
          "Which shading model each part uses (counted with COUNTIF / COUNTIFS over sheet 3) and what it does. The "
          "screenshots are crops of the program's own captures in report/figures (same camera, shading overridden).",
          span=9)
    hdr = ["Shading model", "Parts using it", "Objects using it (parts)", "Normal used", "Lighting evaluated",
           "Cost", "Visual effect", "Why used here"]
    header(ws, 4, hdr)
    models = list(dict.fromkeys(c["shading"] for c in comps if c["shading"] != "-"))
    models += [m for m in SHADING_TEXT if m.startswith("PBR") and m not in models]
    objects = list(dict.fromkeys(c["object"] for c in comps))
    rng = ctx["obj"]["range"]
    for i, m in enumerate(models):
        r = 5 + i
        put(ws, r, 1, m, font=BOLD)
        if m.startswith("PBR"):
            put(ws, r, 2, f"=COUNTIF({rng('S')},\">0\")", "0")
            put(ws, r, 3, "Parts with metalness > 0 are shown in column B; all lit parts can use the PBR path")
        else:
            put(ws, r, 2, f"=COUNTIF({rng('P')},A{r})", "0")
            parts = [f"\"{o} (\"&COUNTIFS({rng('B')},\"{o}\",{rng('P')},A{r})&\")\"" for o in objects
                     if any(c["object"] == o and c["shading"] == m for c in comps)]
            put(ws, r, 3, "=" + "&\", \"&".join(parts) if parts else "-")
        text = SHADING_TEXT.get(m, ("-", "-", "-", "-", "-"))
        for k, v in enumerate(text):
            put(ws, r, 4 + k, v)
        ws.row_dimensions[r].height = 60
    last = 4 + len(models)
    put(ws, last + 1, 1, "Total lit parts", font=BOLD, fill=TOTAL_FILL)
    put(ws, last + 1, 2, f"=SUM(B5:B{last - 1 if models[-1].startswith('PBR') else last})", "0", TOTAL_FILL, BOLD)
    top = last + 4
    section(ws, top - 1, 1, "Same bunker, same camera: flat vs Gouraud vs Phong (crops of the program's captures)")
    anchors = ["A", "D", "G"]
    for (cap, rel), a in zip(FIGURES, anchors):
        ws[f"{a}{top}"] = cap
        ws[f"{a}{top}"].font = BOLD
        image(ws, f"{a}{top + 1}", path(rel), box=BUNKER_CROP, width=300)
    widths(ws, [24, 10, 34, 30, 30, 9, 40, 44])
    finish(ws, 4)
