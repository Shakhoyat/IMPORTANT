"""Sheets 4 (Transformation Matrices M1, M2, ... with M = T * Ry * Rx * Rz * S) and 5 (Hierarchy / world matrices).

All matrix elements are element-wise Excel formulas (no MMULT / array formulas), so they recalculate in any
Excel or LibreOffice version. The code-exported matrices sit next to them with a max-abs-error column.
"""
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import wb_common  # noqa: E402
from wb_common import (BOLD, FAIL_FILL, FMT_3, FMT_SCI, HILITE_FILL, INPUT_FILL, PASS_FILL, S_HIER, S_MAT,
                       TOTAL_FILL, articulation_chains, matrix_components, col, define_name, finish, nowrap, header, prim_fill, put, section, title, widths)

JOINT_FILL = PatternFill("solid", fgColor="DBEAFE")
TOLERANCE = 1e-4  # acceptance limit for |Excel - code|: float32 round-off of the program, not a measurement
C_S, C_R, C_T = 7, 10, 13          # first column of S, R(deg), T on sheet 4
C_ROT, C_M, C_CODE = 16, 25, 37    # first column of R (3x3), M (3x4 formula), M (3x4 code)
C_ERR = 49
C_EX = 51                          # worked example block
M_NAMES = [f"M{i + 1}{j + 1}" for i in range(3) for j in range(4)]


def _rot_formulas(r):
    cx, sx = f"COS(RADIANS({col(C_R)}{r}))", f"SIN(RADIANS({col(C_R)}{r}))"
    cy, sy = f"COS(RADIANS({col(C_R + 1)}{r}))", f"SIN(RADIANS({col(C_R + 1)}{r}))"
    cz, sz = f"COS(RADIANS({col(C_R + 2)}{r}))", f"SIN(RADIANS({col(C_R + 2)}{r}))"
    # R = Ry * Rx * Rz expanded element by element.
    return [f"={cy}*{cz}+{sy}*{sx}*{sz}", f"=-{cy}*{sz}+{sy}*{sx}*{cz}", f"={sy}*{cx}",
            f"={cx}*{sz}", f"={cx}*{cz}", f"=-{sx}",
            f"=-{sy}*{cz}+{cy}*{sx}*{sz}", f"={sy}*{sz}+{cy}*{sx}*{cz}", f"={cy}*{cx}"]


def _m_formulas(r):
    out = []
    for i in range(3):
        for j in range(3):
            out.append(f"={col(C_ROT + i * 3 + j)}{r}*{col(C_S + j)}{r}")  # M_ij = R_ij * S_j
        out.append(f"={col(C_T + i)}{r}")                                  # M_i4 = T_i
    return out


def build_matrices(wb, ctx):
    ws = wb[S_MAT]
    all_comps = ctx["scene"]["components"]
    comps = matrix_components(all_comps)
    ctx["mat_subset"] = comps
    note = (f"Formula volume is bounded: this sheet carries every group node and the first {wb_common.LEAVES_PER_ROOT} leaves of "
            f"each root tree ({len(comps)} of {len(all_comps)} components, tree order); the remaining parts are listed "
            "with S / R / T on sheet 3 and are checked against the code's matrices by tools/verify_workbook.py. ")
    title(ws, "4  Transformation Matrices (M1, M2, M3 ...): M = T x R x S with R = Ry x Rx x Rz",
          "One row per component (Lab-04 sheet p3: 'table maintain'). Blue = S, R, T exported by the program. "
          "R (3x3) and M (3x4; the 4th row is always 0 0 0 1) are Excel formulas: R_ij expanded from Ry*Rx*Rz, "
          "M_ij = R_ij * S_j and M_i4 = T_i. The code's own matrix (Transform::matrix(), exported row-major) is "
          "imported next to it; the last column is max|M_excel - M_code|. " + note + "A step-by-step 4x4 worked example "
          f"(select any ID) is at column {col(C_EX)}.", span=20)
    h, first = 5, 6
    last = first + len(comps) - 1
    hdr = (["Matrix", "ID", "Object", "Component", "Primitive", "Parent ID", "Sx", "Sy", "Sz", "Rx (deg)",
            "Ry (deg)", "Rz (deg)", "Tx", "Ty", "Tz"] + [f"R{i + 1}{j + 1}" for i in range(3) for j in range(3)]
           + [f"{n} (Excel)" for n in M_NAMES] + [f"{n} (code)" for n in M_NAMES] + ["max|Excel - code|"])
    header(ws, h, hdr)
    for i, c in enumerate(comps):
        r = first + i
        base = [f"M{i + 1}", c["id"], c["object"], c["name"], c["primitive"], c.get("parent_id") or "World"]
        for k, v in enumerate(base):
            put(ws, r, 1 + k, v, font=BOLD if k == 0 else None)
        ws.cell(row=r, column=5).fill = prim_fill(c["primitive"])
        for k, v in enumerate(c["S"] + c["R_deg"] + c["T"]):
            put(ws, r, C_S + k, v, FMT_3, INPUT_FILL)
        for k, f in enumerate(_rot_formulas(r)):
            put(ws, r, C_ROT + k, f, "0.0000")
        for k, f in enumerate(_m_formulas(r)):
            put(ws, r, C_M + k, f, "0.0000")
        for k, v in enumerate(c["local_matrix"][:12]):
            put(ws, r, C_CODE + k, v, "0.0000", INPUT_FILL)
        diffs = ",".join(f"ABS({col(C_M + k)}{r}-{col(C_CODE + k)}{r})" for k in range(12))
        put(ws, r, C_ERR, f"=MAX({diffs})", FMT_SCI)
    err_rng = f"{col(C_ERR)}{first}:{col(C_ERR)}{last}"
    ws.conditional_formatting.add(err_rng, CellIsRule(operator="lessThanOrEqual", formula=["$H$3"], fill=PASS_FILL))
    ws.conditional_formatting.add(err_rng, CellIsRule(operator="greaterThan", formula=["$H$3"], fill=FAIL_FILL))
    nowrap(ws, first, last, 1, C_ERR)
    _summary_row(ws, first, last)
    ctx["mat"] = {"first": first, "last": last, "rows": {c["id"]: first + i for i, c in enumerate(comps)}}
    ctx["refs"]["mat_err"] = f"'{S_MAT}'!$D$3"
    define_name(wb, "MatrixMaxError", S_MAT, "D3")
    _worked_example(ws, comps, first, last)
    widths(ws, [6, 9, 14, 22, 12, 9] + [7] * 9 + [8] * 9 + [8] * 12 + [8] * 12 + [12, 3])
    finish(ws, h, freeze="E6")
    ws.print_area = f"A1:{col(C_ERR)}{last}"
    ws.page_setup.fitToWidth = 2


def _summary_row(ws, first, last):
    put(ws, 3, 1, "Rows", font=BOLD)
    put(ws, 3, 2, f"=COUNTA(B{first}:B{last})", "0")
    put(ws, 3, 3, "max|Excel - code| (all rows)", font=BOLD)
    put(ws, 3, 4, f"=MAX({col(C_ERR)}{first}:{col(C_ERR)}{last})", FMT_SCI)
    put(ws, 3, 5, "Verdict", font=BOLD)
    put(ws, 3, 6, '=IF(D3<=H3,"Excel = code","CHECK")', font=BOLD)
    put(ws, 3, 7, "Tolerance", font=BOLD)
    put(ws, 3, 8, TOLERANCE, FMT_SCI)
    ws.conditional_formatting.add("F3", CellIsRule(operator="equal", formula=['"Excel = code"'], fill=PASS_FILL))


# ---- worked example -------------------------------------------------------------------------------------
def _matrix(ws, top, left, label, cells, fmt="0.0000"):
    """cells: 4x4 list of values/formulas. Returns a function giving the address of element (i, j)."""
    c = ws.cell(row=top, column=left, value=label)
    c.font = BOLD
    c.alignment = Alignment(vertical="bottom")
    for i in range(4):
        for j in range(4):
            put(ws, top + 1 + i, left + j, cells[i][j], fmt)
    return lambda i, j: f"{col(left + j)}{top + 1 + i}"


def _product(a, b):
    return [["=" + "+".join(f"{a(i, k)}*{b(k, j)}" for k in range(4)) for j in range(4)] for i in range(4)]


def _default_example(comps):
    """Most instructive row: most rotation axes used, then non-uniform scale, then a translation."""
    def score(c):
        rot = sum(1 for a in c["R_deg"] if abs(a) > 1e-6)
        scale = len({round(s, 6) for s in c["S"]})
        trans = sum(1 for t in c["T"] if abs(t) > 1e-6)
        return (rot, scale, trans)
    leaves = [c for c in comps if c["primitive"] != "Group"] or comps
    return max(leaves, key=score)["id"]


def _worked_example(ws, comps, first, last):
    L, top = C_EX, 6
    section(ws, top - 2, L, "Worked example: pick a component ID (drop-down); every matrix below is recomputed")
    put(ws, top, L, "Component ID", font=BOLD, align=Alignment(vertical="center"))
    ws.merge_cells(start_row=top, start_column=L, end_row=top, end_column=L + 1)
    sel = f"${col(L + 2)}${top}"
    put(ws, top, L + 2, _default_example(comps), fill=HILITE_FILL, font=Font(bold=True, size=12))
    dv = DataValidation(type="list", formula1=f"=$B${first}:$B${last}", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(sel.replace("$", ""))
    match = f"MATCH({sel},$B${first}:$B${last},0)"
    put(ws, top, L + 4, f"=INDEX($D${first}:$D${last},{match})&\" (\"&INDEX($C${first}:$C${last},{match})&\")\"")
    ws.merge_cells(start_row=top, start_column=L + 4, end_row=top, end_column=L + 9)
    vals = {}
    for k, (name, c0) in enumerate((("S", C_S), ("R (deg)", C_R), ("T", C_T))):
        r = top + 1 + k
        put(ws, r, L, name, font=BOLD)
        for a in range(3):
            letter = col(c0 + a)
            put(ws, r, L + 1 + a, f"=INDEX(${letter}${first}:${letter}${last},{match})", FMT_3, INPUT_FILL)
            vals[(name, a)] = f"{col(L + 1 + a)}{r}"
    sx_, sy_, sz_ = (vals[("S", a)] for a in range(3))
    ax, ay, az = (f"RADIANS({vals[('R (deg)', a)]})" for a in range(3))
    tx, ty, tz = (vals[("T", a)] for a in range(3))
    b1 = top + 5
    T = _matrix(ws, b1, L, "T (translation)", [[1, 0, 0, f"={tx}"], [0, 1, 0, f"={ty}"], [0, 0, 1, f"={tz}"],
                                               [0, 0, 0, 1]])
    Ry = _matrix(ws, b1, L + 5, "Ry (yaw about Y)", [[f"=COS({ay})", 0, f"=SIN({ay})", 0], [0, 1, 0, 0],
                                                     [f"=-SIN({ay})", 0, f"=COS({ay})", 0], [0, 0, 0, 1]])
    Rx = _matrix(ws, b1, L + 10, "Rx (pitch about X)", [[1, 0, 0, 0], [0, f"=COS({ax})", f"=-SIN({ax})", 0],
                                                         [0, f"=SIN({ax})", f"=COS({ax})", 0], [0, 0, 0, 1]])
    Rz = _matrix(ws, b1, L + 15, "Rz (roll about Z)", [[f"=COS({az})", f"=-SIN({az})", 0, 0],
                                                        [f"=SIN({az})", f"=COS({az})", 0, 0], [0, 0, 1, 0],
                                                        [0, 0, 0, 1]])
    S = _matrix(ws, b1, L + 20, "S (scale)", [[f"={sx_}", 0, 0, 0], [0, f"={sy_}", 0, 0], [0, 0, f"={sz_}", 0],
                                              [0, 0, 0, 1]])
    b2 = b1 + 7
    RyRx = _matrix(ws, b2, L, "Ry x Rx", _product(Ry, Rx))
    R = _matrix(ws, b2, L + 5, "R = Ry x Rx x Rz", _product(RyRx, Rz))
    TR = _matrix(ws, b2, L + 10, "T x R", _product(T, R))
    M = _matrix(ws, b2, L + 15, "M = T x R x S", _product(TR, S))
    code = [[f"=INDEX(${col(C_CODE + i * 4 + j)}${first}:${col(C_CODE + i * 4 + j)}${last},{match})"
             for j in range(4)] for i in range(3)] + [[0, 0, 0, 1]]
    MC = _matrix(ws, b2, L + 20, "M exported by the code", code)
    b3 = b2 + 7
    E = _matrix(ws, b3, L + 15, "|M - M_code|", [[f"=ABS({M(i, j)}-{MC(i, j)})" for j in range(4)]
                                                 for i in range(4)], FMT_SCI)
    put(ws, b3 + 1, L, "max |M - M_code|", font=BOLD, align=Alignment(vertical="top"))
    ws.merge_cells(start_row=b3 + 1, start_column=L, end_row=b3 + 1, end_column=L + 2)
    put(ws, b3 + 1, L + 3, "=MAX(" + ",".join(E(i, j) for i in range(4) for j in range(4)) + ")", FMT_SCI)
    note = ws.cell(row=b3 + 2, column=L, value=(
        "Order of operations on a vertex p: p' = T * Ry * Rx * Rz * S * p, i.e. scale first, then roll (Z), pitch "
        "(X) and yaw (Y), then translate. In the hierarchy, world = parent_world * M (sheet 5)."))
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=b3 + 2, start_column=L, end_row=b3 + 4, end_column=L + 13)
    for k in range(25):
        ws.column_dimensions[col(L + k)].width = 9


# ---- 5 Hierarchy ----------------------------------------------------------------------------------------
C_WX, C_WEX, C_WXL, C_WERR = 6, 9, 21, 33


def _lamp_chain(comps):
    by_id = {c["id"]: c for c in comps}
    lamp = next((c for c in comps if "lamp" in c["name"].lower() and "uav" in c["object"].lower()), None)
    if lamp is None:
        lamp = next((c for c in comps if "lamp" in c["name"].lower()), None)
    chain = set()
    while lamp is not None:
        chain.add(lamp["id"])
        lamp = by_id.get(lamp.get("parent_id"))
    return chain


def build_hierarchy(wb, ctx):
    ws = wb[S_HIER]
    comps = ctx["mat_subset"]
    all_comps = ctx["scene"]["components"]
    mat_rows = ctx["mat"]["rows"]
    title(ws, "5  Hierarchy: parent -> child tree with world matrices (world = parent_world x M)",
          "Tree in depth-first order (indent = depth). Blue: world matrix exported by the program at t = 0. White: "
          "the same world matrix rebuilt in Excel from the M formulas of sheet 4, multiplying down the tree "
          "(root: world = M). The highlighted chain is the hierarchical searchlight: UAV -> gimbal yaw -> gimbal "
          f"pitch -> lamp. Same subset as sheet 4 ({len(comps)} of {len(all_comps)} nodes: all group nodes + the first "
          f"{wb_common.LEAVES_PER_ROOT} leaves per root). v4: blue rows are animated joints (tank turret / gun / tracks, helicopter "
          "rotor head, radar head, blast-door leaves, flag segments ...), column AI gives the exported animation text "
          "and the table at column AK the deepest joint chain of every controllable object.", span=16)
    h, first = 4, 5
    hdr = (["Node (tree)", "ID", "Primitive", "Parent ID", "Depth", "World x", "World y", "World z"]
           + [f"W{n[1:]} (code)" for n in M_NAMES] + [f"W{n[1:]} (Excel)" for n in M_NAMES]
           + ["max|Excel - code|"])
    header(ws, h, hdr)
    rows = {c["id"]: first + i for i, c in enumerate(comps)}
    chain = _lamp_chain(comps)
    for i, c in enumerate(comps):
        r = first + i
        depth = c.get("depth", 0)
        cell = put(ws, r, 1, ("- " if depth else "") + c["name"])
        cell.alignment = Alignment(indent=depth * 2, vertical="top")
        for k, v in enumerate([c["id"], c["primitive"], c.get("parent_id") or "World", depth]):
            put(ws, r, 2 + k, v)
        ws.cell(row=r, column=3).fill = prim_fill(c["primitive"])
        for a in range(3):
            put(ws, r, C_WX + a, f"={col(C_WEX + a * 4 + 3)}{r}", FMT_3)
        for k, v in enumerate(c["world_matrix"][:12]):
            put(ws, r, C_WEX + k, v, "0.0000", INPUT_FILL)
        for k, f in enumerate(_world_formulas(c, r, rows, mat_rows[c["id"]])):
            put(ws, r, C_WXL + k, f, "0.0000")
        diffs = ",".join(f"ABS({col(C_WEX + k)}{r}-{col(C_WXL + k)}{r})" for k in range(12))
        put(ws, r, C_WERR, f"=MAX({diffs})", FMT_SCI)
        put(ws, r, C_WERR + 2, c["animation"] or "")
        if c["id"] in chain:
            for k in range(1, 6):
                ws.cell(row=r, column=k).fill = HILITE_FILL
                ws.cell(row=r, column=k).font = BOLD
        elif c["primitive"] == "Group" and c["animation"]:
            for k in range(1, 6):
                ws.cell(row=r, column=k).fill = JOINT_FILL
        elif depth == 0:
            for k in (1, 2, 4, 5):
                ws.cell(row=r, column=k).fill = TOTAL_FILL
            ws.cell(row=r, column=1).font = BOLD
    last = first + len(comps) - 1
    nowrap(ws, first, last, 1, C_WERR)
    err_rng = f"{col(C_WERR)}{first}:{col(C_WERR)}{last}"
    put(ws, 3, 1, "max|Excel - code| over all world matrices", font=BOLD)
    put(ws, 3, 2, f"=MAX({err_rng})", FMT_SCI)
    put(ws, 3, 3, f"=IF(B3<='{S_MAT}'!$H$3,\"Excel = code\",\"CHECK\")", font=BOLD)
    put(ws, 3, 4, "Highlighted chain", font=BOLD)
    put(ws, 3, 5, " -> ".join(c["name"] for c in comps if c["id"] in chain), fill=HILITE_FILL)
    ws.merge_cells(start_row=3, start_column=5, end_row=3, end_column=12)
    ws.conditional_formatting.add(err_rng, CellIsRule(operator="lessThanOrEqual", formula=[f"'{S_MAT}'!$H$3"],
                                                      fill=PASS_FILL))
    ws.conditional_formatting.add(err_rng, CellIsRule(operator="greaterThan", formula=[f"'{S_MAT}'!$H$3"],
                                                      fill=FAIL_FILL))
    ctx["refs"]["world_err"] = f"'{S_HIER}'!$B$3"
    define_name(wb, "WorldMaxError", S_HIER, "B3")
    header(ws, h, ["Animation (exported)"], start_col=C_WERR + 2)
    _joint_chain_table(ws, ctx, C_WERR + 4)
    widths(ws, [34, 9, 12, 9, 6, 9, 9, 9] + [9] * 24 + [12, 3, 50, 3, 22, 80, 8, 8])
    finish(ws, h, freeze="C5")


def _joint_chain_table(ws, ctx, left):
    section(ws, 3, left, "Deepest joint chain of every controllable (exported tree)")
    header(ws, 4, ["Controllable", "Chain (root -> deepest animated group)", "Levels", "Animated groups"],
           start_col=left)
    for i, ch in enumerate(articulation_chains(ctx["scene"]["components"], ctx["scene"].get("controllables", []))):
        r = 5 + i
        put(ws, r, left, ch["name"], font=BOLD)
        put(ws, r, left + 1, " -> ".join(ch["path"]))
        put(ws, r, left + 2, ch["levels"], "0", INPUT_FILL)
        put(ws, r, left + 3, ch["joints"], "0", INPUT_FILL)


def _world_formulas(c, r, rows, mrow):
    def m(i, j):
        return f"'{S_MAT}'!{col(C_M + i * 4 + j)}{mrow}"

    pid = c.get("parent_id")
    if not pid or pid not in rows:
        return [f"={m(i, j)}" for i in range(3) for j in range(4)]
    prow = rows[pid]

    def p(i, j):
        return f"{col(C_WXL + i * 4 + j)}{prow}"

    out = []
    for i in range(3):
        for j in range(4):
            terms = [f"{p(i, k)}*{m(k, j)}" for k in range(3)]
            if j == 3:
                terms.append(p(i, 3))
            out.append("=" + "+".join(terms))
    return out
