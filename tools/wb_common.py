"""Shared theme, loaders and cell helpers for the ThermalSense workbook generator.

Every number written by the sheet builders comes from the program's exported files (JSON / CSV) or is an
Excel formula over those cells. Only explanatory text is authored here.
"""
import json
import os

from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.hyperlink import Hyperlink

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "ThermalSense_Graphics_Documentation.xlsx")
SCRATCH = os.path.join(ROOT, "build", "workbook_img")

S_COVER = "0 Cover & Index"
S_INV = "1 Primitive Inventory"
S_VERT = "2 Primitive Vertex Data"
S_OBJ = "3 Object Construction"
S_MAT = "4 Transform Matrices"
S_HIER = "5 Hierarchy"
S_GEOM = "6 Geometry Statistics"
S_LIGHT = "7 Lighting"
S_SHADE = "8 Shading"
S_PIPE = "9 Render Pipeline"
S_TEST = "10 Testing"
S_PERF = "11 Performance"
S_BEHAV = "12 Behaviour"
S_TRACE = "13 Traceability"
S_CTRL = "14 Controllable Objects"
S_BASE = "15 Bunker & Base"
S_COMBAT = "16 Combat & Effects"
S_AUDIT = "17 Audit"
S_CV = "18 Computer Vision"
S_SQUAD = "19 Squad & Mission"
S_AUDIO = "20 Audio & HUD"
S_DAY = "21 Day Cycle & Lighting"
S_SOLD = "22 Soldiers & Routes"
S_CVN = "23 CV Noise & Seeds"
S_RT = "24 Ray Tracing"
S_FX = "25 Frame Explorer"
S_GAME = "26 Game Camera & Controls"
S_V7T = "27 v7 Tests"
SHEET_ORDER = [S_COVER, S_INV, S_VERT, S_OBJ, S_MAT, S_HIER, S_GEOM, S_LIGHT, S_SHADE, S_PIPE, S_TEST, S_PERF,
               S_BEHAV, S_TRACE, S_CTRL, S_BASE, S_COMBAT, S_AUDIT, S_CV, S_SQUAD, S_AUDIO, S_DAY, S_SOLD, S_CVN,
               S_RT, S_FX, S_GAME, S_V7T]

AUTHOR = "Md. Shakhoyat Rahman Shujon"
ROLL = "2107104"

# ---- theme ----------------------------------------------------------------------------------------------
DARK = "1F2937"
ACCENT = "0E7490"
MUTED = "6B7280"
HEADER_FILL = PatternFill("solid", fgColor=DARK)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color=DARK)
SUBTITLE_FONT = Font(italic=True, size=10, color=MUTED)
SECTION_FONT = Font(bold=True, size=12, color=ACCENT)
BOLD = Font(bold=True)
LINK_FONT = Font(color="1D4ED8", underline="single")
# start + end colour so the same fill also works inside conditional-formatting rules (DXF uses bgColor)
PASS_FILL = PatternFill(fill_type="solid", start_color="D1FAE5", end_color="D1FAE5")
FAIL_FILL = PatternFill(fill_type="solid", start_color="FEE2E2", end_color="FEE2E2")
PENDING_FILL = PatternFill(fill_type="solid", start_color="FEF9C3", end_color="FEF9C3")
INPUT_FILL = PatternFill("solid", fgColor="EFF6FF")      # values imported from the program's exports
FORMULA_FILL = PatternFill("solid", fgColor="FFFFFF")    # computed by Excel
NOTE_FILL = PatternFill("solid", fgColor="FEF3C7")       # hand-maintained text "from design notes" (not exported)
HILITE_FILL = PatternFill("solid", fgColor="FDE68A")
TOTAL_FILL = PatternFill("solid", fgColor="E5E7EB")
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Primitive colour code (matches the renderer's diagnostic view: cube cyan, sphere magenta, cylinder amber).
PRIM_COLOURS = {
    "Cube": ("CFF4FC", "06B6D4"),
    "Sphere": ("FCE7F3", "EC4899"),
    "Cylinder": ("FEF3C7", "F59E0B"),
    "Cylinder (frustum)": ("FFEDD5", "EA580C"),
    "Cone": ("ECFCCB", "65A30D"),
    "Terrain disc": ("E2E8E4", "4D6B54"),
    "Group": ("F3F4F6", "6B7280"),
}
OTHER_COLOURS = ("EDE9FE", "7C3AED")

FMT_INT = "#,##0"
FMT_3 = "0.000"
FMT_6 = "0.000000"
FMT_PCT = "0.0%"
FMT_SCI = "0.00E+00"
FMT_MS = "0.00"

WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def prim_fill(name, strong=False):
    light, dark = PRIM_COLOURS.get(name, OTHER_COLOURS)
    return PatternFill("solid", fgColor=dark if strong else light)


# ---- loaders --------------------------------------------------------------------------------------------
def path(rel):
    return os.path.join(ROOT, rel)


def load_json(rel, default=None):
    p = path(rel)
    if not os.path.exists(p):
        return default
    with open(p, encoding="utf-8") as f:
        return json.load(f)


# ---- cell helpers ---------------------------------------------------------------------------------------
def col(c):
    return get_column_letter(c)


def ref(ws_title, cell, absolute=True):
    """Cross-sheet reference like '3 Object Construction'!$B$5."""
    if absolute:
        letters = "".join(ch for ch in cell if ch.isalpha())
        digits = "".join(ch for ch in cell if ch.isdigit())
        cell = f"${letters}${digits}"
    return f"'{ws_title}'!{cell}"


def put(ws, row, column, value, fmt=None, fill=None, font=None, align=None, border=True):
    c = ws.cell(row=row, column=column, value=value)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if font:
        c.font = font
    c.alignment = align or WRAP_TOP
    if border:
        c.border = BORDER
    return c


def header(ws, row, headers, start_col=1, height=32):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[row].height = height


def title(ws, text, subtitle=None, span=10):
    ws.cell(row=1, column=1, value=text).font = TITLE_FONT
    ws.row_dimensions[1].height = 26
    if subtitle:
        c = ws.cell(row=2, column=1, value=subtitle)
        c.font = SUBTITLE_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
        ws.row_dimensions[2].height = max(30, 15 * (1 + len(subtitle) // (span * 14)))


def section(ws, row, column, text):
    c = ws.cell(row=row, column=column, value=text)
    c.font = SECTION_FONT
    return c


def widths(ws, values, start_col=1):
    for i, w in enumerate(values):
        ws.column_dimensions[col(start_col + i)].width = w


def finish(ws, header_row, freeze=None, landscape=True):
    """Frozen header, print setup (landscape, fit to width, repeat header row)."""
    ws.freeze_panes = freeze or f"A{header_row + 1}"
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = "ThermalSense  -  " + AUTHOR + " (" + ROLL + ")  -  page &P of &N"
    ws.sheet_view.zoomScale = 90


def style_chart(ch, legend="b"):
    """Visible axes (openpyxl hides them by default in recent Excel) and a bottom legend."""
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    if ch.legend is not None:
        ch.legend.position = legend
    return ch


def nowrap(ws, first_row, last_row, first_col=1, last_col=None):
    """Single-line cells for big numeric tables (keeps an indent if one was set)."""
    last_col = last_col or ws.max_column
    for row in ws.iter_rows(min_row=first_row, max_row=last_row, min_col=first_col, max_col=last_col):
        for c in row:
            c.alignment = Alignment(vertical="top", indent=c.alignment.indent)


def link(cell, sheet_title, target="A1", text=None):
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet_title}'!{target}", display=text)
    if text is not None:
        cell.value = text
    cell.font = LINK_FONT


def file_link(cell, rel_path, text=None):
    cell.hyperlink = rel_path.replace("\\", "/")
    cell.value = text or rel_path
    cell.font = LINK_FONT


def define_name(wb, name, sheet_title, cell_range):
    wb.defined_names[name] = DefinedName(name, attr_text=ref(sheet_title, cell_range))


def image(ws, anchor, src, box=None, width=None):
    """Embeds a (cropped, resized) copy of a report figure; returns the embedded height in pixels or 0."""
    from PIL import Image as PILImage

    if not os.path.exists(src):
        return 0
    os.makedirs(SCRATCH, exist_ok=True)
    im = PILImage.open(src).convert("RGB")
    if box:
        im = im.crop(box)
    if width and im.width > width:
        im = im.resize((width, round(im.height * width / im.width)), PILImage.LANCZOS)
    stem = os.path.splitext(os.path.basename(src))[0]
    out = os.path.join(SCRATCH, f"{stem}_{box[0] if box else 0}_{width or 0}.png")
    im.save(out)
    ws.add_image(XLImage(out), anchor)
    return im.height


# ---- transform sheets: bounded formula volume -----------------------------------------------------------
LEAVES_PER_ROOT = 80   # sheets 4 and 5 carry Excel formulas for every group node + the first N leaves of each root


GROUPS_PER_ROOT = None   # None = all group nodes; a number only in the reduced verification workbook


def set_leaves(n, groups=None):
    """Used by `build_workbook.py --leaves N [--groups G]`: reduced workbook for the slow pure-Python formula engine."""
    global LEAVES_PER_ROOT, GROUPS_PER_ROOT
    LEAVES_PER_ROOT = int(n)
    GROUPS_PER_ROOT = int(groups) if groups else None


def matrix_components(comps, per_root=None):
    """Subset of the exported components that gets Excel matrix formulas (sheets 4 + 5): every group node (so
    each articulation chain is complete) and the first `per_root` leaves of every root tree, in tree order.
    A leaf's parent is always a group, so the subset is closed under 'parent of'. Sheet 3 lists all components."""
    per_root = LEAVES_PER_ROOT if per_root is None else per_root
    keep, kept_ids, taken, groups = [], set(), 0, 0
    for c in comps:
        if c.get("depth", 0) == 0:
            taken = groups = 0
        if c["primitive"] == "Group":
            groups += 1
            if GROUPS_PER_ROOT is not None and groups > GROUPS_PER_ROOT:
                continue
            keep.append(c)
            kept_ids.add(c["id"])
        elif c.get("parent_id") in kept_ids or c.get("depth", 0) == 0:
            taken += 1
            if taken <= per_root:
                keep.append(c)
                kept_ids.add(c["id"])
    return keep



def articulation_chains(comps, controllables):
    """For every controllable: the deepest chain of animated GROUP nodes (joints), as a list of node names from its
    root, plus the number of animated groups. Everything is read from the exported component tree."""
    by_id = {c["id"]: c for c in comps}
    out = []
    for ctl in controllables:
        rid = ctl.get("root_id")
        if rid not in by_id:
            continue
        # subtree = following comps until the next node of depth <= root depth
        start = next(i for i, c in enumerate(comps) if c["id"] == rid)
        d0 = comps[start].get("depth", 0)
        sub = [comps[start]]
        for c in comps[start + 1:]:
            if c.get("depth", 0) <= d0:
                break
            sub.append(c)
        joints = [c for c in sub if c["primitive"] == "Group" and c["animation"]]
        deepest = max(joints, key=lambda c: c.get("depth", 0), default=comps[start])
        path, cur = [], deepest
        while cur is not None:
            path.append(cur["name"])
            cur = by_id.get(cur.get("parent_id")) if cur["id"] != rid else None
        path = list(reversed(path))
        runs = []  # "Flag strip x12" instead of twelve identical names
        for nme in path:
            if runs and runs[-1][0] == nme:
                runs[-1][1] += 1
            else:
                runs.append([nme, 1])
        out.append({"name": ctl["name"], "key": ctl.get("key", "-"),
                    "path": [n if k == 1 else f"{n} x{k}" for n, k in runs],
                    "levels": len(path), "joints": len(joints), "root_id": rid})
    return out
