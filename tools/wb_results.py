"""Sheets 10 (Testing), 11 (Performance) and 12 (Behaviour)."""
import csv

from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.formatting.rule import CellIsRule, FormulaRule

from wb_pipeline import phase_label
from wb_common import (BOLD, FAIL_FILL, FMT_3, FMT_INT, FMT_MS, FMT_PCT, INPUT_FILL, PASS_FILL, PENDING_FILL,
                       S_BEHAV, S_PERF, S_TEST, TOTAL_FILL, col, define_name, file_link, style_chart, finish, header, path, put,
                       section, title, widths)

TRACK_REQUIREMENT_M = 0.5   # requirement B01 (median ground error while detected), a specification, not a result
PERCENTILE_Q = 0.9
DATA_COL = 10               # first column of the raw trace on sheet 12


# ---- 12 Behaviour ---------------------------------------------------------------------------------------
def _read_trace(rel):
    try:
        with open(path(rel), newline="") as f:
            rows = list(csv.reader(f))
    except FileNotFoundError:
        return [], []
    return (rows[0], rows[1:]) if rows else ([], [])


def _as_number(s):
    try:
        v = float(s)
        return int(v) if v.is_integer() and "." not in s else v
    except ValueError:
        return s


def build_behaviour(wb, ctx):
    ws = wb[S_BEHAV]
    hdr, rows = _read_trace("build/trace.csv")
    summ = ctx["trace_summary"]
    title(ws, "12  Behaviour (UAV tracking vs hiding rover)",
          "Raw samples of ThermalSense.exe --trace (build/trace.csv) are imported from column "
          f"{col(DATA_COL)}. Every KPI is an Excel formula over those columns; the last column repeats the value "
          "computed independently by tools/summarize_trace.py as a cross-check.", span=8)
    if not rows:
        put(ws, 4, 1, "PENDING - build/trace.csv not found; run ThermalSense.exe --trace build/trace.csv 180",
            fill=PENDING_FILL, font=BOLD)
        finish(ws, 4)
        return
    need = ["t", "uav_mode", "detected", "rover_hidden", "est_error"]
    ci = {name: hdr.index(name) for name in need if name in hdr}
    h, first = 4, 5
    last = first + len(rows) - 1
    extra = ["valid error (m)", "hide starts", "REACQUIRE starts"]
    header(ws, h, hdr + extra, start_col=DATA_COL)
    L = {name: col(DATA_COL + i) for name, i in ci.items()}
    xv, xh, xr = (col(DATA_COL + len(hdr) + k) for k in range(3))
    for i, row in enumerate(rows):
        r = first + i
        for k, v in enumerate(row):
            ws.cell(row=r, column=DATA_COL + k, value=_as_number(v))
        ws.cell(row=r, column=DATA_COL + len(hdr), value=f'=IF({L["est_error"]}{r}>=0,{L["est_error"]}{r},"")')
        prev_h = f"{L['rover_hidden']}{r - 1}" if i else "0"
        prev_m = f"{L['uav_mode']}{r - 1}" if i else '""'
        ws.cell(row=r, column=DATA_COL + len(hdr) + 1,
                value=f"=IF(AND({L['rover_hidden']}{r}=1,{prev_h}=0),1,0)")
        ws.cell(row=r, column=DATA_COL + len(hdr) + 2,
                value=f'=IF(AND({L["uav_mode"]}{r}="REACQUIRE",{prev_m}<>"REACQUIRE"),1,0)')
    R = lambda letter: f"${letter}${first}:${letter}${last}"  # noqa: E731
    reacq_py = sum(1 for _, m in summ["transitions"] if m == "REACQUIRE")
    kpis = [
        ("Samples", f"=COUNT({R(L['t'])})", FMT_INT, summ["samples"]),
        ("Duration (s)", f"=MAX({R(L['t'])})", "0.00", summ["duration_s"]),
        ("Sample interval (s)", f"={L['t']}{first + 1}-{L['t']}{first}", "0.00", None),
        ("Detection rate (share of samples with a detection)", f"=AVERAGE({R(L['detected'])})", FMT_PCT,
         summ["detected_share"]),
        ("Rover hidden (share of samples)", f"=AVERAGE({R(L['rover_hidden'])})", FMT_PCT,
         summ["rover_hidden_share"]),
        ("Tracking error median (m, while detected)", f"=MEDIAN({R(xv)})", FMT_3, summ["est_error_median"]),
        ("Tracking error p90 (m)", f"=PERCENTILE({R(xv)},$B$16)", FMT_3, summ["est_error_p90"]),
        ("Tracking error max (m)", f"=MAX({R(xv)})", FMT_3, summ["est_error_max"]),
        ("Hide events (rover_hidden 0 -> 1)", f"=SUM({R(xh)})", FMT_INT, None),
        ("REACQUIRE episodes", f"=SUM({R(xr)})", FMT_INT, reacq_py),
    ]
    header(ws, h, ["Metric", "Excel formula", "summarize_trace.py"])
    for i, (label, f, fmt, py) in enumerate(kpis):
        r = first + i
        put(ws, r, 1, label, font=BOLD)
        put(ws, r, 2, f, fmt)
        put(ws, r, 3, py if py is not None else "-", fmt if py is not None else None, INPUT_FILL)
    put(ws, 16, 1, "Percentile used for p90", font=BOLD)
    put(ws, 16, 2, PERCENTILE_Q, "0.00")
    put(ws, 17, 1, "Requirement: median error below (m)", font=BOLD)
    put(ws, 17, 2, TRACK_REQUIREMENT_M, FMT_3)
    put(ws, 18, 1, "Tracking verdict", font=BOLD)
    put(ws, 18, 2, '=IF(B10<B17,"PASS","FAIL")', font=BOLD)
    ws.conditional_formatting.add("B18", CellIsRule(operator="equal", formula=['"PASS"'], fill=PASS_FILL))
    ws.conditional_formatting.add("B18", CellIsRule(operator="equal", formula=['"FAIL"'], fill=FAIL_FILL))
    refs = ctx["refs"]
    refs.update(beh_median=f"'{S_BEHAV}'!$B$10", beh_p90=f"'{S_BEHAV}'!$B$11", beh_max=f"'{S_BEHAV}'!$B$12",
                beh_hide=f"'{S_BEHAV}'!$B$13", beh_reacq=f"'{S_BEHAV}'!$B$14", beh_req=f"'{S_BEHAV}'!$B$17",
                beh_detect=f"'{S_BEHAV}'!$B$8")
    define_name(wb, "DetectionRate", S_BEHAV, "B8")
    define_name(wb, "TrackingErrorMedian", S_BEHAV, "B10")
    _mode_table(ws, summ, R(L["uav_mode"]), 21)
    _behaviour_charts(ws, first, last, L, xv)
    widths(ws, [44, 16, 16, 3, 3, 3, 3, 3, 3] + [9] * (len(hdr) + 3))
    ws.column_dimensions[col(DATA_COL + 1)].width = 12
    finish(ws, h, freeze=f"A{first}")


def _mode_table(ws, summ, mode_rng, top):
    section(ws, top - 1, 1, "UAV mode shares (COUNTIF over uav_mode)")
    header(ws, top, ["UAV mode", "Share of samples", "Time in mode (s)"])
    for i, m in enumerate(summ["mode_share"]):
        r = top + 1 + i
        put(ws, r, 1, m, font=BOLD)
        put(ws, r, 2, f"=COUNTIF({mode_rng},A{r})/$B$5", FMT_PCT)
        put(ws, r, 3, f"=COUNTIF({mode_rng},A{r})*$B$7", "0.0")


def _behaviour_charts(ws, first, last, L, xv):
    xs = Reference(ws, min_col=_idx(L["t"]), min_row=first, max_row=last)
    ch = ScatterChart()
    ch.title, ch.style = "Tracking error over time (m, only while detected)", 13
    ch.x_axis.title, ch.y_axis.title = "t (s)", "error (m)"
    s = Series(Reference(ws, min_col=_idx(xv), min_row=first - 1, max_row=last), xs, title_from_data=True)
    s.marker.symbol, s.marker.size = "circle", 3
    s.graphicalProperties.line.noFill = True
    ch.series.append(s)
    ch.height, ch.width = 8, 16
    ch.legend = None
    ws.add_chart(style_chart(ch), "A28")
    ch2 = ScatterChart()
    ch2.title, ch2.style = "Detection and hiding over time (1 = yes)", 13
    ch2.x_axis.title = "t (s)"
    for name in ("detected", "rover_hidden"):
        s2 = Series(Reference(ws, min_col=_idx(L[name]), min_row=first - 1, max_row=last), xs, title_from_data=True)
        s2.marker.symbol = "none"
        ch2.series.append(s2)
    ch2.height, ch2.width = 7, 16
    ws.add_chart(style_chart(ch2), "A45")


def _idx(letter):
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(letter)


# ---- 10 Testing -----------------------------------------------------------------------------------------
def build_testing(wb, ctx):
    ws = wb[S_TEST]
    tests, visual, refs = ctx["tests"], ctx["visual"], ctx["refs"]
    title(ws, "10  Testing",
          "Automated unit/integration tests (build/tests.exe -> build/test_results.json), visual checks of real "
          "captures (docs/visual_tests.json) and behaviour checks computed live from the trace (sheet 12). The pass "
          "rate is a formula; status cells are conditionally formatted.", span=7)
    h, first = 5, 6
    header(ws, h, ["Test ID", "Category", "Description", "Expected result", "Observed result", "Status",
                   "Evidence"])
    r = first
    for t in tests:
        _test_row(ws, r, t["id"], "Automated (tests/test_main.cpp)", t, "build/test_results.json")
        r += 1
    for v in visual:
        ev = v.get("evidence", "")
        _test_row(ws, r, v["id"], "Visual check of a real capture", v, ev.split(",")[0].strip(), ev)
        r += 1
    if "beh_median" in refs:
        put(ws, r, 1, "B01", font=BOLD)
        put(ws, r, 2, "Measured (--trace), formula")
        put(ws, r, 3, "Computer-vision tracking accuracy over the traced run")
        put(ws, r, 4, f'="Median ground-position error < "&{refs["beh_req"]}&" m while detected"')
        put(ws, r, 5, f'="median "&ROUND({refs["beh_median"]},3)&" m, p90 "&ROUND({refs["beh_p90"]},3)&'
                      f'" m, max "&ROUND({refs["beh_max"]},3)&" m"')
        put(ws, r, 6, f'=IF({refs["beh_median"]}<{refs["beh_req"]},"PASS","FAIL")', font=BOLD)
        file_link(put(ws, r, 7, ""), "build/trace.csv")
        r += 1
        put(ws, r, 1, "B02", font=BOLD)
        put(ws, r, 2, "Measured (--trace), formula")
        put(ws, r, 3, "Rover breaks line of sight and the UAV re-acquires it")
        put(ws, r, 4, "At least one REACQUIRE episode")
        put(ws, r, 5, f'={refs["beh_reacq"]}&" REACQUIRE episodes, "&{refs["beh_hide"]}&" hide events"')
        put(ws, r, 6, f'=IF({refs["beh_reacq"]}>0,"PASS","FAIL")', font=BOLD)
        file_link(put(ws, r, 7, ""), "build/trace.csv")
        r += 1
    last = r - 1
    st = f"$F${first}:$F${last}"
    ws.conditional_formatting.add(st, CellIsRule(operator="equal", formula=['"PASS"'], fill=PASS_FILL))
    ws.conditional_formatting.add(st, CellIsRule(operator="equal", formula=['"FAIL"'], fill=FAIL_FILL))
    labels = [("Tests", f"=COUNTA($A${first}:$A${last})", FMT_INT), ("Passed", f'=COUNTIF({st},"PASS")', FMT_INT),
              ("Failed", f'=COUNTIF({st},"FAIL")', FMT_INT), ("Pass rate", "=D3/B3", FMT_PCT)]
    for k, (lab, f, fmt) in enumerate(labels):
        put(ws, 3, 1 + 2 * k, lab, font=BOLD, fill=TOTAL_FILL)
        put(ws, 3, 2 + 2 * k, f, fmt, font=BOLD)
    ws.conditional_formatting.add("H3", CellIsRule(operator="equal", formula=["1"], fill=PASS_FILL))
    ws.conditional_formatting.add("H3", CellIsRule(operator="lessThan", formula=["1"], fill=FAIL_FILL))
    refs.update(tests_total=f"'{S_TEST}'!$B$3", tests_passed=f"'{S_TEST}'!$D$3", pass_rate=f"'{S_TEST}'!$H$3")
    define_name(wb, "TestPassRate", S_TEST, "H3")
    ctx["test_rows"] = (first, last)
    widths(ws, [9, 30, 44, 40, 48, 9, 34])
    finish(ws, h)


def _test_row(ws, r, tid, category, t, evidence, evidence_text=None):
    put(ws, r, 1, tid, font=BOLD)
    put(ws, r, 2, category)
    put(ws, r, 3, t["description"])
    put(ws, r, 4, t["expected"])
    put(ws, r, 5, t["observed"])
    put(ws, r, 6, t["status"], font=BOLD)
    c = put(ws, r, 7, "")
    if evidence:
        file_link(c, evidence, evidence_text or evidence)


# ---- 11 Performance -------------------------------------------------------------------------------------
def build_performance(wb, ctx):
    ws = wb[S_PERF]
    bench = ctx["bench"] or {}
    phases = bench.get("phases", [])
    title(ws, "11  Performance",
          "Measured by ThermalSense.exe --benchmark (build/benchmark.json). Frame time and speed-up columns are "
          "formulas. Test setup: " + ", ".join(f"{k} = {v}" for k, v in bench.items() if not isinstance(v, list)),
          span=12)
    if not phases:
        put(ws, 4, 1, "PENDING - build/benchmark.json has no phases; rerun ThermalSense.exe --benchmark",
            fill=PENDING_FILL, font=BOLD)
        finish(ws, 4)
        return
    keys = [k for k in dict.fromkeys(k for p in phases for k in p)
            if k != "name" and not any(isinstance(p.get(k), (list, dict)) for p in phases)]  # passes -> sheet 9
    h, first = 4, 5
    last = first + len(phases) - 1
    hdr = ["Configuration"] + keys + ["1000 / FPS (ms)", "FPS relative to first row"]
    header(ws, h, hdr, height=45)
    fps_col = col(2 + keys.index("fps")) if "fps" in keys else None
    for i, p in enumerate(phases):
        r = first + i
        put(ws, r, 1, phase_label(p, i), font=BOLD)
        for k, key in enumerate(keys):
            v = p.get(key, "n/a")
            fmt = FMT_MS if isinstance(v, float) else FMT_INT
            put(ws, r, 2 + k, v, fmt, INPUT_FILL)
        if fps_col:
            put(ws, r, 2 + len(keys), f"=1000/{fps_col}{r}", FMT_MS)
            put(ws, r, 3 + len(keys), f"={fps_col}{r}/{fps_col}${first}", FMT_PCT)
    if fps_col:
        ctx["refs"]["fps"] = f"MAX('{S_PERF}'!${fps_col}${first}:${fps_col}${last})"
        define_name(wb, "BestFPS", S_PERF, f"{fps_col}{first}:{fps_col}{last}")
    if "scene_parts" in keys and "parts" in ctx["refs"]:
        pc = col(2 + keys.index("scene_parts"))
        put(ws, last + 2, 1, "Freshness check", font=BOLD)
        put(ws, last + 2, 2, f'=IF({pc}{first}={ctx["refs"]["parts"]},"benchmark scene = current scene ("&'
                             f'{pc}{first}&" parts)","STALE: measured with "&{pc}{first}&" parts, scene now has "&'
                             f'{ctx["refs"]["parts"]}&" - rerun ThermalSense.exe --benchmark")')
        ws.merge_cells(start_row=last + 2, start_column=2, end_row=last + 2, end_column=10)
        ws.conditional_formatting.add(f"B{last + 2}", FormulaRule(formula=[f'LEFT(B{last + 2},5)="STALE"'],
                                                                  fill=PENDING_FILL))
    anchor_row = last + 4
    for i, key in enumerate(k for k in ("fps", "frame_ms", "vertices_main", "draw_calls_main") if k in keys):
        ch = BarChart()
        ch.type = "col"
        ch.title = key.replace("_", " ")
        c = 2 + keys.index(key)
        ch.add_data(Reference(ws, min_col=c, min_row=h, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
        ch.legend = None
        ch.height, ch.width = 7, 12
        ws.add_chart(style_chart(ch), f"{'A' if i % 2 == 0 else 'G'}{anchor_row + 15 * (i // 2)}")
    widths(ws, [30] + [12] * (len(keys) + 2))
    finish(ws, h)
