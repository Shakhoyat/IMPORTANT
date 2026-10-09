"""Sheet 32 of the workbook: v9 M3 character motor, spring camera and locomotion.

Sources (written by the program or its tests; nothing typed by hand except wording and the design constants, which are
the values in src/scene/CharacterMotor.h / CameraRig.h):
  docs/v9_notes/m3/m3_video.log    ThermalSense.exe --input-script docs/v9_notes/m3/m3_video.txt ('script log' lines:
                                   sim time, ground speed, sprint / crouch blends, FOV kick, spring-arm length)
  build/test_results.json          T350-T359 (status via sheet 10)
  docs/v9_notes/m3/shots_*/        play_* captures before / after (headless, 1920x1080 EPIC, seed 1)
  docs/v9_notes/m3/exe_frames/     frames of the exe session
Blue = imported, white = Excel formula, amber = design-note text.

Usage (tools/build_workbook.py imports the v9 modules; this one adds):
    from wb_v9_m3 import S_V9M3, V9M3_INFO, build_v9_m3
    SHEET_INFO[S_V9M3] = V9M3_INFO;  build_v9_m3(wb, ctx)   (creates the sheet when it is not there yet)
"""
import os
import re

from openpyxl.chart import Reference, ScatterChart, Series

from wb_common import BOLD, INPUT_FILL, NOTE_FILL, TOTAL_FILL, define_name, file_link, finish, header, path, put, section, style_chart, \
    title, widths
from wb_v6 import pending
from wb_v9 import _images, _kpi, tests_block

S_V9M3 = "32 v9 Character Motor & Camera"
V9M3_INFO = ("v9 M3: critically damped character motor (speed springs, coyote time / jump buffer, step-up, slope limit), spring-arm "
             "camera, distance-driven gait phase (no foot sliding); exe speed curves recomputed by formula; tests T350-T359")
INFO = {S_V9M3: V9M3_INFO}

LOG = "docs/v9_notes/m3/m3_video.log"
LOG_RE = re.compile(r"script log (\S+) \| .*?cam CONTROL (\S+) .*?pos (\S+) (\S+) (\S+) yaw (\S+) speed (\S+) stance (\d) (\S+) "
                    r"eye (\S+) (\S+) (\S+) simt (\S+) sprint (\S+) crouch (\S+) fovkick (\S+) arm (\S+)")

# Design constants (src/scene/CharacterMotor.h MotorParams, src/scene/CameraRig.h ThirdPersonParams)
PARAMS = [
    ("walk speed", 1.3, "m/s", "WalkSpeed"),
    ("sprint speed", 3.4, "m/s", "RunSpeed"),
    ("crouch speed", 0.8, "m/s", "CrouchSpeed"),
    ("omega accelerate", 17.0, "1/s", "OmegaAcc"),
    ("omega brake", 30.0, "1/s", "OmegaStop"),
    ("air control", 0.2, "x omega", "AirControl"),
    ("stance blend time", 0.2, "s", "BlendTime"),
    ("coyote time", 0.12, "s", "Coyote"),
    ("jump buffer", 0.15, "s", "JumpBuffer"),
    ("slope limit", 40.0, "deg", "SlopeLimit"),
    ("step height", 0.45, "m", "StepHeight"),
    ("arm length", 3.1, "m", "ArmLength"),
    ("arm ease-out tau", 0.28, "s (omega = 2 / tau)", "ArmTau"),
    ("sprint FOV kick", 6.0, "deg", "FovKick"),
]

CONCEPTS = [
    ("Critically damped spring",
     "x'' = -2 w x' - w^2 (x - goal) (zeta = 1). Exact step over dt for a constant goal: e(t) = (e0 + (v0 + w e0) t) e^(-w t). From rest the "
     "step response 1 - (1 + w t) e^(-w t) rises monotonically: no overshoot, zero velocity at both ends, the fastest approach without "
     "oscillation. 90 % of a step after w t = 3.89.",
     "CharacterMotor.cpp critDampStep; speed, camera pivot, arm ease-out, step smoothing"),
    ("Velocity response",
     "The horizontal velocity follows the wish velocity on that spring: w = 17 1/s accelerating (90 % of walk in 0.23 s), 30 1/s braking "
     "(below 5 cm/s in 0.17 s), in the air a fifth of the rate and no braking. v1 had linear ramps of 2 m/s^2 (0.65 s to walk speed) and "
     "an abrupt stop.",
     "CharacterMotor::step (3)"),
    ("Coyote time / input buffer",
     "A jump press is remembered for 0.15 s and fires on the first grounded tick (pressed just before landing); after walking off a ledge "
     "a jump still fires for 0.12 s (not after a jump: no double jump). Both are counted in fixed 1/60 s ticks, so a replay reproduces them.",
     "CharacterMotor::step (5)"),
    ("Step-up, slope limit, step smoothing",
     "Box tops up to 0.45 m above the feet are floor (0.35 m steps pass, 0.5 m blocks); ground steeper than 40 deg removes the uphill part "
     "of the step (multi-scale gradient: a step is not a slope). A step moves the collision feet at once; the drawn feet follow on a "
     "spring (no pop).",
     "WalkSpace stepFoot, groundSlope, CharacterMotor::step (4, 7)"),
    ("Spring-arm camera",
     "Pivot follows the soldier on critically damped springs (w = 2 / tau); the arm is sphere-cast (NearBox oriented boxes + sphereCast): "
     "it pulls in at once to the first contact (never inside a wall) and eases back out on the spring (w = 7.1 1/s). Shoulder offset "
     "0.62 m; mouse look is applied directly.",
     "CameraRig ThirdPersonCam::update, AppPlay thirdPersonTarget"),
    ("Distance-driven gait phase",
     "phase += distance / stride(speed): in stance the ankle moves back relative to the hip by stride x d(phase) = distance, exactly the body "
     "displacement, so the stance foot stands still at every speed. A time-driven phase (or a stride that jumps between walk and run "
     "parameters) makes the feet slide. Gait parameters are blended idle -> walk -> run -> crouch walk by the actual ground speed; stance "
     "feet are world-locked and solved with two-link IK + hip roll (strafe / backwards), standing feet sit on slopes.",
     "Gait locomotionGait, Garrison Soldier::animate"),
    ("Head bob, FOV kick",
     "First person: two dips per stride from the gait phase, 1.2 cm walking .. 3 cm sprinting (App::headBob_, off -> stable eye); +6 deg "
     "field of view at full sprint speed, eased; the v8 landing dip stays.",
     "CharacterMotor headBobOffset, AppGame targetPose / updateGameCamera"),
]


def _logs():
    p = path(LOG)
    if not os.path.exists(p):
        return []
    out = []
    with open(p, encoding="utf-8", errors="replace") as f:
        for line in f:
            m = LOG_RE.search(line)
            if not m:
                continue
            g = m.groups()
            out.append({"label": g[0], "view": g[1], "x": float(g[2]), "y": float(g[3]), "z": float(g[4]), "speed": abs(float(g[6])),
                        "air": g[8] == "AIR", "eye_y": float(g[10]), "t": float(g[12]), "sprint": float(g[13]), "crouch": float(g[14]),
                        "fov": float(g[15]), "arm": float(g[16])})
    return out


def _group(logs, prefix):
    rows = [L for L in logs if re.fullmatch(prefix + r"_\d+", L["label"])]
    return rows


def _scatter(ws, anchor, title_text, xcol, ycols, first, last, hdr_row, ytitle, xtitle="t (s)"):
    ch = ScatterChart()
    ch.title, ch.style = title_text, 13
    ch.x_axis.title, ch.y_axis.title = xtitle, ytitle
    xs = Reference(ws, min_col=xcol, min_row=first, max_row=last)
    for c in ycols:
        s = Series(Reference(ws, min_col=c, min_row=hdr_row, max_row=last), xs, title_from_data=True)
        s.marker.symbol, s.smooth = "circle", False
        ch.series.append(s)
    ch.height, ch.width = 7.5, 13
    style_chart(ch)
    ws.add_chart(ch, anchor)


def _curve(ws, r, rows, label, model, threshold=None):
    """Measured speed rows (t from the first sample, which comes one 1/60 s tick after the key event) + the model curve by
    formula; column F: where the measured curve crosses 'threshold' (an Excel formula: linear interpolation between the two
    samples around it). model: 'rise' / 'stop' / None. Returns (next row, first data row, last data row, header row)."""
    header(ws, r, ["t (s)", "measured v (m/s)", "model v (m/s)", "sprint blend", "label", "crossing (s)"], height=22)
    h = r
    a = r + 1
    t0 = rows[0]["t"]
    for i, L in enumerate(rows):
        q = a + i
        put(ws, q, 1, round(L["t"] - t0, 4), "0.000", INPUT_FILL)
        put(ws, q, 2, L["speed"], "0.00", INPUT_FILL)
        if model == "rise":   # first sample is one tick after the press: shift by it
            put(ws, q, 3, f"=WalkSpeed*(1-(1+OmegaAcc*(A{q}+1/60))*EXP(-OmegaAcc*(A{q}+1/60)))", "0.00")
        elif model == "stop":
            put(ws, q, 3, f"=WalkSpeed*(1+OmegaStop*(A{q}+1/60))*EXP(-OmegaStop*(A{q}+1/60))", "0.00")
        else:
            put(ws, q, 3, None)
        put(ws, q, 4, L["sprint"], "0.00", INPUT_FILL)
        put(ws, q, 5, L["label"])
        if threshold and i > 0:
            th, p0 = threshold, q - 1
            cond = f"B{p0}<{th},B{q}>={th}" if model not in ("stop", "stop_free") else f"B{p0}>={th},B{q}<{th}"
            put(ws, q, 6, f'=IF(AND({cond}),A{p0}+({th}-B{p0})/(B{q}-B{p0})*(A{q}-A{p0})+1/60,"")', "0.000")
    b = a + len(rows) - 1
    return b + 2, a, b, h


def build_v9_m3(wb, ctx):
    ws = wb[S_V9M3] if S_V9M3 in wb.sheetnames else wb.create_sheet(S_V9M3)
    title(ws, "32  v9 M3: Character Motor, Spring Camera, Locomotion",
          "The possessed soldier moves on a character motor (scene/CharacterMotor.h) run once per fixed 1/60 s tick: critically damped "
          "speed springs instead of linear ramps, sprint / crouch blends over 0.2 s, coyote time and a jump buffer, step-up and a 40 deg "
          "slope limit on the WalkSpace collision, smooth step-down. The third-person camera is a critically damped spring arm with "
          "sphere-cast collision; the gait phase advances by distance travelled so the feet never slide. Below: the concepts, the "
          "design constants (named cells), the speed curves logged by the real exe recomputed with Excel formulas, the tests and the "
          "captures.", span=14)
    r = 5
    section(ws, r, 1, "Concepts")
    r += 1
    header(ws, r, ["Concept", "How it works", "Code"], height=20)
    for i, (a, b, c) in enumerate(CONCEPTS):
        q = r + 1 + i
        put(ws, q, 1, a, fill=NOTE_FILL, font=BOLD)
        put(ws, q, 2, b, fill=NOTE_FILL)
        put(ws, q, 3, c, fill=NOTE_FILL)
        ws.row_dimensions[q].height = 78
    r += len(CONCEPTS) + 2

    section(ws, r, 1, "Design constants (CharacterMotor.h, CameraRig.h) and what they imply (formulas)")
    r += 1
    header(ws, r, ["Constant", "Value", "Unit", "Name"], height=20)
    for i, (lab, v, unit, name) in enumerate(PARAMS):
        q = r + 1 + i
        put(ws, q, 1, lab, font=BOLD)
        put(ws, q, 2, v, "0.###", NOTE_FILL)
        put(ws, q, 3, unit)
        put(ws, q, 4, name)
        define_name(wb, name, S_V9M3, f"B{q}")
    q = r + len(PARAMS) + 2
    derived = [
        ("90 % rise time, accelerating (3.89 / w)", "=3.8897/OmegaAcc", "0.000 s"),
        ("walk speed after 0.25 s (step response)", "=WalkSpeed*(1-(1+OmegaAcc*0.25)*EXP(-OmegaAcc*0.25))", "0.000 m/s"),
        ("speed 0.2 s after release from walk", "=WalkSpeed*(1+OmegaStop*0.2)*EXP(-OmegaStop*0.2)", "0.0000 m/s"),
        ("tan(slope limit) = max rise per metre", "=TAN(RADIANS(SlopeLimit))", "0.000"),
        ("jump apex (v^2 / 2 g, 4.8 m/s, 16 m/s^2)", "=4.8^2/(2*16)", "0.00 m"),
        ("arm ease-out omega (2 / tau)", "=2/ArmTau", "0.00 1/s"),
        ("arm 90 % ease-out time", "=3.8897/(2/ArmTau)", "0.000 s"),
        ("stride per cycle at walk / run (m)", "1.30 / 2.00", None),
    ]
    header(ws, q, ["Derived", "Value"], height=20)
    for i, (lab, f, fmt) in enumerate(derived):
        put(ws, q + 1 + i, 1, lab, font=BOLD)
        put(ws, q + 1 + i, 2, f, fmt.split(" ")[0] if fmt else None)
        if fmt and " " in fmt:
            put(ws, q + 1 + i, 3, fmt.split(" ", 1)[1])
    r = q + len(derived) + 3

    logs = _logs()
    section(ws, r, 1, "ThermalSense.exe input-script session (docs/v9_notes/m3/m3_video.txt): logged speed curves vs the spring model")
    r += 1
    file_link(put(ws, r, 1, None), LOG, "exe log (m3_video.log)")
    file_link(put(ws, r, 3, None), "docs/v9_notes/m3/m3_session.mp4", "session video (MP4)")
    r += 2
    if not logs:
        r = pending(ws, r, LOG, span=6)
    else:
        kpis = []
        for prefix, model, label in [("walk_rise", "rise", "Walk: W pressed (rise)"), ("walk_stop", "stop", "Walk: W released (stop)"),
                                     ("sprint_rise", None, "Sprint: Shift + W pressed"), ("sprint_stop", None, "Sprint: released")]:
            rows = _group(logs, prefix)
            if not rows:
                continue
            section(ws, r, 1, label)
            r += 1
            top = r
            th = {"walk_rise": "0.9*WalkSpeed", "walk_stop": "0.05", "sprint_rise": "0.9*RunSpeed", "sprint_stop": "0.05"}[prefix]
            r, a, b, h = _curve(ws, r, rows, label, model if prefix != "sprint_stop" else "stop_free", threshold=th)
            name = {"walk_rise": "Exe walk 90 % rise (s)", "walk_stop": "Exe walk stop < 5 cm/s (s)",
                    "sprint_rise": "Exe sprint 90 % rise (s)", "sprint_stop": "Exe sprint stop < 5 cm/s (s)"}[prefix]
            kpis.append((name, f"=MIN(F{a}:F{b})", "0.000"))
            _scatter(ws, f"G{top}", label, 1, [2, 3] if model else [2, 4], a, b, h, "m/s")
            r = max(r, top + 16)
        _kpi(ws, kpis[:4])
        # crouch / stand blend and the eye height
        for prefix, label in [("crouch_blend", "Crouch (Left Alt): blend and first-person eye height"),
                              ("stand_blend", "Stand up: blend and eye height")]:
            rows = _group(logs, prefix)
            if not rows:
                continue
            section(ws, r, 1, label)
            r += 1
            top = r
            header(ws, r, ["t (s)", "crouch blend", "eye height above feet (m)", "smoothstep model"], height=22)
            h, a, t0 = r, r + 1, rows[0]["t"]
            for i, L in enumerate(rows):
                q = a + i
                put(ws, q, 1, round(L["t"] - t0, 4), "0.000", INPUT_FILL)
                put(ws, q, 2, L["crouch"], "0.00", INPUT_FILL)
                put(ws, q, 3, round(L["eye_y"] - L["y"], 3), "0.00", INPUT_FILL)
                ramp = f"MIN(1,MAX(0,(A{q})/BlendTime))"
                f = f"={ramp}^2*(3-2*{ramp})" if prefix == "crouch_blend" else f"=1-{ramp}^2*(3-2*{ramp})"
                put(ws, q, 4, f, "0.00")
            b = a + len(rows) - 1
            _scatter(ws, f"G{top}", label, 1, [2, 4], a, b, h, "blend")
            r = max(b + 2, top + 16)
        # jump with the buffered second press
        rows = _group(logs, "buffer_jump")
        if rows:
            section(ws, r, 1, "Jump + a second press 0.1 s before landing (buffered: fires on touch-down, no ground sample between)")
            r += 1
            top = r
            header(ws, r, ["t (s)", "feet height (m)", "in the air"], height=22)
            h, a, t0 = r, r + 1, rows[0]["t"]
            for i, L in enumerate(rows):
                q = a + i
                put(ws, q, 1, round(L["t"] - t0, 4), "0.000", INPUT_FILL)
                put(ws, q, 2, L["y"], "0.00", INPUT_FILL)
                put(ws, q, 3, 1 if L["air"] else 0, "0", INPUT_FILL)
            b = a + len(rows) - 1
            put(ws, b + 1, 1, "air samples", font=BOLD, fill=TOTAL_FILL)
            put(ws, b + 1, 2, f"=SUM(C{a}:C{b})", "0")
            _scatter(ws, f"G{top}", "Buffered double jump: feet height", 1, [2], a, b, h, "m")
            r = max(b + 3, top + 16)
        # orbit: arm length
        rows = _group(logs, "orbit")
        if rows:
            section(ws, r, 1, "Third-person orbit beside the barrier: arm length (pull-in at once, critically damped ease-out)")
            r += 1
            top = r
            header(ws, r, ["t (s)", "arm (m)", "wanted (m)"], height=22)
            h, a, t0 = r, r + 1, rows[0]["t"]
            for i, L in enumerate(rows):
                q = a + i
                put(ws, q, 1, round(L["t"] - t0, 4), "0.000", INPUT_FILL)
                put(ws, q, 2, L["arm"], "0.00", INPUT_FILL)
                put(ws, q, 3, "=ArmLength", "0.00")
            b = a + len(rows) - 1
            put(ws, b + 1, 1, "shortest arm", font=BOLD, fill=TOTAL_FILL)
            put(ws, b + 1, 2, f"=MIN(B{a}:B{b})", "0.00")
            put(ws, b + 2, 1, "never longer than wanted", font=BOLD, fill=TOTAL_FILL)
            put(ws, b + 2, 2, f'=IF(MAX(B{a}:B{b})<=ArmLength+0.001,"YES","NO")')
            _scatter(ws, f"G{top}", "Spring arm during the orbit", 1, [2, 3], a, b, h, "m")
            r = max(b + 4, top + 16)

    r, ta, tb = tests_block(ws, r, ctx, 350, 359, "M3 tests (T350-T359)")
    section(ws, r, 1, "Captures: headless play_* before (M1) / after (M3), 1920x1080 EPIC, seed 1")
    r = _images(ws, r + 1, [("M1: first person", "docs/v9_notes/m3/shots_before/play_first_person.jpg"),
                            ("M3: first person", "docs/v9_notes/m3/shots_after/play_first_person.jpg"),
                            ("M1: third person", "docs/v9_notes/m3/shots_before/play_third_person.jpg"),
                            ("M3: third person", "docs/v9_notes/m3/shots_after/play_third_person.jpg"),
                            ("M3: sprint, +6 deg FOV kick (3.4 m/s)", "docs/v9_notes/m3/shots_after/play_sprint_fov.jpg"),
                            ("M3: crouch walk (0.8 m/s)", "docs/v9_notes/m3/shots_after/play_crouch.jpg"),
                            ("M3: jump", "docs/v9_notes/m3/shots_after/play_jump.jpg"),
                            ("M3: indoor arm pulled in", "docs/v9_notes/m3/shots_after/play_third_person_indoor.jpg")])
    section(ws, r, 1, "ThermalSense.exe session frames")
    r = _images(ws, r + 1, [("Sprint FOV kick", "docs/v9_notes/m3/exe_frames/m3_06_fp_sprint_fov.jpg"),
                            ("Crouch walk", "docs/v9_notes/m3/exe_frames/m3_07_fp_crouch_walk.jpg"),
                            ("Orbit near walls", "docs/v9_notes/m3/exe_frames/m3_10_tp_orbit_b.jpg"),
                            ("Third-person sprint jump", "docs/v9_notes/m3/exe_frames/m3_14_tp_sprint_jump.jpg")])
    widths(ws, [34, 60, 34, 14, 14, 4, 12, 12, 12, 12, 12, 12, 12, 12])
    finish(ws, 6, freeze="A4")
