"""Sheet 37 of the workbook: v9 M8 "lighting depth" -- CSM + PCSS, contact shadows, GTAO, aerial perspective, analytic
sky, AgX / ACES tone mapping.

Sources (written by the program, its tests or its source constants; nothing typed by hand except wording):
  src/math/ShadowMath.h, src/math/ToneMap.h, src/math/SkyModel.h   constants (parsed: the same numbers the GPU runs)
  docs/v9_notes/m8/night_ssim.txt                                  ffmpeg SSIM of the night shots (M8 vs the M1 renders)
  docs/v9_notes/m8/bench_shots_*.json                              --bench-shots: GPU ms per pass, before (M1) / after (M8)
  build/test_results.json                                          T380-T389 (via sheet 10)
  docs/v9_notes/m8/*.jpg                                           before / after captures
Blue = imported, white = Excel formula, amber = design-note text. Missing inputs degrade to "pending: <what>" rows.
"""
import json
import os
import re

from openpyxl.chart import LineChart, Reference, ScatterChart, Series
from openpyxl.formatting.rule import CellIsRule

from wb_common import (BOLD, INPUT_FILL, NOTE_FILL, PASS_FILL, PENDING_FILL, TOTAL_FILL, finish, header, path, put,
                       section, title, widths)
from wb_v6 import pending
import wb_v9

S_V9M8 = "37 v9 Lighting & Shadows"
INFO = {S_V9M8: "v9 M8: CSM + PCSS contact-hardening shadows, cascade blending, contact shadows, GTAO, aerial perspective, "
                "Preetham sky + SH ambient, AgX / ACES curves recomputed in Excel; GPU ms, night SSIM, tests T380-T389"}

TECHNIQUES = [
    ("Cascaded shadow maps", "2 cascades (near: frustum slice to 14 m, world: camp sphere), texel-snapped ortho fits, cross-fade over the last 20 % of the split",
     "CSM, stable cascades, cascade blending", "ShadowMath.cpp, ShaderSources.cpp moonShadowMap"),
    ("PCSS (sun)", "blocker search (16 raw-depth taps through a no-compare sampler object) -> average blocker depth -> penumbra w = gap x 2 tan(theta/2) -> Poisson PCF radius w/2; the moonlit night keeps the v8 PCF",
     "percentage-closer soft shadows (Fernando 2005), contact hardening", "ShaderSources.cpp cascadeShadow, ShadowMath.cpp"),
    ("Contact shadows", "per pixel a 0.6 m ray towards the sun through the pre-pass depth (16 steps); a hit hidden behind the surface by < thickness darkens the key light",
     "screen-space ray marching, depth-buffer occlusion", "ShadersLightingDepth.cpp kContactFrag"),
    ("GTAO (EPIC, day)", "3 slices x 8 steps both ways; horizon angles h0, h1; visibility = |n_proj| sum (cos n + 2h sin n - cos(2h - n)) / 4; rotated per frame, bilateral blur + TAA",
     "ground-truth ambient occlusion (Jimenez 2016) vs SSAO (hemisphere point sampling)", "ShadersLightingDepth.cpp kGtaoFrag"),
    ("Analytic sky", "Preetham / Perez F(theta, gamma) = (1 + A e^(B / cos theta))(1 + C e^(D gamma) + E cos^2 gamma), turbidity 2.5, limb-darkened sun disc, cumulus lit by the model",
     "physically based sky model", "SkyModel.cpp, ShadersPost.cpp daySky"),
    ("SH ambient", "the same sky + sunlit ground projected on 9 spherical harmonics, cosine-convolved (pi, 2pi/3, pi/4): dayAmbient(N) and the haze colours",
     "irradiance environment map (Ramamoorthi & Hanrahan 2001)", "SkyModel.cpp ambientSH, DayCycle.cpp"),
    ("Aerial perspective", "closed-form optical depth of an exponential height medium, in-scatter = sky-model radiance + HG(0.76) sun lobe, by day",
     "participating media, Beer-Lambert, Henyey-Greenstein", "ShadersPostFx.cpp aerialPerspective"),
    ("AgX / ACES", "AgX (inset, log2 over 16.5 stops, sigmoid, outset) by day, the v8 ACES fit at night; TS_TONEMAP=aces / RenderSettings::toneMap",
     "display rendering transforms, path to white", "ToneMap.cpp, ShadersPostFx.cpp composite"),
]


def _consts(rel, names):
    """Float constants parsed from a C++ header (name -> value)."""
    p = path(rel)
    out = {}
    if not os.path.exists(p):
        return out
    text = open(p, encoding="utf-8").read()
    for n in names:
        m = re.search(n + r"\s*=\s*(-?[\d.]+(?:e-?\d+)?)f?", text)
        if m:
            out[n] = float(m.group(1))
    return out


def _poly(rel):
    p = path(rel)
    if not os.path.exists(p):
        return []
    m = re.search(r"kAgxPoly\[7\]\s*=\s*\{([^}]*)\}", open(p, encoding="utf-8").read())
    return [float(x.strip().rstrip("f")) for x in m.group(1).split(",")] if m else []


def _load(rel):
    p = path(rel)
    if not os.path.exists(p):
        return None
    with open(p, encoding="utf-8") as f:
        return json.load(f)


def _ssim(rel):
    """Rows (shot, M8 vs v8, M8 vs M1, M1 vs v8) of night_ssim.txt."""
    p = path(rel)
    if not os.path.exists(p):
        return []
    rows = []
    for line in open(p, encoding="utf-8"):
        parts = line.split()
        if len(parts) >= 4 and not line.startswith("#"):
            try:
                rows.append((parts[0], float(parts[1]), float(parts[2]), float(parts[3])))
            except ValueError:
                pass
    return rows


def _param(ws, r, label, value, fmt, note):
    put(ws, r, 1, label, font=BOLD)
    put(ws, r, 2, value, fmt, INPUT_FILL)
    put(ws, r, 3, note, fill=NOTE_FILL)
    return f"$B${r}"


def build_v9_m8(wb, ctx):
    ws = wb[S_V9M8] if S_V9M8 in wb.sheetnames else wb.create_sheet(S_V9M8)
    title(ws, "37  v9 M8: Lighting Depth (CSM + PCSS, contact shadows, GTAO, aerial perspective, analytic sky, AgX)",
          "Objects sit on the ground and the day reads in depth: soft shadows widen with the blocker distance and harden at contact "
          "(PCSS), a screen-space ray march adds the contact line the shadow map cannot resolve, GTAO integrates the visible "
          "horizon arc, a Preetham sky drives both the sky that is drawn and the ambient / haze light, and AgX maps the HDR "
          "day to the display without clipping saturated highlights. Night keeps the v8 look (SSIM below).", span=14)
    r = 5
    section(ws, r, 1, "Techniques")
    r += 1
    header(ws, r, ["Technique", "What the renderer does", "Graphics concept", "Code"], height=20)
    for i, row in enumerate(TECHNIQUES):
        for k, v in enumerate(row):
            put(ws, r + 1 + i, 1 + k, v, fill=NOTE_FILL, font=BOLD if k == 0 else None)
        ws.row_dimensions[r + 1 + i].height = 42
    r += len(TECHNIQUES) + 2

    # ---- PCSS penumbra recomputed -------------------------------------------------------------------------------
    sh = _consts("src/math/ShadowMath.h", ["kSunSourceDeg", "kMoonSourceDeg", "kPcssMaxBlockerDistance", "kCascadeBlendBand"])
    section(ws, r, 1, "PCSS: penumbra w = (d_receiver - d_blocker) x light size / d_blocker; light at infinity: w = gap x 2 tan(theta / 2)")
    r += 1
    if len(sh) < 4:
        r = pending(ws, r, "constants in src/math/ShadowMath.h", span=6) + 1
    else:
        sun = _param(ws, r, "Sun effective disc (deg)", sh["kSunSourceDeg"], "0.00", "0.53 deg disc x 3: circumsolar aureole widens the source")
        moon = _param(ws, r + 1, "Moon disc (deg)", sh["kMoonSourceDeg"], "0.00", "the moon's disc")
        search = _param(ws, r + 2, "Max blocker distance (m)", sh["kPcssMaxBlockerDistance"], "0.0", "the blocker search covers blockers up to this far above")
        put(ws, r + 3, 1, "Blocker search radius (m)", font=BOLD)
        put(ws, r + 3, 2, f"=0.5*{search}*2*TAN(RADIANS({sun})/2)", "0.000")
        texel = _param(ws, r + 4, "Near-cascade texel (m, 4096 map)", 0.005, "0.0000", "typical: 20 m slice sphere / 4096 texels (setting)")
        r += 6
        header(ws, r, ["Gap receiver - blocker (m)", "Sun penumbra (m)", "Moon penumbra (m)", "PCF radius (texels)"], height=30)
        a = r + 1
        gaps = [0, 0.05, 0.1, 0.25, 0.5, 1, 2, 3, 5, 7.5, 10]
        for i, g in enumerate(gaps):
            q = a + i
            put(ws, q, 1, g, "0.00")
            put(ws, q, 2, f"=A{q}*2*TAN(RADIANS({sun})/2)", "0.0000")
            put(ws, q, 3, f"=A{q}*2*TAN(RADIANS({moon})/2)", "0.0000")
            put(ws, q, 4, f"=MIN(MAX(0.5*B{q}/{texel},1.6),40)", "0.0")
        b = a + len(gaps) - 1
        ch = ScatterChart()
        ch.title = "Penumbra width grows with the blocker distance"
        ch.x_axis.title = "gap (m)"
        ch.y_axis.title = "penumbra (m)"
        xs = Reference(ws, min_col=1, min_row=a, max_row=b)
        for c, name in ((2, "sun"), (3, "moon")):
            s = Series(Reference(ws, min_col=c, min_row=a, max_row=b), xs, title=name)
            s.marker.symbol = "circle"
            ch.series.append(s)
        ch.height, ch.width = 7, 13
        ws.add_chart(ch, f"F{a - 7}")
        r = b + 2
        # Cascade blend weight
        section(ws, r, 1, "Cascade blending: weight of the world cascade = smoothstep(split (1 - band), split, distance)")
        r += 1
        band = _param(ws, r, "Blend band (fraction of the split)", sh["kCascadeBlendBand"], "0.00", "ShadowMath.h kCascadeBlendBand")
        split = _param(ws, r + 1, "Split distance EPIC (m)", 14.0, "0.0", "QualityProfile::cascadeSplit (setting)")
        r += 3
        header(ws, r, ["View distance (m)", "t", "Weight (world cascade)"], height=20)
        a = r + 1
        for i in range(13):
            q = a + i
            put(ws, q, 1, round(10.0 + 0.4 * i, 2), "0.0")
            put(ws, q, 2, f"=MIN(MAX((A{q}-{split}*(1-{band}))/({split}-{split}*(1-{band})),0),1)", "0.000")
            put(ws, q, 3, f"=B{q}*B{q}*(3-2*B{q})", "0.000")
        r = a + 14

    # ---- Tone curves recomputed ---------------------------------------------------------------------------------
    tm = _consts("src/math/ToneMap.h", ["kAgxMinEv", "kAgxMaxEv", "kAcesExposure", "kAgxExposure", "kAgxDayContrast", "kAcesDayContrast"])
    poly = _poly("src/math/ToneMap.h")
    section(ws, r, 1, "Tone curves (grey axis): AgX = sRGB(((sigmoid((log2 x - minEv) / (maxEv - minEv)))^2.2), ACES = sRGB(Hill fit of RRT + ODT)")
    r += 1
    if len(tm) < 6 or len(poly) != 7:
        r = pending(ws, r, "constants in src/math/ToneMap.h", span=6) + 1
    else:
        mn = _param(ws, r, "AgX min EV", tm["kAgxMinEv"], "0.0000", "log2 range of the encoding (16.5 stops)")
        mx = _param(ws, r + 1, "AgX max EV", tm["kAgxMaxEv"], "0.0000", "")
        ea = _param(ws, r + 2, "ACES pre-exposure", tm["kAcesExposure"], "0.00", "v3-v8 value")
        eg = _param(ws, r + 3, "AgX pre-exposure", tm["kAgxExposure"], "0.00", "mid grey maps like the ACES setting (T382)")
        put(ws, r + 4, 1, "AgX sigmoid coefficients x^6..x^0", font=BOLD)
        for k, c in enumerate(poly):
            put(ws, r + 4, 2 + k, c, "0.0000", INPUT_FILL)
        prow = r + 4
        P = [f"${chr(ord('B') + k)}${prow}" for k in range(7)]
        r += 6
        header(ws, r, ["Exposure (EV from 0.18)", "Scene linear", "AgX log-encoded", "AgX sigmoid", "AgX display (sRGB)",
                       "ACES RRT+ODT", "ACES display (sRGB)"], height=30)
        a = r + 1
        evs = [x * 0.5 for x in range(-16, 13)]
        for i, ev in enumerate(evs):
            q = a + i
            put(ws, q, 1, ev, "0.0")
            put(ws, q, 2, f"=0.18*2^A{q}", "0.00000")
            put(ws, q, 3, f"=(MIN(MAX(LOG(B{q}*{eg},2),{mn}),{mx})-{mn})/({mx}-{mn})", "0.0000")
            x = f"C{q}"
            put(ws, q, 4, f"={P[0]}*{x}^6+{P[1]}*{x}^5+{P[2]}*{x}^4+{P[3]}*{x}^3+{P[4]}*{x}^2+{P[5]}*{x}+{P[6]}", "0.0000")
            lin = f"MIN(MAX(D{q},0),1)^2.2"
            put(ws, q, 5, f"=IF({lin}<=0.0031308,12.92*{lin},1.055*{lin}^(1/2.4)-0.055)", "0.000")
            v = f"(B{q}*{ea})"
            put(ws, q, 6, f"=MIN(MAX(({v}*({v}+0.0245786)-0.000090537)/({v}*(0.983729*{v}+0.432951)+0.238081),0),1)", "0.0000")
            put(ws, q, 7, f"=IF(F{q}<=0.0031308,12.92*F{q},1.055*F{q}^(1/2.4)-0.055)", "0.000")
        b = a + len(evs) - 1
        ch = LineChart()
        ch.title = "AgX vs ACES display response (grey)"
        ch.y_axis.title = "display (sRGB)"
        ch.x_axis.title = "EV from mid grey"
        ch.add_data(Reference(ws, min_col=5, min_row=a - 1, max_row=b), titles_from_data=True)
        ch.add_data(Reference(ws, min_col=7, min_row=a - 1, max_row=b), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=a, max_row=b))
        ch.height, ch.width = 8, 15
        ws.add_chart(ch, f"I{a}")
        r = b + 2
        put(ws, r, 1, "Toe (EV -4): AgX - ACES", font=BOLD, fill=TOTAL_FILL)
        put(ws, r, 2, f"=INDEX(E{a}:E{b},MATCH(-4,A{a}:A{b},0))-INDEX(G{a}:G{b},MATCH(-4,A{a}:A{b},0))", "0.000", font=BOLD)
        put(ws, r, 3, "AgX lifts the darks: why the night keeps the ACES toe (night grade authored for it)", fill=NOTE_FILL)
        r += 2

    # ---- Sky model: zenith luminance vs sun elevation --------------------------------------------------------------
    sk = _consts("src/math/SkyModel.h", ["kSkyTurbidity", "kSkyRadianceScale"])
    section(ws, r, 1, "Preetham zenith luminance Y_z = (4.0453 T - 4.9710) tan chi - 0.2155 T + 2.4192, chi = (4/9 - T/120)(pi - 2 theta_s)")
    r += 1
    if len(sk) < 2:
        r = pending(ws, r, "constants in src/math/SkyModel.h", span=6) + 1
    else:
        T = _param(ws, r, "Turbidity T", sk["kSkyTurbidity"], "0.0", "SkyModel.h (2 very clear .. 6 hazy)")
        sc = _param(ws, r + 1, "kcd/m^2 -> HDR radiance", sk["kSkyRadianceScale"], "0.000", "calibrated to the v8 noon ambient (T388)")
        r += 3
        header(ws, r, ["Sun elevation (deg)", "theta_s (rad)", "chi", "Y_z (kcd/m^2)", "Zenith luminance (HDR)"], height=22)
        a = r + 1
        for i, e in enumerate([1, 5, 10, 20, 30, 45, 60, 75, 90]):
            q = a + i
            put(ws, q, 1, e, "0")
            put(ws, q, 2, f"=RADIANS(90-A{q})", "0.000")
            put(ws, q, 3, f"=(4/9-{T}/120)*(PI()-2*B{q})", "0.000")
            put(ws, q, 4, f"=MAX((4.0453*{T}-4.971)*TAN(C{q})-0.2155*{T}+2.4192,0.05)", "0.00")
            put(ws, q, 5, f"=D{q}*{sc}", "0.000")
        b = a + 8
        ch = LineChart()
        ch.title = "Zenith brightness vs sun elevation"
        ch.add_data(Reference(ws, min_col=4, min_row=a - 1, max_row=b), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=a, max_row=b))
        ch.height, ch.width = 7, 13
        ws.add_chart(ch, f"G{a - 2}")
        r = b + 3

    # ---- GPU ms ---------------------------------------------------------------------------------------------------
    section(ws, r, 1, "GPU ms per pass (GpuTimer, --bench-shots, 1920x1080 EPIC; GPU shared with other work: repeated runs, mean)")
    r += 1
    runs = [("WSL d3d12, M1 build", "docs/v9_notes/m8/bench_shots_m1_wsl.json"), ("WSL d3d12, M8 build", "docs/v9_notes/m8/bench_shots_m8_wsl.json"),
            ("exe NVIDIA, M1", "docs/v9_notes/m8/bench_shots_m1_win.json"), ("exe NVIDIA, M8", "docs/v9_notes/m8/bench_shots_m8_win.json")]
    data = [(lab, _load(rel)) for lab, rel in runs]
    passes = ["shadow maps", "pre-pass", "ssao", "gtao", "contact shadows", "volumetric fog", "scene", "combine + ssr", "composite"]
    header(ws, r, ["Run", "Shot", "Time"] + passes + ["GPU sum", "frame ms"], height=36)
    q = r + 1
    any_rows = False
    for lab, d in data:
        if not d:
            put(ws, q, 1, lab, font=BOLD)
            put(ws, q, 2, "pending: " + dict(runs)[lab], fill=PENDING_FILL)
            q += 1
            continue
        for x in d["shots"]:
            any_rows = True
            put(ws, q, 1, lab, font=BOLD)
            put(ws, q, 2, x["shot"])
            put(ws, q, 3, x["time"])
            for k, pname in enumerate(passes):
                put(ws, q, 4 + k, x.get("passes", {}).get(pname), "0.000", INPUT_FILL)
            put(ws, q, 4 + len(passes), x.get("gpu_sum_ms"), "0.00", INPUT_FILL)
            put(ws, q, 5 + len(passes), x.get("frame_ms"), "0.00", INPUT_FILL)
            q += 1
    first, last = r + 1, q - 1
    r = q + 1
    if any_rows:
        put(ws, r, 1, "Mean M8 contact shadows (ms, day)", font=BOLD, fill=TOTAL_FILL)
        put(ws, r, 2, f'=IFERROR(AVERAGEIFS(H{first}:H{last},A{first}:A{last},"*M8*",C{first}:C{last},"day"),"-")', "0.000", font=BOLD)
        put(ws, r + 1, 1, "Mean M8 GTAO (ms)", font=BOLD, fill=TOTAL_FILL)
        put(ws, r + 1, 2, f'=IFERROR(AVERAGEIFS(G{first}:G{last},A{first}:A{last},"*M8*"),"-")', "0.000", font=BOLD)
        put(ws, r + 2, 1, "Mean SSAO it replaced (M1, ms)", font=BOLD, fill=TOTAL_FILL)
        put(ws, r + 2, 2, f'=IFERROR(AVERAGEIFS(F{first}:F{last},A{first}:A{last},"*M1*"),"-")', "0.000", font=BOLD)
        put(ws, r + 3, 1, "Scene pass M8 - M1 (PCSS, sky SH; ms)", font=BOLD, fill=TOTAL_FILL)
        put(ws, r + 3, 2, f'=IFERROR(AVERAGEIFS(J{first}:J{last},A{first}:A{last},"*M8*")-AVERAGEIFS(J{first}:J{last},A{first}:A{last},"*M1*"),"-")', "0.000", font=BOLD)
        r += 5

    # ---- Night SSIM -----------------------------------------------------------------------------------------------
    section(ws, r, 1, "Night parity: ffmpeg SSIM (All), same seed / shot / size; v8 = 138163d, M1 = 96782ee (the base of M8)")
    r += 1
    rows = _ssim("docs/v9_notes/m8/night_ssim.txt")
    header(ws, r, ["Shot", "M8 vs v8", "M8 vs M1", "M1 vs v8", "Threshold", "Status vs v8", "M8's own change"], height=30)
    a = r + 1
    if not rows:
        r = pending(ws, a, "docs/v9_notes/m8/night_ssim.txt", span=6) + 1
    else:
        for i, (shot, v8, m1, m1v8) in enumerate(rows):
            q = a + i
            put(ws, q, 1, shot, font=BOLD)
            put(ws, q, 2, v8, "0.0000", INPUT_FILL)
            put(ws, q, 3, m1, "0.0000", INPUT_FILL)
            put(ws, q, 4, m1v8, "0.0000", INPUT_FILL)
            put(ws, q, 5, 0.97, "0.00")
            # below 0.97 vs v8 while M1 already differs as much: the drop predates M8 (M1 grass strip / tick timing)
            put(ws, q, 6, f'=IF(B{q}>=E{q},"PASS",IF(D{q}<E{q},"FAIL (M1 already "&TEXT(D{q},"0.000")&")","FAIL"))', font=BOLD)
            put(ws, q, 7, f'=IF(C{q}>=E{q},"PASS","FAIL")', font=BOLD)
        b = a + len(rows) - 1
        for rng in (f"F{a}:F{b}", f"G{a}:G{b}"):
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'], fill=PASS_FILL))
            ws.conditional_formatting.add(rng, CellIsRule(operator="notEqual", formula=['"PASS"'], fill=PENDING_FILL))
        put(ws, b + 1, 1, "Minimum", font=BOLD, fill=TOTAL_FILL)
        for c in "BCD":
            put(ws, b + 1, " BCD".index(c) + 1, f"=MIN({c}{a}:{c}{b})", "0.0000", font=BOLD)
        r = b + 3

    r, _, _ = wb_v9.tests_block(ws, r, ctx, 380, 389, "M8 tests (T380-T389)")
    section(ws, r, 1, "Captures (headless EPIC, RTX 4070): before = M1, after = M8")
    r = wb_v9._images(ws, r + 1, [
        ("Before: day overview", "docs/v9_notes/m8/before_day_overview.jpg"), ("After: day overview", "docs/v9_notes/m8/after_day_overview.jpg"),
        ("Before: soldier close-up", "docs/v9_notes/m8/before_soldier_closeup.jpg"), ("After: soldier close-up (PCSS + contact)", "docs/v9_notes/m8/after_soldier_closeup.jpg"),
        ("Feet: PCSS + contact off / on (same build)", "docs/v9_notes/m8/feet_contact_off_on.jpg"), ("Mechanic: off / on", "docs/v9_notes/m8/mechanic_off_on.jpg"),
        ("Before: noon sky", "docs/v9_notes/m8/before_sky_noon.jpg"), ("After: noon sky (Preetham + clouds)", "docs/v9_notes/m8/after_sky_noon.jpg"),
        ("Before: sunset", "docs/v9_notes/m8/before_sky_sunset.jpg"), ("After: sunset (AgX, aerial perspective)", "docs/v9_notes/m8/after_sky_sunset.jpg"),
        ("Night forest M1", "docs/v9_notes/m8/before_walk_forest.jpg"), ("Night forest M8", "docs/v9_notes/m8/after_walk_forest.jpg"),
    ])
    widths(ws, [34, 40, 22, 18, 14, 14, 14, 12, 12, 12, 12, 12, 10, 10])
    finish(ws, 6, freeze="A4")
