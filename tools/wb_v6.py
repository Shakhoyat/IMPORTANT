"""Sheets 21-23 of the workbook (v6): Day Cycle & Lighting, Soldiers & Routes, CV Noise & Seeds.

Sources: scene_export.json key "v6" (DocExport, WP-E: day_cycle / garrison / sensor / seeds), the soldier components
of the ordinary export (sheet 3 picks them up automatically), build/figdata_out/cv_noise.json (WP-D noise sweep) and the
test results (sheet 10). Every block degrades to a clear "pending: <what>" row when its data is not exported yet, so the
workbook builds on old exports. Numbers are imported (blue) or Excel formulas (white); only wording is authored here.
"""
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.formatting.rule import CellIsRule

from wb_common import (S_CVN, S_DAY, S_SOLD, BOLD, FAIL_FILL, FMT_INT, INPUT_FILL, NOTE_FILL, PASS_FILL, PENDING_FILL, S_TEST, TOTAL_FILL, col,
                       define_name, finish, header, load_json, put, section, style_chart, title, widths)

V6_SHEETS = [S_DAY, S_SOLD, S_CVN]


def pending(ws, r, what, c=1, span=8):
    """One clearly marked placeholder row."""
    put(ws, r, c, "pending: " + what, fill=PENDING_FILL, font=BOLD)
    if span > 1:
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
    return r + 1


def _kv(ws, r, rows, fmt="0.###"):
    """(label, value, unit/note) rows; returns (next row, {label: absolute cell})."""
    cells = {}
    for lab, val, note in rows:
        put(ws, r, 1, lab, font=BOLD)
        put(ws, r, 2, val, fmt if isinstance(val, (int, float)) and not isinstance(val, bool) else None, INPUT_FILL)
        put(ws, r, 3, note)
        cells[lab] = f"$B${r}"
        r += 1
    return r, cells


def _rgb(v, i):
    return v[i] if isinstance(v, (list, tuple)) and len(v) > i else None


# ---- 21 Day Cycle & Lighting ------------------------------------------------------------------------------
GRID_COLS = [  # (header, key, number format)
    ("Hour", "hour", "0.00"), ("Clock", "clock", None), ("Phase", "phase", None),
    ("Sun elevation (deg, exported)", "elevation_deg", "0.000"), ("Sun azimuth (deg)", "azimuth_deg", "0.0"),
    ("Daylight 0..1", "daylight", "0.000"), ("Key intensity", "key_intensity", "0.000"),
    ("Ambient luma", "ambient_luma", "0.0000"), ("Fog density", "fog_density", "0.0000"),
    ("Volumetric fog", "vol_fog_density", "0.0000"), ("Exposure key (log2)", "exposure_key_log2", "0.000"),
    ("Purkinje", "purkinje", "0.000"), ("Solar loading", "solar", "0.000"), ("Practicals scale", "practical_scale", "0.000"),
    ("Beams on", "beams_on", None), ("Waves allowed", "waves_allowed", None),
    ("Cloud shadow mean", "cloud_shadow_mean", "0.000"),
]


def build_daycycle(wb, ctx):
    ws = wb[S_DAY]
    refs = ctx["refs"]
    v6 = ctx["scene"].get("v6") or {}
    dc = v6.get("day_cycle") or {}
    title(ws, "21  Day Cycle & Lighting",
          "v6 time of day: the exported day-cycle grid (0..24 h every 0.25 h) with the sun, key light, ambient, fog, "
          "exposure, solar loading, searchlight beams, wave gating and the mean cloud-shadow factor. The white columns "
          "recompute the sun elevation in Excel from latitude, declination and hour: "
          "sin e = sin(lat) sin(dec) + cos(lat) cos(dec) cos(15 (h - 12)), and check it against the exported value "
          "(tolerance below). Blue = exported by the program (DocExport v6.day_cycle).", span=12)
    cfg = dc.get("config") or {}
    section(ws, 4, 1, "Configuration (DayCycleConfig)")
    if not cfg:
        r = pending(ws, 5, "build/scene_export.json has no v6.day_cycle.config (WP-E exporter) - re-run `make docs`", span=6)
        cells = {}
    else:
        header(ws, 5, ["Constant", "Value", "Note"], height=20)
        notes = {"latitude_deg": "site latitude (Khulna)", "declination_deg": "solar declination, early October",
                 "day_sec_per_hour": "real seconds per game hour by day", "night_sec_per_hour": "real seconds per game hour by night",
                 "twilight_slow": "|elevation| < 6 deg plays slower by this factor",
                 "scrub_hours_per_sec": "manual scrub speed", "shadow_quant_deg": "shadow-cache direction quantum",
                 "cycle_seconds": "length of one auto day + night cycle"}
        r, cells = _kv(ws, 6, [(k, v, notes.get(k, "")) for k, v in cfg.items()])
    r += 1
    # tolerance for the formula check (an authored setting, labelled)
    put(ws, r, 1, "Elevation check tolerance (deg)", font=BOLD)
    put(ws, r, 2, 0.05, "0.000", NOTE_FILL)
    put(ws, r, 3, "authored setting; the exported value is a float32 evaluation of the same formula")
    tol = f"$B${r}"
    r += 2
    section(ws, r, 1, "Phase events of one manual sweep 0..24 h")
    r += 1
    ev = dc.get("events") or []
    if not ev:
        r = pending(ws, r, "v6.day_cycle.events (Dawn / DayBreak / Dusk / NightFalls) not exported yet", span=6)
    else:
        header(ws, r, ["Hour", "Event"], height=20)
        r += 1
        for e in ev:
            put(ws, r, 1, e.get("hour"), "0.00", INPUT_FILL)
            put(ws, r, 2, e.get("event"), fill=INPUT_FILL)
            r += 1
    r += 1
    section(ws, r, 1, "Day-cycle grid with the Excel recomputation of the sun elevation")
    r += 1
    grid = dc.get("grid") or []
    if not grid:
        pending(ws, r, "v6.day_cycle.grid (97 rows) not exported yet - the sun-elevation formulas, check column and chart "
                       "appear when WP-E's DocExport lands", span=10)
        ctx["refs"]["day_rows"] = None
        widths(ws, [26, 14, 16] + [12] * 17)
        finish(ws, 5)
        return
    hdr = [g[0] for g in GRID_COLS] + ["Sun elevation (deg, Excel formula)", "Excel - exported (deg)", "Check"]
    h = r
    header(ws, h, hdr, height=48)
    first = h + 1
    lat, dec = cells.get("latitude_deg"), cells.get("declination_deg")
    cE, cX, cC = len(GRID_COLS) + 1, len(GRID_COLS) + 2, len(GRID_COLS) + 3
    for i, g in enumerate(grid):
        rr = first + i
        for k, (_, key, fmt) in enumerate(GRID_COLS):
            v = g.get(key)
            if isinstance(v, bool):
                v = "yes" if v else "no"
            put(ws, rr, 1 + k, v, fmt, INPUT_FILL)
        if lat and dec:
            f = (f"=DEGREES(ASIN(SIN(RADIANS({lat}))*SIN(RADIANS({dec}))+COS(RADIANS({lat}))*COS(RADIANS({dec}))"
                 f"*COS(RADIANS(15*(A{rr}-12)))))")
            put(ws, rr, cE, f, "0.000")
            put(ws, rr, cX, f"={col(cE)}{rr}-D{rr}", "0.0000")
            put(ws, rr, cC, f'=IF(ABS({col(cX)}{rr})<={tol},"OK","CHECK")')
        else:
            put(ws, rr, cE, "pending: latitude / declination")
    last = first + len(grid) - 1
    ws.conditional_formatting.add(f"{col(cC)}{first}:{col(cC)}{last}", CellIsRule(operator="equal", formula=['"OK"'], fill=PASS_FILL))
    ws.conditional_formatting.add(f"{col(cC)}{first}:{col(cC)}{last}", CellIsRule(operator="equal", formula=['"CHECK"'], fill=FAIL_FILL))
    # KPI strip
    put(ws, 3, 1, "Grid rows", font=BOLD, fill=TOTAL_FILL)
    put(ws, 3, 2, f"=COUNT(A{first}:A{last})", FMT_INT)
    put(ws, 3, 3, "Max |Excel - exported| (deg)", font=BOLD, fill=TOTAL_FILL)
    put(ws, 3, 4, f"=MAX(MAX({col(cX)}{first}:{col(cX)}{last}),-MIN({col(cX)}{first}:{col(cX)}{last}))" if lat else "pending", "0.0000")
    put(ws, 3, 5, "Rows OK", font=BOLD, fill=TOTAL_FILL)
    put(ws, 3, 6, f'=COUNTIF({col(cC)}{first}:{col(cC)}{last},"OK")', FMT_INT)
    put(ws, 3, 7, "Peak elevation (deg)", font=BOLD, fill=TOTAL_FILL)
    put(ws, 3, 8, f"=MAX(D{first}:D{last})", "0.00")
    put(ws, 3, 9, "Peak key intensity", font=BOLD, fill=TOTAL_FILL)
    put(ws, 3, 10, f"=MAX(G{first}:G{last})", "0.000")
    put(ws, 3, 11, "Hours with beams on", font=BOLD, fill=TOTAL_FILL)
    put(ws, 3, 12, f'=COUNTIF(O{first}:O{last},"yes")*0.25', "0.00")
    ctx["refs"]["day_rows"] = f"'{S_DAY}'!$F$3"
    define_name(wb, "DayCycleRowsOK", S_DAY, "F3")
    # chart: elevation (left axis) and key intensity (right axis) against hour
    c1 = ScatterChart()
    c1.title, c1.style = "Sun elevation and key intensity vs hour", 13
    c1.x_axis.title, c1.y_axis.title = "hour", "elevation (deg)"
    xs = Reference(ws, min_col=1, min_row=first, max_row=last)
    s1 = Series(Reference(ws, min_col=4, min_row=h, max_row=last), xs, title_from_data=True)
    s1.marker.symbol, s1.smooth = "none", False
    c1.series.append(s1)
    c2 = ScatterChart()
    s2 = Series(Reference(ws, min_col=7, min_row=h, max_row=last), xs, title_from_data=True)
    s2.marker.symbol, s2.smooth = "none", False
    c2.series.append(s2)
    c2.y_axis.axId = 200
    c2.y_axis.title = "key intensity"
    c2.y_axis.crosses = "max"
    c2.y_axis.delete = False
    style_chart(c1)
    c1 += c2
    c1.height, c1.width = 9, 20
    ws.add_chart(c1, f"{col(cC + 2)}{h}")
    widths(ws, [26, 14, 16] + [12] * 14 + [16, 14, 10])
    finish(ws, h, freeze=f"B{first}")


# ---- 22 Soldiers & Routes ---------------------------------------------------------------------------------
PRIMS = ["Cube", "Sphere", "Cylinder", "Cylinder (frustum)", "Cone"]


def build_soldiers(wb, ctx):
    ws = wb[S_SOLD]
    g = (ctx["scene"].get("v6") or {}).get("garrison") or {}
    title(ws, "22  Soldiers & Routes",
          "v6 garrison: the 8 friendly soldiers (possessable personnel, never targets, no damage model): roster with "
          "day / night duty, primitive counts (limbs are capsules = a cylinder + joint spheres), the joint tree, the "
          "gait constants and cycle, and the length of each Catmull-Rom route. Blue = exported (DocExport v6.garrison); "
          "totals, counts and the route statistics are Excel formulas. Sheet 3 lists every soldier part.", span=12)
    soldiers = g.get("soldiers") or []
    r = 4
    section(ws, r, 1, "Roster")
    r += 1
    if not soldiers:
        r = pending(ws, r, "v6.garrison.soldiers not exported yet (re-run `make docs` after WP-E's DocExport lands); "
                           "the soldier parts will also appear on sheet 3 automatically", span=10)
        first = last = None
    else:
        hdr = ["#", "Name", "Controllable #", "Day duty", "Night duty", "x (m)", "y (m)", "z (m)", "Parts (leaves)", "Groups"] \
              + PRIMS + ["Friendly radius (m)", "Status text"]
        header(ws, r, hdr, height=40)
        first = r + 1
        for i, s in enumerate(soldiers):
            rr = first + i
            pos = s.get("position") or []
            vals = [s.get("index"), s.get("name"), s.get("controllable_index"), s.get("roster_day"), s.get("roster_night"),
                    _rgb(pos, 0), _rgb(pos, 1), _rgb(pos, 2), s.get("parts"), s.get("groups")]
            for k, v in enumerate(vals):
                put(ws, rr, 1 + k, v, "0.00" if 5 <= k <= 7 else None, INPUT_FILL, BOLD if k == 1 else None)
            pr = s.get("primitives") or {}
            for k, p in enumerate(PRIMS):
                put(ws, rr, 11 + k, pr.get(p, 0), FMT_INT, INPUT_FILL)
            put(ws, rr, 16, s.get("friendly_radius"), "0.00", INPUT_FILL)
            put(ws, rr, 17, s.get("status"), fill=INPUT_FILL)
        last = first + len(soldiers) - 1
        t = last + 1
        put(ws, t, 2, "Total", fill=TOTAL_FILL, font=BOLD)
        for c in [9, 10] + list(range(11, 16)):
            put(ws, t, c, f"=SUM({col(c)}{first}:{col(c)}{last})", FMT_INT, TOTAL_FILL, BOLD)
        put(ws, t + 1, 2, "Primitives counted from the list above", font=BOLD)
        put(ws, t + 1, 9, f"=SUM(K{t}:O{t})", FMT_INT)
        put(ws, t + 1, 10, f'=IF(I{t + 1}=I{t},"matches parts","differs")')
        put(ws, t + 2, 2, "Capsule limb rows (cylinder rows)", font=BOLD)
        put(ws, t + 2, 9, f"=M{t}", FMT_INT)
        put(ws, t + 2, 10, "each limb segment = one cylinder (+ spheres at its joints)")
        put(ws, t + 3, 2, "Soldier leaf rows on sheet 3 (cross-check)", font=BOLD)
        put(ws, t + 3, 9, f"=COUNTIFS('3 Object Construction'!$B$5:$B$20000,\"*oldier*\",'3 Object Construction'!$C$5:$C$20000,\"<>Group\")", FMT_INT)
        put(ws, t + 3, 10, f'=IF(I{t + 3}=I{t},"sheet 3 agrees","sheet 3 differs (object names)")')
        put(ws, 3, 1, "Soldiers", font=BOLD, fill=TOTAL_FILL)
        put(ws, 3, 2, f"=COUNTA(B{first}:B{last})", "0")
        put(ws, 3, 3, "On duty by day", font=BOLD, fill=TOTAL_FILL)
        put(ws, 3, 4, f'=COUNTA(D{first}:D{last})-COUNTIF(D{first}:D{last},"off*")-COUNTIF(D{first}:D{last},"-")', "0")
        put(ws, 3, 5, "On duty at night", font=BOLD, fill=TOTAL_FILL)
        put(ws, 3, 6, f'=COUNTA(E{first}:E{last})-COUNTIF(E{first}:E{last},"off*")-COUNTIF(E{first}:E{last},"-")', "0")
        ctx["refs"]["soldiers"] = f"'{S_SOLD}'!$B$3"
        define_name(wb, "SoldierCount", S_SOLD, "B3")
        r = t + 5
    # routes
    section(ws, r, 1, "Route lengths (Catmull-Rom splines, alpha = 0.5)")
    r += 1
    rl = g.get("route_lengths_m") or []
    if not rl:
        r = pending(ws, r, "v6.garrison.route_lengths_m not exported yet", span=6)
    else:
        header(ws, r, ["Route", "Length (m)"], height=20)
        a = r + 1
        for i, v in enumerate(rl):
            put(ws, a + i, 1, i + 1, "0")
            put(ws, a + i, 2, v, "0.0", INPUT_FILL)
        b = a + len(rl) - 1
        for k, (lab, fn) in enumerate((("Total", "SUM"), ("Shortest", "MIN"), ("Longest", "MAX"), ("Mean", "AVERAGE"))):
            put(ws, b + 1 + k, 1, lab, font=BOLD, fill=TOTAL_FILL)
            put(ws, b + 1 + k, 2, f"={fn}(B{a}:B{b})", "0.0", TOTAL_FILL)
        r = b + 6
    # joints
    section(ws, r, 1, "Joint tree of soldier 1 (node hierarchy)")
    r += 1
    joints = g.get("joints") or []
    if not joints:
        r = pending(ws, r, "v6.garrison.joints not exported yet", span=6)
    else:
        header(ws, r, ["Id", "Name", "Parent", "Depth", "Primitive"], height=20)
        a = r + 1
        for i, j in enumerate(joints):
            put(ws, a + i, 1, j.get("id"), fill=INPUT_FILL)
            put(ws, a + i, 2, "  " * int(j.get("depth", 0) or 0) + str(j.get("name")), fill=INPUT_FILL)
            put(ws, a + i, 3, j.get("parent"), fill=INPUT_FILL)
            put(ws, a + i, 4, j.get("depth"), "0", INPUT_FILL)
            put(ws, a + i, 5, j.get("primitive"), fill=INPUT_FILL)
        b = a + len(joints) - 1
        put(ws, b + 1, 1, "Nodes", font=BOLD, fill=TOTAL_FILL)
        put(ws, b + 1, 4, f"=COUNTA(A{a}:A{b})", "0", TOTAL_FILL)
        put(ws, b + 2, 1, "Max depth", font=BOLD, fill=TOTAL_FILL)
        put(ws, b + 2, 4, f"=MAX(D{a}:D{b})", "0", TOTAL_FILL)
        put(ws, b + 3, 1, "Group nodes", font=BOLD, fill=TOTAL_FILL)
        put(ws, b + 3, 4, f'=COUNTIF(E{a}:E{b},"Group")', "0", TOTAL_FILL)
        r = b + 5
    # gait
    section(ws, r, 1, "Gait constants (walk / run)")
    r += 1
    gait = g.get("gait") or {}
    walk, run = gait.get("walk") or {}, gait.get("run") or {}
    if not walk and not run:
        r = pending(ws, r, "v6.garrison.gait not exported yet", span=6)
    else:
        header(ws, r, ["Parameter", "Walk", "Run"], height=20)
        r += 1
        for k in list(dict.fromkeys(list(walk) + list(run))):
            put(ws, r, 1, k, font=BOLD)
            put(ws, r, 2, walk.get(k), "0.###", INPUT_FILL)
            put(ws, r, 3, run.get(k), "0.###", INPUT_FILL)
            r += 1
    r += 1
    section(ws, r, 1, "Walk cycle (forward kinematics output, phase 0..1)")
    r += 1
    cyc = gait.get("cycle") or []
    if not cyc:
        pending(ws, r, "v6.garrison.gait.cycle not exported yet", span=6)
    else:
        keys = list(cyc[0].keys())
        header(ws, r, keys, height=34)
        a = r + 1
        for i, row in enumerate(cyc):
            for k, key in enumerate(keys):
                put(ws, a + i, 1 + k, row.get(key), "0.000", INPUT_FILL)
        b = a + len(cyc) - 1
        ch = ScatterChart()
        ch.title, ch.style = "Walk cycle joint angles vs phase", 13
        ch.x_axis.title, ch.y_axis.title = "phase", "deg / m"
        xs = Reference(ws, min_col=1, min_row=a, max_row=b)
        for k, key in enumerate(keys):
            if key in ("hip_l", "knee_l", "hip_r", "shoulder_l"):
                s = Series(Reference(ws, min_col=1 + k, min_row=r, max_row=b), xs, title_from_data=True)
                s.marker.symbol, s.smooth = "none", True
                ch.series.append(s)
        ch.height, ch.width = 8.5, 17
        ws.add_chart(style_chart(ch), f"{col(len(keys) + 2)}{r}")
    widths(ws, [24, 22, 12, 14, 14, 9, 9, 9, 10, 9, 8, 8, 9, 10, 8, 11, 34])
    finish(ws, 5, freeze="A4")


# ---- 23 CV Noise & Seeds ----------------------------------------------------------------------------------
def _flatten_json(data):
    """Generic view of figdata JSON: (scalars {name: value}, tables {name: [row dicts]})."""
    scal, tabs = {}, {}
    if isinstance(data, list):
        if data and isinstance(data[0], dict):
            tabs["rows"] = data
        return scal, tabs
    for k, v in (data or {}).items():
        if isinstance(v, list) and v and isinstance(v[0], dict):
            tabs[k] = v
        elif isinstance(v, dict):
            for k2, v2 in v.items():
                if isinstance(v2, (int, float, str)):
                    scal[f"{k}.{k2}"] = v2
        elif isinstance(v, (int, float, str)):
            scal[k] = v
        elif isinstance(v, list) and v and isinstance(v[0], (int, float)):
            tabs[k] = [{"i": i, k: x} for i, x in enumerate(v)]
    return scal, tabs


def _noise_generic(ws, r, noise):
    """Fallback for an unknown cv_noise.json layout: scalars and row tables as found."""
    scal, tabs = _flatten_json(noise)
    if scal:
        header(ws, r, ["Statistic", "Value"], height=20)
        r += 1
        for k, v in scal.items():
            put(ws, r, 1, k, font=BOLD)
            put(ws, r, 2, v, "0.####" if isinstance(v, (int, float)) else None, INPUT_FILL)
            r += 1
        r += 1
    for name, rows in tabs.items():
        keys = list(dict.fromkeys(k for row in rows for k, v in row.items() if not isinstance(v, (list, dict))))
        put(ws, r, 1, name, font=BOLD, border=False)
        r += 1
        header(ws, r, keys, height=30)
        a = r + 1
        for i, row in enumerate(rows):
            for k, key in enumerate(keys):
                v = row.get(key)
                put(ws, a + i, 1 + k, v, "0.####" if isinstance(v, float) else None, INPUT_FILL)
        b = a + len(rows) - 1
        numeric = [k for k, key in enumerate(keys) if all(isinstance(row.get(key), (int, float)) for row in rows)]
        if numeric and len(rows) > 2:
            put(ws, b + 1, 1, "mean", font=BOLD, fill=TOTAL_FILL)
            for k in numeric:
                put(ws, b + 1, 1 + k, f"=AVERAGE({col(1 + k)}{a}:{col(1 + k)}{b})", "0.####", TOTAL_FILL)
            b += 1
        r = b + 2
    if not scal and not tabs:
        r = pending(ws, r, "cv_noise.json has no rows or scalars I can tabulate (format changed?)", span=6)
    return r


SWEEP_COLS = [("netd", "NETD (DN)", "0"), ("pd", "P_d (rover)", "0.000"), ("pdFar", "P_d (half-size rover)", "0.000"),
              ("falsePerFrame", "False targets / frame", "0.000"), ("personAsTarget", "Soldier classified as target", "0.000"),
              ("personRecognised", "Person recognised", "0.000"), ("centroidRms", "Centroid RMS (px)", "0.000")]


def _noise_sweep(ws, r, nz):
    """WP-D layout: parallel arrays per metric over the NETD points + noNuc row + config."""
    put(ws, r, 1, "Frames per point", font=BOLD)
    put(ws, r, 2, nz.get("frames"), "0", INPUT_FILL)
    r += 1
    header(ws, r, [c[1] for c in SWEEP_COLS], height=40)
    h = r
    n = len(nz["netd"])
    for i in range(n):
        for k, (key, _, fmt) in enumerate(SWEEP_COLS):
            arr = nz.get(key) or []
            put(ws, h + 1 + i, 1 + k, arr[i] if i < len(arr) else None, fmt, INPUT_FILL)
    last = h + n
    nn = nz.get("noNuc") or {}
    if nn:
        for k, (key, _, fmt) in enumerate(SWEEP_COLS):
            v = nn.get(key)
            put(ws, last + 1, 1 + k, v, fmt, NOTE_FILL, BOLD if k == 0 else None)
        put(ws, last + 1, 8, "last row: NUC switched off (at NETD 2)")
    put(ws, last + 2, 1, "Worst", font=BOLD, fill=TOTAL_FILL)
    put(ws, last + 2, 2, f"=MIN(B{h + 1}:B{last})", "0.000", TOTAL_FILL)
    put(ws, last + 2, 3, f"=MIN(C{h + 1}:C{last})", "0.000", TOTAL_FILL)
    for c in (4, 5, 7):
        put(ws, last + 2, c, f"=MAX({col(c)}{h + 1}:{col(c)}{last})", "0.000", TOTAL_FILL)
    put(ws, last + 2, 6, f"=MIN(F{h + 1}:F{last})", "0.000", TOTAL_FILL)
    put(ws, last + 3, 1, "P_d >= 0.97 up to NETD (DN)", font=BOLD)
    put(ws, last + 3, 2, f"=SUMPRODUCT(MAX((B{h + 1}:B{last}>=0.97)*A{h + 1}:A{last}))", "0")
    ch = ScatterChart()
    ch.title, ch.style = "Detection and false alarms vs NETD", 13
    ch.x_axis.title, ch.y_axis.title = "NETD (DN)", "probability / rate"
    xs = Reference(ws, min_col=1, min_row=h + 1, max_row=last)
    for c in (2, 3, 4, 5):
        sr = Series(Reference(ws, min_col=c, min_row=h, max_row=last), xs, title_from_data=True)
        sr.marker.symbol, sr.smooth = "circle", False
        ch.series.append(sr)
    ch.height, ch.width = 8.5, 17
    ws.add_chart(style_chart(ch), f"J{h}")
    r = last + 5
    cfg = nz.get("config") or {}
    if cfg:
        put(ws, r, 1, "Sensor config used by the sweep", font=BOLD, border=False)
        r += 1
        header(ws, r, ["Key", "Value"], height=20)
        for i, (k, v) in enumerate(cfg.items()):
            put(ws, r + 1 + i, 1, k, font=BOLD)
            put(ws, r + 1 + i, 2, v, "0.####", INPUT_FILL)
        r += len(cfg) + 2
    ex = nz.get("examples") or []
    if ex:
        put(ws, r, 1, f"{len(ex)} example crops ({nz.get('cropW')} x {nz.get('cropH')} px at {nz.get('cropX')}, {nz.get('cropY')}) "
                      "are stored in cv_noise.json (shown in the report figure).", border=False)
        r += 1
    return r


def build_cvnoise(wb, ctx):
    ws = wb[S_CVN]
    v6 = ctx["scene"].get("v6") or {}
    title(ws, "23  CV Noise & Seeds",
          "v6 real computer vision: the thermal sensor noise model applied between the render and the detector "
          "(NETD, fixed pattern, NUC residual, PSF, stuck pixels, jitter, frame drops), the P_d / false-alarm noise sweep "
          "(build/figdata_out/cv_noise.json), the per-stream seeds derived from the run seed, and the v6 tests "
          "(T190-T239) with their status looked up from sheet 10. Blue = exported; status cells are formulas.", span=8)
    r = 4
    section(ws, r, 1, "Sensor and detector configuration (SensorNoiseConfig + CV thresholds)")
    r += 1
    sens = v6.get("sensor") or []
    if not sens:
        r = pending(ws, r, "v6.sensor not exported yet (WP-E DocExport; the config comes from WP-D's SensorModel.h)", span=6)
    else:
        header(ws, r, ["Key", "Value", "Unit", "Source"], height=20)
        r += 1
        for s in sens:
            put(ws, r, 1, s.get("key"), font=BOLD)
            put(ws, r, 2, s.get("value"), "0.###", INPUT_FILL)
            put(ws, r, 3, s.get("unit"))
            put(ws, r, 4, s.get("source"))
            r += 1
    r += 1
    section(ws, r, 1, "Noise sweep (P_d and false alarms against NETD)")
    r += 1
    noise = load_json("build/figdata_out/cv_noise.json")
    if noise is None:
        r = pending(ws, r, "build/figdata_out/cv_noise.json (WP-D `figdata` noise sweep) not generated yet", span=6)
    else:
        if isinstance(noise, dict) and isinstance(noise.get("netd"), list):
            r = _noise_sweep(ws, r, noise)
        else:
            r = _noise_generic(ws, r, noise)
    r += 1
    section(ws, r, 1, "Seeds (deriveSeed(base, stream))")
    r += 1
    seeds = v6.get("seeds") or {}
    if not seeds:
        r = pending(ws, r, "v6.seeds not exported yet", span=6)
    else:
        put(ws, r, 1, "Example run seed", font=BOLD)
        put(ws, r, 2, seeds.get("example_base"), "0", INPUT_FILL)
        r += 1
        header(ws, r, ["Stream", "Seed"], height=20)
        a = r + 1
        for i, s in enumerate(seeds.get("streams") or []):
            put(ws, a + i, 1, s.get("stream"), font=BOLD)
            put(ws, a + i, 2, s.get("seed"), "0", INPUT_FILL)
        n = len(seeds.get("streams") or [])
        r = a + n
        if n:
            put(ws, r, 1, "Distinct seeds", font=BOLD, fill=TOTAL_FILL)
            put(ws, r, 2, f"=SUMPRODUCT(1/COUNTIF(B{a}:B{a + n - 1},B{a}:B{a + n - 1}))", "0", TOTAL_FILL)
            put(ws, r, 3, f'=IF(B{r}=COUNTA(A{a}:A{a + n - 1}),"all streams independent","collision")')
            r += 1
    r += 1
    section(ws, r, 1, "v6 tests T190-T239 (status from sheet 10)")
    r += 1
    header(ws, r, ["Test", "Block", "What is checked (build/test_results.json)", "Observed", "Status (sheet 10)"], height=24)
    by_id = {t["id"]: t for t in ctx["tests"]}
    t_first, t_last = ctx["test_rows"]
    ids, st = f"'{S_TEST}'!$A${t_first}:$A${t_last}", f"'{S_TEST}'!$F${t_first}:$F${t_last}"
    blocks = [(190, 199, "day cycle / sky / clouds"), (200, 214, "soldiers"), (215, 224, "random paths + day rules"),
              (225, 236, "sensor model + real CV"), (237, 239, "app v6")]
    a = r + 1
    rr = a
    for lo, hi, blk in blocks:
        for n in range(lo, hi + 1):
            tid = f"T{n}"
            t = by_id.get(tid)
            put(ws, rr, 1, tid, font=BOLD)
            put(ws, rr, 2, blk)
            put(ws, rr, 3, t["description"] if t else "pending: test not in build/test_results.json yet",
                fill=INPUT_FILL if t else PENDING_FILL)
            put(ws, rr, 4, t["observed"] if t else None, fill=INPUT_FILL if t else None)
            put(ws, rr, 5, f'=IFERROR(INDEX({st},MATCH(A{rr},{ids},0)),"pending")', font=BOLD)
            rr += 1
    b = rr - 1
    ws.conditional_formatting.add(f"E{a}:E{b}", CellIsRule(operator="equal", formula=['"PASS"'], fill=PASS_FILL))
    ws.conditional_formatting.add(f"E{a}:E{b}", CellIsRule(operator="equal", formula=['"FAIL"'], fill=FAIL_FILL))
    ws.conditional_formatting.add(f"E{a}:E{b}", CellIsRule(operator="equal", formula=['"pending"'], fill=PENDING_FILL))
    for k, (lab, f) in enumerate((("v6 tests", f"=ROWS(E{a}:E{b})"), ("Passed", f'=COUNTIF(E{a}:E{b},"PASS")'),
                                  ("Failed", f'=COUNTIF(E{a}:E{b},"FAIL")'), ("Pending", f'=COUNTIF(E{a}:E{b},"pending")'))):
        put(ws, 3, 1 + 2 * k, lab, font=BOLD, fill=TOTAL_FILL)
        put(ws, 3, 2 + 2 * k, f, "0", font=BOLD)
    ctx["refs"].update(v6_tests_passed=f"'{S_CVN}'!$D$3", v6_tests_pending=f"'{S_CVN}'!$H$3")
    widths(ws, [30, 24, 56, 44, 18, 12, 12, 12])
    finish(ws, 5, freeze="A4")


def build_v6(wb, ctx):
    build_daycycle(wb, ctx)
    build_soldiers(wb, ctx)
    build_cvnoise(wb, ctx)
