#!/usr/bin/env python3
"""v10 teaching workbook: report/workbooks/cg_workbook.xlsx (11 formula-driven sheets) + per-sheet PNG snaps.

Run:  ~/.local/bin/uv run -q --with openpyxl --with matplotlib python tools/v10/build_cg_workbook.py
      [--export FILE.json]  (default: runs ./build/ts-v9 --export-docs into a temp file)
      [--no-snap]           (only write the xlsx)

Every derived number in the workbook is a live Excel formula. Yellow cells are inputs copied from the code or
from the GPU-free scene export (ts --export-docs); worked textbook examples are labelled as such.
The snap step recalculates the workbook with LibreOffice (headless), exports a PDF, checks every formula cell
for errors, and writes report/figures/workbook/<Sheet>.png (200 dpi, white margins cropped).
Re-runnable: the Performance sheet imports report/data/bench.csv when it exists, otherwise it keeps a TODO row.
"""
import argparse
import csv
import json
import math
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OUT_XLSX = ROOT / "report" / "workbooks" / "cg_workbook.xlsx"
SNAP_DIR = ROOT / "report" / "figures" / "workbook"
BENCH_CSV = ROOT / "report" / "data" / "bench.csv"
TS_BIN = ROOT / "build" / "ts-v9"
SOFFICE = Path.home() / ".local/opt/squashfs-root/opt/libreoffice25.8/program/soffice"
UV = Path.home() / ".local/bin/uv"
NSS_LIB = Path.home() / ".local/opt/nss/root/usr/lib/x86_64-linux-gnu"

SHEETS = ["Primitives", "Decomposition", "Transform_Trace", "View_Projection", "Clipping", "Illumination",
          "Color_Interp", "Ray_Tracing", "Thermal_Kernels", "Duty_Cycle", "Performance"]

# --------------------------------------------------------------------------------------------- workbook helpers
if __name__ == "__main__" and "--snap-only" not in sys.argv:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.pagebreak import Break

    YELLOW = PatternFill("solid", fgColor="FFF2CC")
    HEAD = PatternFill("solid", fgColor="D9E1F2")
    THIN = Side(style="thin", color="7F7F7F")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    NUM = "0.0000"


def A(r, c, absolute=False, sheet=None):
    """Cell address; absolute=True gives $A$1; sheet prefixes 'Sheet'!."""
    from openpyxl.utils import get_column_letter as gcl
    s = f"${gcl(c)}${r}" if absolute else f"{gcl(c)}{r}"
    return f"'{sheet}'!{s}" if sheet else s


def rng(r0, c0, r1, c1, sheet=None):
    from openpyxl.utils import get_column_letter as gcl
    s = f"{gcl(c0)}{r0}:{gcl(c1)}{r1}"
    return f"'{sheet}'!{s}" if sheet else s


def smooth(x, e0, e1):
    """Inline GLSL smoothstep(e0, e1, x) as an Excel expression."""
    t = f"MIN(MAX(({x}-({e0}))/(({e1})-({e0})),0),1)"
    return f"({t}^2*(3-2*{t}))"


class Sheet:
    def __init__(self, wb, name, title, source):
        self.ws = wb.create_sheet(name)
        self.name = name
        self.ws.sheet_view.showGridLines = False
        self.maxr, self.maxc = 3, 1
        self.put(1, 1, title, border=False, bold=True, size=13)
        self.put(2, 1, "Legend: yellow = input from the code; white = formula.", border=False, italic=True)
        self.put(3, 1, "Source: " + source, border=False, italic=True)

    def put(self, r, c, v, inp=False, fmt=None, bold=False, border=True, head=False, italic=False, size=10,
            wrap=False, center=False, track=True):
        cell = self.ws.cell(row=r, column=c, value=v)
        cell.font = Font(name="Calibri", size=size, bold=bold or head, italic=italic)
        if head:
            cell.fill = HEAD
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        elif inp:
            cell.fill = YELLOW
        if border:
            cell.border = BORDER
        if wrap and not head:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        elif center and not head:
            cell.alignment = Alignment(horizontal="center")
        numeric = (isinstance(v, (int, float)) and not isinstance(v, bool)) or (isinstance(v, str) and v.startswith("="))
        if fmt:
            cell.number_format = fmt
        elif numeric and not head:
            cell.number_format = "0" if isinstance(v, int) else NUM
        if track:
            self.maxr = max(self.maxr, r)
            if border or c == 1:
                self.maxc = max(self.maxc, c)
        return A(r, c)

    def note(self, r, c, text, bold=False):
        return self.put(r, c, text, border=False, italic=not bold, bold=bold)

    def heads(self, r, c, labels, height=None):
        for i, lab in enumerate(labels):
            self.put(r, c + i, lab, head=True)
        if height:
            self.ws.row_dimensions[r].height = height

    def widths(self, spec):
        for col, w in spec.items():
            self.ws.column_dimensions[col].width = w

    def grid(self, r, c, title, cells, inp=False, fmt=None):
        """Titled 4x4 (or n x m) grid. cells: list of rows (values or formulas). Returns the top-left cell (row, col)."""
        self.put(r, c, title, border=False, bold=True)
        for i, row in enumerate(cells):
            for j, v in enumerate(row):
                self.put(r + 1 + i, c + j, v, inp=inp, fmt=fmt, center=True)
        return r + 1, c

    def finish(self, fit_height=1, last_col=None, last_row=None):
        ws = self.ws
        lc = last_col or self.maxc
        lr = last_row or self.maxr
        ws.print_area = rng(1, 1, lr, lc)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = fit_height
        ws.print_options.gridLines = False
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.4
        ws.page_margins.header = ws.page_margins.footer = 0.2


def matmul_cells(a, b):
    """Formulas of the 4x4 product a*b; a, b are (row, col) of the grids' top-left cells."""
    (ar, ac), (br, bc) = a, b
    out = []
    for i in range(4):
        row = []
        for j in range(4):
            terms = [f"{A(ar + i, ac + k)}*{A(br + k, bc + j)}" for k in range(4)]
            row.append("=" + "+".join(terms))
        out.append(row)
    return out


def matvec(m, v):
    """Formulas of the 4x4 grid m times the column vector v (list of 4 cell addresses)."""
    mr, mc = m
    return ["=" + "+".join(f"{A(mr + i, mc + k)}*{v[k]}" for k in range(4)) for i in range(4)]


def rowmajor(m16):
    return [m16[i * 4:(i + 1) * 4] for i in range(4)]


# --------------------------------------------------------------------------------------------- data
def load_export(path):
    if path:
        return json.loads(Path(path).read_text())
    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "export.json"
        subprocess.run([str(TS_BIN), "--export-docs", str(out)], cwd=ROOT, check=True, stdout=subprocess.DEVNULL)
        return json.loads(out.read_text())


def comp_index(exp):
    return {c["id"]: c for c in exp["components"]}


# --------------------------------------------------------------------------------------------- sheets
def sheet_primitives(wb, exp):
    s = Sheet(wb, "Primitives", "Primitives: vertex and triangle counts of the real mesh generators",
              "src/primitives/{Cube,Sphere,Cylinder}.cpp, objects/Terrain.cpp createTerrainDisc, rendering/Grass.cpp; "
              "exported counts: ts --export-docs 'primitives'.")
    s.widths({"A": 19, "B": 33, "C": 7, "D": 8, "E": 12, "F": 8, "G": 11, "H": 8, "I": 8, "J": 11, "K": 11,
              "L": 8, "M": 8, "N": 6, "O": 6, "P": 9, "Q": 11})
    r = 5
    s.heads(r, 1, ["Primitive", "Generator call", "slices s", "stacks / rings k", "V formula", "V",
                   "T formula", "T", "indices I = 3T", "bytes indexed 24V+4I", "bytes non-indexed 72T",
                   "V export", "T export", "dV", "dT", "instances (export)", "scene triangles T export x inst"], height=42)
    prims = {p["type"]: p for p in exp["primitives"]}
    rows = [
        ("Cube", "createCube()", None, None, "6 faces x 4", "=6*4", "6 faces x 2", "=6*2"),
        ("Sphere", "createSphere(36, 18)", 36, 18, "(k-1)s + 2", "=(D{r}-1)*C{r}+2", "2s(k-1)", "=2*C{r}*(D{r}-1)"),
        ("Cylinder", "createCylinder(36) (capped)", 36, None, "4s + 2", "=4*C{r}+2", "4s", "=4*C{r}"),
        ("Cylinder (frustum)", "createCylinder(36, 0.03, 0.5, false)", 36, None, "2s", "=2*C{r}", "2s", "=2*C{r}"),
        ("Cone", "createCylinder(36, 0.5, 0, true)", 36, None, "3s + 1", "=3*C{r}+1", "2s", "=2*C{r}"),
        ("Terrain disc", "createTerrainDisc(130, 384, 36)", 384, 130, "1 + k s", "=1+D{r}*C{r}", "s(2k-1)",
         "=C{r}*(2*D{r}-1)"),
        ("Grass blade", "Grass.cpp instanced strip (3 segments)", None, 3, "2k + 1", "=2*D{r}+1", "2k - 1",
         "=2*D{r}-1"),
    ]
    first = r + 1
    for name, call, sl, st, vtxt, vf, ttxt, tf in rows:
        r += 1
        s.put(r, 1, name, bold=True)
        s.put(r, 2, call, inp=True)
        s.put(r, 3, sl if sl is not None else "-", inp=sl is not None, center=True)
        s.put(r, 4, st if st is not None else "-", inp=st is not None, center=True)
        s.put(r, 5, vtxt, center=True)
        s.put(r, 6, vf.format(r=r), fmt="0")
        s.put(r, 7, ttxt, center=True)
        s.put(r, 8, tf.format(r=r), fmt="0")
        s.put(r, 9, f"=3*H{r}", fmt="0")
        s.put(r, 10, f"=24*F{r}+4*I{r}", fmt="0")
        s.put(r, 11, f"=72*H{r}", fmt="0")
        if name in prims:
            p = prims[name]
            s.put(r, 12, p["vertices"], inp=True)
            s.put(r, 13, p["triangles"], inp=True)
            s.put(r, 16, p["instances"], inp=True)
            s.put(r, 17, f"=M{r}*P{r}", fmt="0")
        else:  # grass: counts from the Grass.cpp comment (7 vertices, 5 triangles); not in the primitive library
            s.put(r, 12, 7, inp=True)
            s.put(r, 13, 5, inp=True)
            s.put(r, 16, "-", center=True)
            s.put(r, 17, "-", center=True)
        s.put(r, 14, f"=F{r}-L{r}", fmt="0")
        s.put(r, 15, f"=H{r}-M{r}", fmt="0")
    last = r
    r += 1
    s.put(r, 1, "Scene total (library)", bold=True)
    s.put(r, 16, f"=SUM(P{first}:P{last - 1})", fmt="0")
    s.put(r, 17, f"=SUM(Q{first}:Q{last - 1})", fmt="0")
    r += 1
    s.note(r, 1, "dV = dT = 0: the closed-form counts match the exported meshes. Terrain dT = triangles whose centroid "
                 "lies in the entrance-trench hole (dropped by createTerrainDisc).")
    r += 1
    s.note(r, 1, "Bytes: 24-byte vertex (position + normal, 6 floats), 4-byte indices; non-indexed = 3 vertices per "
                 "triangle (72 bytes). Exported bytes_indexed / bytes_non_indexed agree.")
    r += 2
    s.put(r, 1, "Not primitives in this code (docs/audit.md section 2)", border=False, bold=True)
    r += 1
    s.heads(r, 1, ["Shape", "Status in the code"])
    for shape, status in [("Cuboid", "a scaled createCube() (S per part); not a separate generator"),
                          ("Hemisphere", "NOT a primitive (only scaled / clipped spheres in places)"),
                          ("Torus", "NOT a primitive: no generator exists"),
                          ("Wedge / prism", "NOT a primitive: no generator exists"),
                          ("Helix", "NOT a primitive: no generator exists")]:
        r += 1
        s.put(r, 1, shape, bold=True)
        s.put(r, 2, status)
        s.ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    r += 2
    s.put(r, 1, "Parametric equations (local space, unit size)", border=False, bold=True)
    r += 1
    s.heads(r, 1, ["Primitive", "Point p(parameters)", "", "", "", "", "", "", "", "Parameter ranges"])
    s.ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
    s.ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=17)
    eqs = [
        ("Cube", "p = 0.5 n + 0.5 (a u + b v), face frame (n, u, v) with u x v = n",
         "6 faces, a, b in {-1, +1}; 2 CCW triangles (0,1,2), (0,2,3) per face"),
        ("Sphere", "p = 0.5 (sin(th) cos(ph), cos(th), sin(th) sin(ph)), normal = p / 0.5",
         "th = pi k / stacks, k = 1..stacks-1; ph = 2 pi j / slices; 2 poles; south half mirrored (y -> -y)"),
        ("Cylinder / frustum / cone", "p = (r(z) cos(ph), r(z) sin(ph), z), r(-0.5) = r_bottom, r(+0.5) = r_top",
         "ph = 2 pi j / slices; cylinder r = 0.5 / 0.5, frustum 0.5 / 0.03, cone 0.5 / 0 (apex); caps = fans"),
        ("Terrain disc", "p = (r cos(ph), h(x, z), r sin(ph)), r = 36 k / rings",
         "k = 0..130, ph = 2 pi j / 384; h = terrainHeight(x, z); normal by central differences"),
        ("Grass blade", "p = (+-0.625 w(t), t - 0.5, 0), t in {0, 0.38, 0.72}; tip (0, 0.5, 0)",
         "bent and scaled per instance in the vertex shader (wind, lean); one-sided strip"),
    ]
    for name, eq, rg in eqs:
        r += 1
        s.put(r, 1, name, bold=True)
        s.put(r, 2, eq)
        s.ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
        s.put(r, 10, rg)
        s.ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=17)
    for rr in range(r - len(eqs) + 1, r + 1):
        for cc in range(2, 18):
            s.ws.cell(row=rr, column=cc).border = BORDER
    s.finish(last_col=17)


PRIM_ORDER = ["Cube", "Sphere", "Cylinder", "Cylinder (frustum)", "Cone", "Terrain disc"]


def sheet_decomposition(wb, exp):
    s = Sheet(wb, "Decomposition", "Decomposition: objects as hierarchies of primitive parts",
              "ts --export-docs 'components' (S, R_deg, T, parent, colour, specular, shininess) and 'objects'; "
              "objects/*.cpp Builder. Rows below list every drawn part (leaf) of UAV-1 and of Soldier 1.")
    s.widths({"A": 22, "B": 21, "C": 15, "D": 28, "E": 8, "F": 8, "G": 8, "H": 8, "I": 8, "J": 8, "K": 8,
              "L": 8, "M": 8, "N": 17, "O": 10, "P": 9, "Q": 8, "R": 10, "S": 9, "T": 8})
    comps = exp["components"]
    byid = comp_index(exp)
    uav = [c for c in comps if c["id"].startswith("UAV-") and c["primitive"] != "Group"]
    sol = [c for c in comps if c["object"] == "Soldier 1" and c["primitive"] != "Group"]
    uav_ctrl = next(c for c in exp["controllables"] if c.get("root_id") == "UAV-00")
    sol_obj = next(o for o in exp["objects"] if o["object"] == "Soldier 1")

    # Page 1: object table (left) + COUNTIFS summary (right); part tables after it.
    n_obj = len(exp["objects"])
    obj_r = 5
    sum_r, sc = 5, 14
    uav_r = obj_r + n_obj + 7
    sol_r = uav_r + len(uav) + 4
    uav_rng = (uav_r + 2, uav_r + 1 + len(uav))
    sol_rng = (sol_r + 2, sol_r + 1 + len(sol))

    # 1. primitive totals of the two listed objects (COUNTIFS over the part rows) vs the export
    s.put(sum_r, sc, "2. Parts per primitive (COUNTIFS over the part rows)", border=False, bold=True)
    s.heads(sum_r + 1, sc, ["Primitive", "UAV-1 COUNTIFS", "UAV-1 export", "check", "Soldier 1 COUNTIFS",
                            "Soldier 1 export", "check"], height=30)
    r = sum_r + 1
    for p in PRIM_ORDER[:5]:
        r += 1
        s.put(r, sc, p, bold=True)
        s.put(r, sc + 1, f'=COUNTIFS($C${uav_rng[0]}:$C${uav_rng[1]},{A(r, sc)})', fmt="0")
        s.put(r, sc + 2, uav_ctrl["primitives"].get(p, 0), inp=True)
        s.put(r, sc + 3, f'=IF({A(r, sc + 1)}={A(r, sc + 2)},"OK","DIFF")', center=True)
        s.put(r, sc + 4, f'=COUNTIFS($C${sol_rng[0]}:$C${sol_rng[1]},{A(r, sc)})', fmt="0")
        s.put(r, sc + 5, sol_obj["primitives"].get(p, 0), inp=True)
        s.put(r, sc + 6, f'=IF({A(r, sc + 4)}={A(r, sc + 5)},"OK","DIFF")', center=True)
    r += 1
    s.put(r, sc, "Total parts", bold=True)
    for c in (1, 2, 4, 5):
        s.put(r, sc + c, f"=SUM({A(sum_r + 2, sc + c)}:{A(r - 1, sc + c)})", fmt="0")
    s.put(r, sc + 3, f'=IF({A(r, sc + 1)}={A(r, sc + 2)},"OK","DIFF")', center=True)
    s.put(r, sc + 6, f'=IF({A(r, sc + 4)}={A(r, sc + 5)},"OK","DIFF")', center=True)
    s.note(r + 2, sc, "UAV-1 export = controllable 'UAV-1 DRONE' (root UAV-00);")
    s.note(r + 3, sc, "the object 'Autonomous UAV' below counts both UAVs.")

    # 2. object-level table
    s.put(obj_r, 1, "1. Object-level primitive totals (whole scene, from the export)", border=False, bold=True)
    s.heads(obj_r + 1, 1, ["Object"] + PRIM_ORDER + ["parts (sum)", "parts export", "check", "groups", "animated nodes"],
            height=30)
    r = obj_r + 1
    for o in exp["objects"]:
        r += 1
        s.put(r, 1, o["object"], bold=True)
        for j, p in enumerate(PRIM_ORDER):
            s.put(r, 2 + j, o["primitives"].get(p, 0), inp=True)
        s.put(r, 8, f"=SUM(B{r}:G{r})", fmt="0")
        s.put(r, 9, o["parts"], inp=True)
        s.put(r, 10, f'=IF(H{r}=I{r},"OK","DIFF")', center=True)
        s.put(r, 11, o["groups"], inp=True)
        s.put(r, 12, o["animated_nodes"], inp=True)
    o0, o1 = obj_r + 2, r
    r += 1
    s.put(r, 1, "Scene total", bold=True)
    for c in range(2, 10):
        s.put(r, c, f"=SUM({A(o0, c)}:{A(o1, c)})", fmt="0")
    for c in (11, 12):
        s.put(r, c, f"=SUM({A(o0, c)}:{A(o1, c)})", fmt="0")
    r += 1
    s.put(r, 1, "Instances (export)", bold=True)
    prims = {p["type"]: p for p in exp["primitives"]}
    for j, p in enumerate(PRIM_ORDER):
        s.put(r, 2 + j, prims[p]["instances"], inp=True)
    s.put(r, 8, f"=SUM(B{r}:G{r})", fmt="0")
    s.put(r, 9, exp["leaf_count"], inp=True)
    r += 1
    s.put(r, 1, "check", bold=True)
    for c in range(2, 10):
        s.put(r, c, f'=IF({A(r - 2, c)}={A(r - 1, c)},"OK","DIFF")', center=True)
    s.ws.row_breaks.append(Break(id=r + 1))

    # 3. part tables
    hdr = ["ID", "Part", "Primitive", "Parent (group)", "Sx", "Sy", "Sz", "Rx (deg)", "Ry (deg)", "Rz (deg)",
           "Tx (m)", "Ty (m)", "Tz (m)", "R", "G", "B", "ks", "shininess n"]

    def parts(r0, title, rows):
        s.put(r0, 1, title, border=False, bold=True)
        s.heads(r0 + 1, 1, hdr, height=28)
        r = r0 + 1
        for c in rows:
            r += 1
            par = byid.get(c["parent_id"], {})
            vals = [c["id"], c["name"], c["primitive"], f'{par.get("name", "-")} ({c["parent_id"]})'] + \
                   [round(v, 4) for v in c["S"]] + [round(v, 4) for v in c["R_deg"]] + \
                   [round(v, 4) for v in c["T"]] + [round(v, 4) for v in c["color"]] + \
                   [round(c["specular"], 4), round(c["shininess"], 4)]
            for j, v in enumerate(vals):
                s.put(r, 1 + j, v, inp=True, fmt=NUM if j >= 4 else None)
        return r

    end = parts(uav_r, f"3. UAV-1 (root UAV-00 'Autonomous UAV'): {len(uav)} drawn parts. M = T Ry Rx Rz S, "
                       "world = parent_world M (the export's object 'Autonomous UAV' holds 2 UAVs).", uav)
    s.ws.row_breaks.append(Break(id=end + 1))
    parts(sol_r, f"4. Soldier 1 (root SLD1-00): {len(sol)} drawn parts (joint angles live in the parent groups).", sol)
    s.finish(fit_height=0, last_col=20)


def sheet_view_projection(wb, exp):
    s = Sheet(wb, "View_Projection", "View_Projection: default orbit camera, lookAt basis, perspective and depth",
              "scene/Camera.h (target, yaw 32, pitch 27, distance 34, fov 50), Camera.cpp compute() (orbit eye), "
              "Math3D.cpp lookAt / perspective, near 0.05 far 600, 24-bit depth (FrameBuffers.cpp).")
    s.widths({"A": 30, "B": 10, "C": 10, "D": 10, "E": 10, "F": 3, "G": 10, "H": 11, "I": 11, "J": 12, "K": 11,
              "L": 11, "M": 12})
    S = s.put
    r = 5
    s.heads(r, 1, ["Input / derived", "x", "y", "z"])
    inputs = [("target (m)", [0, 0.8, 0]), ("up", [0, 1, 0])]
    cells = {}
    for lab, v in inputs:
        r += 1
        S(r, 1, lab, bold=True)
        for j, x in enumerate(v):
            S(r, 2 + j, x, inp=True, fmt=NUM)
        cells[lab] = r
    scal = [("yaw (deg)", 32), ("pitch (deg)", 27), ("distance (m)", 34), ("fov y (deg)", 50), ("width (px)", 1920),
            ("height (px)", 1080), ("near (m)", 0.05), ("far (m)", 600), ("near, comparison (m)", 0.5)]
    for lab, v in scal:
        r += 1
        S(r, 1, lab, bold=True)
        S(r, 2, v, inp=True, fmt="0" if isinstance(v, int) else NUM)
        cells[lab] = r
    yaw, pit, dist = (A(cells[k], 2, True) for k in ("yaw (deg)", "pitch (deg)", "distance (m)"))
    tr = cells["target (m)"]
    up = cells["up"]
    r += 1
    S(r, 1, "aspect = width / height", bold=True)
    S(r, 2, f"={A(cells['width (px)'], 2)}/{A(cells['height (px)'], 2)}")
    cells["aspect"] = r
    r += 1
    S(r, 1, "eye = target + d (cos p sin y, sin p, cos p cos y)", bold=True)
    S(r, 2, f"={A(tr, 2)}+{dist}*COS(RADIANS({pit}))*SIN(RADIANS({yaw}))")
    S(r, 3, f"={A(tr, 3)}+{dist}*SIN(RADIANS({pit}))")
    S(r, 4, f"={A(tr, 4)}+{dist}*COS(RADIANS({pit}))*COS(RADIANS({yaw}))")
    eye = r
    cells["eye"] = r
    r += 1
    S(r, 1, "eye - target", bold=True)
    for j in range(3):
        S(r, 2 + j, f"={A(eye, 2 + j)}-{A(tr, 2 + j)}")
    d = r
    r += 1
    S(r, 1, "n = normalize(eye - target)", bold=True)
    nrm = f"SQRT({A(d, 2)}^2+{A(d, 3)}^2+{A(d, 4)}^2)"
    for j in range(3):
        S(r, 2 + j, f"={A(d, 2 + j)}/{nrm}")
    n = r
    r += 1
    S(r, 1, "up x n", bold=True)
    ux, uy, uz = (A(up, 2), A(up, 3), A(up, 4))
    nx, ny, nz = (A(n, 2), A(n, 3), A(n, 4))
    S(r, 2, f"={uy}*{nz}-{uz}*{ny}")
    S(r, 3, f"={uz}*{nx}-{ux}*{nz}")
    S(r, 4, f"={ux}*{ny}-{uy}*{nx}")
    c = r
    r += 1
    S(r, 1, "u = normalize(up x n)", bold=True)
    nrm = f"SQRT({A(c, 2)}^2+{A(c, 3)}^2+{A(c, 4)}^2)"
    for j in range(3):
        S(r, 2 + j, f"={A(c, 2 + j)}/{nrm}")
    u = r
    r += 1
    S(r, 1, "v = n x u", bold=True)
    vx_, vy_, vz_ = (A(u, 2), A(u, 3), A(u, 4))
    S(r, 2, f"={ny}*{vz_}-{nz}*{vy_}")
    S(r, 3, f"={nz}*{vx_}-{nx}*{vz_}")
    S(r, 4, f"={nx}*{vy_}-{ny}*{vx_}")
    v = r
    r += 1
    S(r, 1, "code: s = normalize(f x up), f = -n; |s - u|", bold=True)
    # f x up with f = -n  ==  up x n, so s == u; shown as a numeric check
    S(r, 2, f"=ABS(-({ny}*{uz}-{nz}*{uy})-{A(u, 2)}*{nrm})+ABS(-({nz}*{ux}-{nx}*{uz})-{A(u, 3)}*{nrm})"
            f"+ABS(-({nx}*{uy}-{ny}*{ux})-{A(u, 4)}*{nrm})")
    r += 1
    S(r, 1, "u . v,  u . n,  v . n (orthogonal: 0)", bold=True)
    dot = lambda a, b: "+".join(f"{A(a, 2 + j)}*{A(b, 2 + j)}" for j in range(3))
    S(r, 2, "=" + dot(u, v))
    S(r, 3, "=" + dot(u, n))
    S(r, 4, "=" + dot(v, n))
    r += 2
    # view matrix
    vm = [[f"={A(u, 2)}", f"={A(u, 3)}", f"={A(u, 4)}", f"=-({dot(u, eye)})"],
          [f"={A(v, 2)}", f"={A(v, 3)}", f"={A(v, 4)}", f"=-({dot(v, eye)})"],
          [f"={A(n, 2)}", f"={A(n, 3)}", f"={A(n, 4)}", f"=-({dot(n, eye)})"],
          [0, 0, 0, 1]]
    vtop = s.grid(r, 2, "View V = lookAt(eye, target, up): rows u, v, n; last column -(axis . eye)", vm)
    S(r + 1, 1, "row 0 (u)", bold=True)
    S(r + 2, 1, "row 1 (v)", bold=True)
    S(r + 3, 1, "row 2 (n)", bold=True)
    S(r + 4, 1, "row 3", bold=True)
    r += 6
    fov, asp, near, far = (A(cells[k], 2, True) for k in ("fov y (deg)", "aspect", "near (m)", "far (m)"))
    S(r, 1, "t = 1 / tan(fov / 2)", bold=True)
    S(r, 2, f"=1/TAN(RADIANS({fov})/2)")
    t = A(r, 2, True)
    r += 1
    pm = [[f"={t}/{asp}", 0, 0, 0], [0, f"={t}", 0, 0],
          [0, 0, f"=-({far}+{near})/({far}-{near})", f"=-2*{far}*{near}/({far}-{near})"], [0, 0, -1, 0]]
    ptop = s.grid(r, 2, "Perspective P (Math3D perspective, OpenGL clip z in [-w, w])", pm)
    for i in range(4):
        S(r + 1 + i, 1, f"row {i}", bold=True)
    r += 6
    s.note(r, 1, "Code (lookAt): f = normalize(target - eye), s = normalize(f x up), u' = s x f; rows s, u', -f. "
                 "Since f = -n: s = u and u' = v, the same matrix.")
    cells["view_top"] = vtop
    cells["proj_top"] = ptop

    # depth table
    r = 5
    c0 = 7
    s.put(r - 1, c0, "Depth mapping (24-bit fixed point: step = 1 / (2^24 - 1))", border=False, bold=True)
    s.heads(r, c0, ["z_eye (m)", "z_ndc (near 0.05)", "depth d (near 0.05)", "metres per depth step (near 0.05)",
                    "z_ndc (near 0.5)", "depth d (near 0.5)", "metres per depth step (near 0.5)"], height=44)
    n2 = A(cells["near, comparison (m)"], 2, True)
    a1, b1 = f"(-({far}+{near})/({far}-{near}))", f"(-2*{far}*{near}/({far}-{near}))"
    a2, b2 = f"(-({far}+{n2})/({far}-{n2}))", f"(-2*{far}*{n2}/({far}-{n2}))"
    step = "(1/(2^24-1))"
    zs = [-0.05, -0.1, -0.5, -1, -2, -5, -10, -34, -50, -100, -200, -400, -600]
    rows = {}
    for zv in zs:
        r += 1
        z = A(r, c0)
        S(r, c0, zv, inp=True, fmt=NUM)
        S(r, c0 + 1, f"=({a1}*{z}+{b1})/(-{z})")
        S(r, c0 + 2, f"=0.5*{A(r, c0 + 1)}+0.5")
        S(r, c0 + 3, f"={step}*({far}-{near})*{z}^2/({far}*{near})", fmt="0.00E+00")
        if -zv >= 0.5:
            S(r, c0 + 4, f"=({a2}*{z}+{b2})/(-{z})")
            S(r, c0 + 5, f"=0.5*{A(r, c0 + 4)}+0.5")
            S(r, c0 + 6, f"={step}*({far}-{n2})*{z}^2/({far}*{n2})", fmt="0.00E+00")
        else:
            for j in (4, 5, 6):
                S(r, c0 + j, "clipped", center=True)
        rows[zv] = r
    r += 2
    S(r, c0, "Summary", border=False, bold=True)
    r += 1
    r100 = rows[-100]
    S(r, c0, "depth step at 100 m (m)", bold=True)
    s.ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 2)
    S(r, c0 + 3, f"={A(r100, c0 + 3)}", fmt="0.00E+00")
    S(r, c0 + 6, f"={A(r100, c0 + 6)}", fmt="0.00E+00")
    r += 1
    S(r, c0, "ratio near 0.05 / near 0.5 at 100 m", bold=True)
    s.ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 2)
    S(r, c0 + 3, f"={A(r100, c0 + 3)}/{A(r100, c0 + 6)}")
    r += 1
    S(r, c0, "share of [0, 1] depth used by the first 1 m", bold=True)
    s.ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 2)
    S(r, c0 + 3, f"={A(rows[-1], c0 + 2)}")
    S(r, c0 + 6, f"={A(rows[-1], c0 + 5)}")
    r += 1
    S(r, c0, "share of depth used beyond 10 m", bold=True)
    s.ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 2)
    S(r, c0 + 3, f"=1-{A(rows[-10], c0 + 2)}")
    S(r, c0 + 6, f"=1-{A(rows[-10], c0 + 5)}")
    r += 2
    for line in ["d = (F / (F - N)) (1 - N / |z|): precision falls with z^2.",
                 "Step size dz = step (F - N) z^2 / (F N): ten times the near plane = ten times finer steps.",
                 "Two surfaces closer than one step at 100 m share a depth value: z-fighting.",
                 "The code keeps near 0.05 m for the first-person views; 24-bit depth."]:
        s.note(r, c0, line)
        r += 1
    s.finish(last_col=13)
    return cells


def sheet_transform_trace(wb, exp, vp):
    byid = comp_index(exp)
    part = byid["UAV-14"]
    parent = byid[part["parent_id"]]
    s = Sheet(wb, "Transform_Trace", "Transform_Trace: one real UAV vertex from local space to the window",
              "ts --export-docs component UAV-14 'Arm Front-right' (S, R_deg, T, local_matrix, world_matrix, parent "
              "UAV-00 world_matrix); objects/Builder M = T Ry Rx Rz S; View_Projection sheet (V, P, viewport).")
    s.widths({"A": 24})
    for col in range(2, 21):
        s.ws.column_dimensions[get_column_letter(col)].width = 9 if (col - 1) % 5 else 2.5
    S = s.put
    r = 5
    S(r, 1, "Part", bold=True)
    S(r, 2, f'{part["id"]} {part["name"]} ({part["primitive"]}, createCube)', inp=True)
    s.ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
    S(r, 11, "Parent", bold=True)
    s.ws.merge_cells(start_row=r, start_column=11, end_row=r, end_column=12)
    S(r, 13, f'{parent["id"]} {parent["name"]} (root, R = (0, 180, 0), T = (0, 7.5, 12.5))', inp=True)
    s.ws.merge_cells(start_row=r, start_column=13, end_row=r, end_column=20)
    r += 1
    s.heads(r, 1, ["", "x", "y", "z", "w"])
    rows = {}
    for lab, vals in [("S (scale)", part["S"]), ("R_deg (pitch x, yaw y, roll z)", part["R_deg"]),
                      ("T (m)", part["T"]), ("local vertex p (cube corner)", [0.5, 0.5, 0.5, 1])]:
        r += 1
        S(r, 1, lab, bold=True)
        for j, x in enumerate(vals):
            S(r, 2 + j, round(x, 7), inp=True)
        rows[lab] = r
    s.note(r, 6, "createCube() vertex 3 = (0.5, 0.5, 0.5), +X face (mesh_samples in the export)")
    Sx, Sy, Sz = (A(rows["S (scale)"], 2 + j, True) for j in range(3))
    Rx, Ry, Rz = (A(rows["R_deg (pitch x, yaw y, roll z)"], 2 + j, True) for j in range(3))
    Tx, Ty, Tz = (A(rows["T (m)"], 2 + j, True) for j in range(3))
    cx, sx = f"COS(RADIANS({Rx}))", f"SIN(RADIANS({Rx}))"
    cy, sy = f"COS(RADIANS({Ry}))", f"SIN(RADIANS({Ry}))"
    cz, sz = f"COS(RADIANS({Rz}))", f"SIN(RADIANS({Rz}))"
    mS = [[f"={Sx}", 0, 0, 0], [0, f"={Sy}", 0, 0], [0, 0, f"={Sz}", 0], [0, 0, 0, 1]]
    mRx = [[1, 0, 0, 0], [0, f"={cx}", f"=-{sx}", 0], [0, f"={sx}", f"={cx}", 0], [0, 0, 0, 1]]
    mRy = [[f"={cy}", 0, f"={sy}", 0], [0, 1, 0, 0], [f"=-{sy}", 0, f"={cy}", 0], [0, 0, 0, 1]]
    mRz = [[f"={cz}", f"=-{sz}", 0, 0], [f"={sz}", f"={cz}", 0, 0], [0, 0, 1, 0], [0, 0, 0, 1]]
    mT = [[1, 0, 0, f"={Tx}"], [0, 1, 0, f"={Ty}"], [0, 0, 1, f"={Tz}"], [0, 0, 0, 1]]
    cols = [2, 7, 12, 17]
    b = 12
    g = {}
    g["S"] = s.grid(b, cols[0], "S = scale(Sx, Sy, Sz)", mS)
    g["Rx"] = s.grid(b, cols[1], "Rx(pitch)", mRx)
    g["Ry"] = s.grid(b, cols[2], "Ry(yaw)", mRy)
    g["Rz"] = s.grid(b, cols[3], "Rz(roll)", mRz)
    b += 6
    g["T"] = s.grid(b, cols[0], "T = translate(Tx, Ty, Tz)", mT)
    g["RyRx"] = s.grid(b, cols[1], "Ry Rx", matmul_cells(g["Ry"], g["Rx"]))
    g["R"] = s.grid(b, cols[2], "R = Ry Rx Rz", matmul_cells(g["RyRx"], g["Rz"]))
    g["RS"] = s.grid(b, cols[3], "R S", matmul_cells(g["R"], g["S"]))
    b += 6
    g["M"] = s.grid(b, cols[0], "M = T R S (local matrix)", matmul_cells(g["T"], g["RS"]))
    g["Mx"] = s.grid(b, cols[1], "M exported (local_matrix)", [[round(x, 9) for x in row] for row in rowmajor(part["local_matrix"])], inp=True)
    g["dM"] = s.grid(b, cols[2], "|M - M exported|",
                     [[f"=ABS({A(g['M'][0] + i, g['M'][1] + j)}-{A(g['Mx'][0] + i, g['Mx'][1] + j)})" for j in range(4)]
                      for i in range(4)])
    g["P"] = s.grid(b, cols[3], "Parent world (UAV-00, exported)", [[round(x, 9) for x in row] for row in rowmajor(parent["world_matrix"])], inp=True)
    b += 6
    g["W"] = s.grid(b, cols[0], "World = Parent M", matmul_cells(g["P"], g["M"]))
    g["Wx"] = s.grid(b, cols[1], "World exported (world_matrix)", [[round(x, 9) for x in row] for row in rowmajor(part["world_matrix"])], inp=True)
    g["dW"] = s.grid(b, cols[2], "|World - World exported|",
                     [[f"=ABS({A(g['W'][0] + i, g['W'][1] + j)}-{A(g['Wx'][0] + i, g['Wx'][1] + j)})" for j in range(4)]
                      for i in range(4)])
    vt, pt = vp["view_top"], vp["proj_top"]
    g["V"] = s.grid(b, cols[3], "View V (View_Projection sheet)",
                    [[f"={A(vt[0] + i, vt[1] + j, sheet='View_Projection')}" for j in range(4)] for i in range(4)])
    b += 6
    g["Pr"] = s.grid(b, cols[0], "Projection P (View_Projection sheet)",
                     [[f"={A(pt[0] + i, pt[1] + j, sheet='View_Projection')}" for j in range(4)] for i in range(4)])
    # pipeline table
    c0 = 7
    S(b, c0, "Vertex pipeline (column vectors)", border=False, bold=True)
    s.heads(b + 1, c0, ["stage", "", "", ""])
    s.heads(b + 1, c0 + 5, ["x", "y", "z", "w"])
    s.ws.merge_cells(start_row=b + 1, start_column=c0, end_row=b + 1, end_column=c0 + 3)
    pr = rows["local vertex p (cube corner)"]
    p_local = [A(pr, 2 + j) for j in range(4)]
    stages = []
    r = b + 1

    def stage(label, vals, merge=True):
        nonlocal r
        r += 1
        S(r, c0, label, bold=True)
        for j in range(1, 4):
            S(r, c0 + j, "", track=False)
        if merge:
            s.ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 3)
        for j, v in enumerate(vals):
            S(r, c0 + 5 + j, v)
        stages.append(r)
        return [A(r, c0 + 5 + j) for j in range(4)]

    pl = stage("local p", [f"={x}" for x in p_local])
    pw = stage("world = World p", matvec(g["W"], pl))
    pe = stage("eye = V world", matvec(g["V"], pw))
    pc = stage("clip = P eye", matvec(g["Pr"], pe))
    w = pc[3]
    pn = stage("NDC = clip / w", [f"={pc[0]}/{w}", f"={pc[1]}/{w}", f"={pc[2]}/{w}", f"={w}/{w}"])
    W, H = (A(vp[k], 2, True, sheet="View_Projection") for k in ("width (px)", "height (px)"))
    stage("window (x px, y px top-left, depth, -)",
          [f"=({pn[0]}*0.5+0.5)*{W}", f"=(1-({pn[1]}*0.5+0.5))*{H}", f"={pn[2]}*0.5+0.5", "-"])
    # checks
    c1 = 17
    S(b, c1, "Checks", border=False, bold=True)
    s.heads(b + 1, c1, ["check", "", "", "value"])
    s.ws.merge_cells(start_row=b + 1, start_column=c1, end_row=b + 1, end_column=c1 + 2)
    checks = [
        ("max |M - export|", f"=MAX({rng(g['dM'][0], g['dM'][1], g['dM'][0] + 3, g['dM'][1] + 3)})", "0.00E+00"),
        ("max |World - export|", f"=MAX({rng(g['dW'][0], g['dW'][1], g['dW'][0] + 3, g['dW'][1] + 3)})", "0.00E+00"),
        ("inside clip volume |x|,|y|,|z| <= w",
         f'=IF(AND(ABS({pc[0]})<={w},ABS({pc[1]})<={w},ABS({pc[2]})<={w}),"yes","no")', None),
        ("eye-space distance -z_eye (m)", f"=-{pe[2]}", None),
    ]
    rr = b + 1
    for lab, f, fmt in checks:
        rr += 1
        S(rr, c1, lab, bold=True)
        S(rr, c1 + 1, "", track=False)
        S(rr, c1 + 2, "", track=False)
        s.ws.merge_cells(start_row=rr, start_column=c1, end_row=rr, end_column=c1 + 2)
        S(rr, c1 + 3, f, fmt=fmt, center=True)
    s.note(r + 1, 7, "Exported matrices are float32: differences of ~1e-7 are rounding. "
                     "Window: x = (ndc.x/2+0.5) W, y = (1-(ndc.y/2+0.5)) H, d = ndc.z/2+0.5 (glDepthRange 0..1).")
    s.finish(last_col=20)


def sheet_clipping(wb):
    s = Sheet(wb, "Clipping", "Clipping: Cohen-Sutherland, Liang-Barsky, Sutherland-Hodgman (textbook worked example)",
              "Worked textbook example (window and lines chosen for teaching). The code itself clips on the GPU "
              "(clip volume -w <= x, y, z <= w) and culls on the CPU with math/Frustum (6 planes, sphere test).")
    s.widths({"A": 9, "B": 8, "C": 8, "D": 8, "E": 8, "F": 8, "G": 8, "H": 8, "I": 8, "J": 8, "K": 8, "L": 8,
              "M": 8, "N": 8, "O": 8, "P": 8, "Q": 8, "R": 9, "S": 9, "T": 9, "U": 10})
    S = s.put
    r = 5
    S(r, 1, "Window", bold=True)
    labs = ["xmin", "xmax", "ymin", "ymax"]
    vals = [0, 10, 0, 8]
    for j, (lab, v) in enumerate(zip(labs, vals)):
        S(r, 2 + 2 * j, lab, bold=True, center=True)
        S(r, 3 + 2 * j, v, inp=True, fmt="0")
    xmin, xmax, ymin, ymax = (A(r, 3 + 2 * j, True) for j in range(4))
    s.note(r, 11, "Outcode bits T B R L = (y > ymax)(y < ymin)(x > xmax)(x < xmin); code = 8T + 4B + 2R + L")
    r += 2
    S(r, 1, "1. Cohen-Sutherland outcodes", border=False, bold=True)
    r += 1
    s.heads(r, 1, ["line", "x0", "y0", "x1", "y1", "T0", "B0", "R0", "L0", "code0", "T1", "B1", "R1", "L1", "code1",
                   "bits 0", "bits 1", "AND of codes", "decision"], height=30)
    lines = [("L1", 2, 2, 8, 6), ("L2", -2, -1, -1, 9), ("L3", -3, 4, 6, 12), ("L4", -2, 3, 12, 5)]
    lrows = []
    for name, x0, y0, x1, y1 in lines:
        r += 1
        lrows.append(r)
        S(r, 1, name, bold=True, center=True)
        for j, v in enumerate((x0, y0, x1, y1)):
            S(r, 2 + j, v, inp=True, fmt="0")
        for e, (xc, yc) in enumerate(((2, 3), (4, 5))):
            x, y = A(r, xc), A(r, yc)
            base = 6 + 5 * e
            S(r, base, f"=IF({y}>{ymax},1,0)", fmt="0", center=True)
            S(r, base + 1, f"=IF({y}<{ymin},1,0)", fmt="0", center=True)
            S(r, base + 2, f"=IF({x}>{xmax},1,0)", fmt="0", center=True)
            S(r, base + 3, f"=IF({x}<{xmin},1,0)", fmt="0", center=True)
            S(r, base + 4, f"=8*{A(r, base)}+4*{A(r, base + 1)}+2*{A(r, base + 2)}+{A(r, base + 3)}", fmt="0",
              center=True)
        S(r, 16, f"={A(r, 6)}&{A(r, 7)}&{A(r, 8)}&{A(r, 9)}", center=True)
        S(r, 17, f"={A(r, 11)}&{A(r, 12)}&{A(r, 13)}&{A(r, 14)}", center=True)
        S(r, 18, f"={A(r, 6)}*{A(r, 11)}+{A(r, 7)}*{A(r, 12)}+{A(r, 8)}*{A(r, 13)}+{A(r, 9)}*{A(r, 14)}", fmt="0",
          center=True)
        S(r, 19, f'=IF({A(r, 10)}+{A(r, 15)}=0,"accept (both inside)",IF({A(r, 18)}>0,"reject (same side)",'
                 f'"clip (subdivide)"))')
        s.ws.merge_cells(start_row=r, start_column=19, end_row=r, end_column=21)
    r += 2
    S(r, 1, "2. Liang-Barsky: x = x0 + t dx, y = y0 + t dy; p1 = -dx, q1 = x0 - xmin; p2 = dx, q2 = xmax - x0; "
            "p3 = -dy, q3 = y0 - ymin; p4 = dy, q4 = ymax - y0", border=False, bold=True)
    r += 1
    s.heads(r, 1, ["line", "dx", "dy", "p1", "p2", "p3", "p4", "q1", "q2", "q3", "q4", "t enter = max(0, q/p, p<0)",
                   "t exit = min(1, q/p, p>0)", "parallel outside", "visible", "x0'", "y0'", "x1'", "y1'"], height=44)
    for lr in lrows:
        r += 1
        x0, y0, x1, y1 = (A(lr, 2 + j) for j in range(4))
        S(r, 1, f"={A(lr, 1)}", bold=True, center=True)
        S(r, 2, f"={x1}-{x0}", fmt="0")
        S(r, 3, f"={y1}-{y0}", fmt="0")
        dx, dy = A(r, 2), A(r, 3)
        for j, f in enumerate([f"=-{dx}", f"={dx}", f"=-{dy}", f"={dy}"]):
            S(r, 4 + j, f, fmt="0")
        for j, f in enumerate([f"={x0}-{xmin}", f"={xmax}-{x0}", f"={y0}-{ymin}", f"={ymax}-{y0}"]):
            S(r, 8 + j, f, fmt="0")
        ps = [A(r, 4 + j) for j in range(4)]
        qs = [A(r, 8 + j) for j in range(4)]
        S(r, 12, "=MAX(0," + ",".join(f"IF({p}<0,{q}/{p},0)" for p, q in zip(ps, qs)) + ")")
        S(r, 13, "=MIN(1," + ",".join(f"IF({p}>0,{q}/{p},1)" for p, q in zip(ps, qs)) + ")")
        S(r, 14, "=IF(OR(" + ",".join(f"AND({p}=0,{q}<0)" for p, q in zip(ps, qs)) + '),"yes","no")', center=True)
        S(r, 15, f'=IF(AND({A(r, 14)}="no",{A(r, 12)}<={A(r, 13)}),"yes","no")', center=True)
        te, tx = A(r, 12), A(r, 13)
        vis = A(r, 15)
        for j, f in enumerate([f"{x0}+{te}*{dx}", f"{y0}+{te}*{dy}", f"{x0}+{tx}*{dx}", f"{y0}+{tx}*{dy}"]):
            S(r, 16 + j, f'=IF({vis}="yes",{f},"-")')
    r += 2
    S(r, 1, "3. Sutherland-Hodgman: triangle clipped against the 4 window edges in turn (output list -> next "
            "stage input)", border=False, bold=True)
    r += 1
    tri = [(-2, 3), (7, -2), (12, 10)]
    S(r, 1, "Triangle", bold=True)
    for j, (x, y) in enumerate(tri):
        S(r, 2 + 3 * j, "ABC"[j], bold=True, center=True)
        S(r, 3 + 3 * j, x, inp=True, fmt="0")
        S(r, 4 + 3 * j, y, inp=True, fmt="0")
    poly = [(A(r, 3 + 3 * j), A(r, 4 + 3 * j)) for j in range(3)]
    polyv = [tuple(map(float, p)) for p in tri]
    r += 1
    # edges: (name, inside test, axis value, coordinate index)
    edges = [("left x >= xmin", lambda x, y: x >= 0, xmin, 0, ">="), ("right x <= xmax", lambda x, y: x <= 10, xmax, 0, "<="),
             ("bottom y >= ymin", lambda x, y: y >= 0, ymin, 1, ">="), ("top y <= ymax", lambda x, y: y <= 8, ymax, 1, "<=")]
    bounds = {xmin: 0, xmax: 10, ymin: 0, ymax: 8}
    # two edge stages side by side per band
    col_starts = [1, 12]
    band_r, max_r = r, r
    for k, (ename, inside, bnd, axis, op) in enumerate(edges):
        c0 = col_starts[k % 2]
        if k == 2:
            band_r = max_r + 2
        rr = band_r
        S(rr, c0, f"Stage {k + 1}: {ename}", border=False, bold=True)
        rr += 1
        s.heads(rr, c0, ["edge", "Sx", "Sy", "Px", "Py", "S in", "P in", "Ix", "Iy", "output"], height=18)
        n = len(poly)
        out, outv = [], []
        for i in range(n):
            rr += 1
            sxa, sya = poly[i]
            pxa, pya = poly[(i + 1) % n]
            sv, pv = polyv[i], polyv[(i + 1) % n]
            S(rr, c0, f"{i + 1}-{(i + 1) % n + 1}", center=True)
            for j, a in enumerate((sxa, sya, pxa, pya)):
                S(rr, c0 + 1 + j, f"={a}")
            sc = sxa if axis == 0 else sya
            pc = pxa if axis == 0 else pya
            S(rr, c0 + 5, f"=IF({A(rr, c0 + 1 + axis)}{op}{bnd},1,0)", fmt="0", center=True)
            S(rr, c0 + 6, f"=IF({A(rr, c0 + 3 + axis)}{op}{bnd},1,0)", fmt="0", center=True)
            Sx_, Sy_, Px_, Py_ = (A(rr, c0 + 1 + j) for j in range(4))
            sin_, pin_ = A(rr, c0 + 5), A(rr, c0 + 6)
            if axis == 0:
                ix = f"{bnd}"
                iy = f"{Sy_}+({bnd}-{Sx_})*({Py_}-{Sy_})/({Px_}-{Sx_})"
            else:
                ix = f"{Sx_}+({bnd}-{Sy_})*({Px_}-{Sx_})/({Py_}-{Sy_})"
                iy = f"{bnd}"
            S(rr, c0 + 7, f'=IF({sin_}<>{pin_},{ix},"-")')
            S(rr, c0 + 8, f'=IF({sin_}<>{pin_},{iy},"-")')
            S(rr, c0 + 9, f'=IF({sin_}+{pin_}=2,"P",IF({pin_}=1,"I, P",IF({sin_}=1,"I","none")))', center=True)
            # python-side structure of the output list (values decide which cells are referenced)
            si, pi_ = inside(*sv), inside(*pv)
            b_ = bounds[bnd]
            if si != pi_:
                if axis == 0:
                    iv = (b_, sv[1] + (b_ - sv[0]) * (pv[1] - sv[1]) / (pv[0] - sv[0]))
                else:
                    iv = (sv[0] + (b_ - sv[1]) * (pv[0] - sv[0]) / (pv[1] - sv[1]), b_)
            if si and pi_:
                out.append((Px_, Py_)); outv.append(pv)
            elif si and not pi_:
                out.append((A(rr, c0 + 7), A(rr, c0 + 8))); outv.append(iv)
            elif not si and pi_:
                out.append((A(rr, c0 + 7), A(rr, c0 + 8))); outv.append(iv)
                out.append((Px_, Py_)); outv.append(pv)
        rr += 1
        S(rr, c0, "out", bold=True, center=True)
        txt = "=" + '&" "&'.join(f'"("&ROUND({x},2)&", "&ROUND({y},2)&")"' for x, y in out)
        S(rr, c0 + 1, txt)
        s.ws.merge_cells(start_row=rr, start_column=c0 + 1, end_row=rr, end_column=c0 + 9)
        S(rr + 1, c0, f"{len(out)} vertices", border=False, italic=True)
        poly, polyv = out, outv
        max_r = max(max_r, rr + 1)
    s.note(max_r + 1, 1, "Each stage keeps the output-list rule: in-in -> P; in-out -> I; out-in -> I, P; out-out -> none. "
                         "The final list is the clipped polygon (fan-triangulated afterwards).")
    s.finish(last_col=21)


def sheet_illumination(wb, exp, vp):
    s = Sheet(wb, "Illumination", "Illumination: ambient + Lambert diffuse + Blinn-Phong specular at one roof point",
              "scene/DayCycle (solar model, transmittance, 3.2 smoothstep(-1, 10, e), ambient mixLog); export "
              "'lights'; ShaderSources.cpp brdf(); BNK-171 'Roof slab A'.")
    s.widths({"A": 36, "B": 11, "C": 11, "D": 11, "E": 11, "F": 11, "G": 3, "H": 33, "I": 9, "J": 9, "K": 9})
    S = s.put
    L = exp["lights"]
    roof = comp_index(exp)["BNK-171"]
    roof_top = roof["T"][1] + roof["S"][1] / 2
    # ---- inputs on the right
    c0 = 8
    r = 5
    s.heads(r, c0, ["Input (from the code)", "x / R", "y / G", "z / B"])
    inp = {}

    def irow(lab, vals):
        nonlocal r
        r += 1
        S(r, c0, lab, bold=True)
        for j, v in enumerate(vals):
            S(r, c0 + 1 + j, v, inp=True)
        inp[lab] = [A(r, c0 + 1 + j, True) for j in range(len(vals))]

    irow("latitude phi, declination delta (deg)", [22.8, -4.6])
    irow("sun intensity, key above (deg)", [3.2, -1.0])
    irow("optical depth tau (R, G, B)", [0.06, 0.11, 0.21])
    irow("moon direction (travel)", L["moon"]["direction"])
    irow("moon colour", L["moon"]["color"])
    irow("moon intensity", [L["moon"]["intensity"]])
    irow("spot position (m)", [round(x, 4) for x in L["spot"]["position"]])
    irow("spot direction", [round(x, 4) for x in L["spot"]["direction"]])
    irow("spot colour", L["spot"]["color"])
    irow("spot intensity, inner, outer (deg)", [L["spot"]["intensity"], L["spot"]["inner_deg"], L["spot"]["outer_deg"]])
    irow("spot attenuation linear, quadratic", [L["spot"]["linear"], L["spot"]["quadratic"]])
    irow("night ambient (LightRig)", L["ambient"])
    irow("day ambient (kDayAmbient)", [0.34, 0.36, 0.40])
    irow("roof albedo kd (BNK-171 colour)", roof["color"])
    irow("roof ks, shininess n", [roof["specular"], round(roof["shininess"], 4)])
    irow("roof point P (top face, m)", [roof["T"][0], round(roof_top, 4), round(roof["T"][2], 4)])
    irow("normal N", [0, 1, 0])
    r += 1
    S(r, c0, "viewer: orbit eye (View_Projection)", bold=True)
    eye_r = vp["eye"]
    for j in range(3):
        S(r, c0 + 1 + j, f"={A(eye_r, 2 + j, True, sheet='View_Projection')}")
    inp["eye"] = [A(r, c0 + 1 + j, True) for j in range(3)]
    r += 2
    for line in ["Model (display-referred Blinn-Phong path), per channel:",
                 "  I = kd Ia + kd Il (N.L) + ks Il (N.H)^n",
                 "Not modelled: shadows, AO, relief, HDR sky SH ambient,",
                 "  moon tint.",
                 "Spot cone (shader): c = clamp((cos a - cos outer) /",
                 "  (cos inner - cos outer)), weight c^2 (not smoothstep).",
                 "Roof point: 20 deg off the spot axis at t = 0, outside",
                 "  the 14 deg cone. Column E: the up-facing point on the",
                 "  beam axis at roof height.",
                 "Sun key direction is quantised to 1 deg in the code."]:
        s.note(r, c0, line)
        r += 1
    ph, de = inp["latitude phi, declination delta (deg)"]
    Isun, keyAbove = inp["sun intensity, key above (deg)"]
    # ---- scenario columns
    r = 5
    s.heads(r, 1, ["Quantity", "Noon", "Sunset", "Night (moon)", "UAV spot (beam axis)", "UAV spot (roof point)"],
            height=30)
    kinds = ["sun", "sun", "moon", "spot", "spot"]
    hours = [12.0, 17.5, 23.0, 23.0, 23.0]
    R = {}

    def row(lab, fn, fmt=None, bold=True):
        nonlocal r
        r += 1
        S(r, 1, lab, bold=bold)
        for k in range(5):
            v = fn(k, 2 + k)
            S(r, 2 + k, v, fmt=fmt, inp=(lab.startswith("hour")), center=isinstance(v, str) and not v.startswith("="))
        R[lab] = r
        return r

    def at(lab, k):
        return A(R[lab], 2 + k)

    row("hour (clock)", lambda k, c: hours[k], fmt="0.00")
    row("H = 15 (hour - 12) (deg)", lambda k, c: f"=15*({at('hour (clock)', k)}-12)")
    row("sin e = sin phi sin d + cos phi cos d cos H",
        lambda k, c: f"=SIN(RADIANS({ph}))*SIN(RADIANS({de}))+COS(RADIANS({ph}))*COS(RADIANS({de}))*COS(RADIANS({at('H = 15 (hour - 12) (deg)', k)}))")
    row("sun elevation e (deg)", lambda k, c: f"=DEGREES(ASIN({at('sin e = sin phi sin d + cos phi cos d cos H', k)}))")
    e = lambda k: at("sun elevation e (deg)", k)
    row("daylight w = smoothstep(-8, 6, e)", lambda k, c: "=" + smooth(e(k), -8, 6))
    Hc = lambda k: f"RADIANS({at('H = 15 (hour - 12) (deg)', k)})"
    P = inp["roof point P (top face, m)"]
    sp, sd = inp["spot position (m)"], inp["spot direction"]
    # surface point
    tb = f"(({P[1]}-{sp[1]})/{sd[1]})"
    for j, ax in enumerate("xyz"):
        row(f"surface point P.{ax} (m)",
            lambda k, c, j=j: (f"={sp[j]}+{tb}*{sd[j]}" if (k == 3 and j != 1) else f"={P[j]}"))

    def Lraw(k, j):
        if kinds[k] == "sun":
            return [f"=-COS(RADIANS({de}))*SIN({Hc(k)})", f"={at('sin e = sin phi sin d + cos phi cos d cos H', k)}",
                    f"=COS(RADIANS({ph}))*SIN(RADIANS({de}))-SIN(RADIANS({ph}))*COS(RADIANS({de}))*COS({Hc(k)})"][j]
        if kinds[k] == "moon":
            return f"=-{inp['moon direction (travel)'][j]}"
        return f"={sp[j]}-{at('surface point P.' + 'xyz'[j] + ' (m)', k)}"

    for j, ax in enumerate("xyz"):
        row(f"to light (unnormalised) {ax}", lambda k, c, j=j: Lraw(k, j))
    lr = [R[f"to light (unnormalised) {ax}"] for ax in "xyz"]
    row("distance d = |to light| (m)", lambda k, c: f"=SQRT({A(lr[0], c)}^2+{A(lr[1], c)}^2+{A(lr[2], c)}^2)")
    dd = lambda k: at("distance d = |to light| (m)", k)
    for j, ax in enumerate("xyz"):
        row(f"L.{ax} (unit)", lambda k, c, j=j: f"={A(lr[j], c)}/{dd(k)}")
    Lu = lambda k, j: at(f"L.{'xyz'[j]} (unit)", k)
    row("angle to spot axis (deg)",
        lambda k, c: (f"=DEGREES(ACOS(-({Lu(k, 0)}*{sd[0]}+{Lu(k, 1)}*{sd[1]}+{Lu(k, 2)}*{sd[2]})))"
                      if kinds[k] == "spot" else "-"))
    si, ci, co = inp["spot intensity, inner, outer (deg)"]
    row("spot cone c^2, c = clamp((cos a - cos o)/(cos i - cos o))",
        lambda k, c: (f"=MIN(MAX((COS(RADIANS({at('angle to spot axis (deg)', k)}))-COS(RADIANS({co})))/"
                      f"(COS(RADIANS({ci}))-COS(RADIANS({co}))),0),1)^2" if kinds[k] == "spot" else "-"))
    kl, kq = inp["spot attenuation linear, quadratic"]
    row("attenuation 1 / (1 + kl d + kq d^2)",
        lambda k, c: (f"=1/(1+{kl}*{dd(k)}+{kq}*{dd(k)}^2)" if kinds[k] == "spot" else "-"))
    row("air mass m (Kasten-Young)",
        lambda k, c: (f"=1/(SIN(RADIANS(MAX({e(k)},-1)))+0.50572*(MAX({e(k)},-1)+6.07995)^(-1.6364))"
                      if kinds[k] == "sun" else "-"))
    row("light intensity (scalar)",
        lambda k, c: {"sun": f"={Isun}*" + smooth(e(k), keyAbove, 10),
                      "moon": f"={inp['moon intensity'][0]}",
                      "spot": f"={si}*{at('spot cone c^2, c = clamp((cos a - cos o)/(cos i - cos o))', k)}"
                              f"*{at('attenuation 1 / (1 + kl d + kq d^2)', k)}"}[kinds[k]])
    N = inp["normal N"]
    row("N . L (clamped at 0)", lambda k, c: f"=MAX(0,{N[0]}*{Lu(k, 0)}+{N[1]}*{Lu(k, 1)}+{N[2]}*{Lu(k, 2)})")
    ey = inp["eye"]
    vr = lambda k, j: f"({ey[j]}-{at('surface point P.' + 'xyz'[j] + ' (m)', k)})"
    vlen = lambda k: f"SQRT({vr(k, 0)}^2+{vr(k, 1)}^2+{vr(k, 2)}^2)"
    for j, ax in enumerate("xyz"):
        row(f"H = normalize(L + V) {ax}", lambda k, c, j=j: f"={Lu(k, j)}+{vr(k, j)}/{vlen(k)}")
    hr = [R[f"H = normalize(L + V) {ax}"] for ax in "xyz"]
    row("N . H", lambda k, c: f"=MAX(0,({N[0]}*{A(hr[0], c)}+{N[1]}*{A(hr[1], c)}+{N[2]}*{A(hr[2], c)})/"
                              f"SQRT({A(hr[0], c)}^2+{A(hr[1], c)}^2+{A(hr[2], c)}^2))")
    ks, nn = inp["roof ks, shininess n"]
    row("(N . H)^n (0 when N . L = 0)", lambda k, c: f"=IF({at('N . L (clamped at 0)', k)}>0,{at('N . H', k)}^{nn},0)")
    S(R["H = normalize(L + V) x"], 1, "L + V (H before normalising) x", bold=True)
    S(R["H = normalize(L + V) y"], 1, "L + V y", bold=True)
    S(R["H = normalize(L + V) z"], 1, "L + V z", bold=True)

    # ---- RGB results (rows = term x channel, columns = scenarios)
    r += 2
    S(r, 1, "Per-channel terms (light colour = transmittance for the sun)", border=False, bold=True)
    r += 1
    s.heads(r, 1, ["Term (channel)"] + ["Noon", "Sunset", "Night (moon)", "UAV spot (beam axis)",
                                        "UAV spot (roof point)"], height=30)
    tau = inp["optical depth tau (R, G, B)"]
    na, da = inp["night ambient (LightRig)"], inp["day ambient (kDayAmbient)"]
    kd = inp["roof albedo kd (BNK-171 colour)"]
    T = {}
    terms = [("Il = light colour x intensity", "Il"), ("ambient Ia", "Ia"), ("ambient kd Ia", "amb"),
             ("diffuse kd Il (N.L)", "dif"), ("specular ks Il (N.H)^n", "spe"), ("total I", "tot")]
    for lab, key in terms:
        for j, ch in enumerate("RGB"):
            r += 1
            T[(key, j)] = r
            S(r, 1, f"{lab}  {ch}", bold=(key == "tot"))
            for k in range(5):
                c = 2 + k
                I = at("light intensity (scalar)", k)
                w = at("daylight w = smoothstep(-8, 6, e)", k)
                if key == "Il":
                    col = {"sun": f"EXP(-{tau[j]}*({at('air mass m (Kasten-Young)', k)}-1))",
                           "moon": inp["moon colour"][j], "spot": inp["spot colour"][j]}[kinds[k]]
                    f = f"={col}*{I}"
                elif key == "Ia":
                    f = f"={na[j]}^(1-{w})*{da[j]}^{w}"
                elif key == "amb":
                    f = f"={kd[j]}*{A(T[('Ia', j)], c)}"
                elif key == "dif":
                    f = f"={kd[j]}*{A(T[('Il', j)], c)}*{at('N . L (clamped at 0)', k)}"
                elif key == "spe":
                    f = f"={ks}*{A(T[('Il', j)], c)}*{at('(N . H)^n (0 when N . L = 0)', k)}"
                else:
                    f = f"={A(T[('amb', j)], c)}+{A(T[('dif', j)], c)}+{A(T[('spe', j)], c)}"
                S(r, c, f, bold=(key == "tot"))
    r += 1
    S(r, 1, "luminance Y of total (Rec. 709)", bold=True)
    for k in range(5):
        c = 2 + k
        S(r, c, f"=0.2126*{A(T[('tot', 0)], c)}+0.7152*{A(T[('tot', 1)], c)}+0.0722*{A(T[('tot', 2)], c)}", bold=True)
    s.note(r + 1, 1, "Ambient Ia = night^(1-w) day^w (DayCycle mixLog); spot columns: beams on at night only "
                     "(sun below -2 deg). Linear HDR units (the sun's irradiance is 3.2).")
    s.finish(last_col=11)


def sheet_color_interp(wb, exp):
    s = Sheet(wb, "Color_Interp", "Color_Interp: barycentric weights, Gouraud vs Phong, perspective-correct "
                                  "interpolation, sRGB blending, sky colour over the day",
              "math/ExplorerMath edgeFunction / interpolateAffine / interpolatePerspective / srgbEncode (sections A-D: "
              "worked example values); section E: DayCycle solar model + export v6.day_cycle.grid (sky_zenith).")
    s.widths({"A": 16, "B": 9, "C": 9, "D": 9, "E": 9, "F": 9, "G": 9, "H": 9, "I": 9, "J": 3, "K": 10, "L": 9,
              "M": 9, "N": 9, "O": 9, "P": 9, "Q": 9, "R": 9, "S": 9, "T": 9, "U": 9})
    S = s.put
    r = 5
    S(r, 1, "A. Screen triangle (worked example) with per-vertex data", border=False, bold=True)
    r += 1
    s.heads(r, 1, ["vertex", "x (px)", "y (px)", "clip w", "tex u", "N.x", "N.y", "N.z", "I vertex"], height=18)
    verts = [(100, 100, 1.0, 0.0, (-0.6, -0.4, 0.69)), (400, 150, 4.0, 1.0, (0.6, -0.3, 0.74)),
             (220, 380, 10.0, 0.5, (0.0, 0.6, 0.8))]
    vrow = []
    for i, (x, y, w, u, nvec) in enumerate(verts):
        r += 1
        vrow.append(r)
        ln = math.sqrt(sum(c * c for c in nvec))
        nn = [round(c / ln, 4) for c in nvec]
        S(r, 1, f"V{i}", bold=True, center=True)
        for j, v in enumerate([x, y, w, u] + nn):
            S(r, 2 + j, v, inp=True, fmt="0" if j < 2 else NUM)
    r += 1
    S(r, 1, "light", bold=True)
    S(r, 2, "L = V = (0, 0, 1), H = (0, 0, 1)", border=False)
    r += 1
    S(r, 1, "kd, ks, n", bold=True)
    for j, v in enumerate([0.6, 0.6, 40]):
        S(r, 2 + j, v, inp=True, fmt=NUM if j < 2 else "0")
    kd, ks, nexp = (A(r, 2 + j, True) for j in range(3))
    for vr_ in vrow:  # I = kd N.L + ks (N.H)^n with L = H = z
        S(vr_, 9, f"={kd}*MAX(0,{A(vr_, 8)})+{ks}*MAX(0,{A(vr_, 8)})^{nexp}")
    X = [A(v, 2, True) for v in vrow]
    Y = [A(v, 3, True) for v in vrow]
    r += 1
    S(r, 1, "2 x area (edge fn)", bold=True)
    S(r, 2, f"=({X[1]}-{X[0]})*({Y[2]}-{Y[0]})-({Y[1]}-{Y[0]})*({X[2]}-{X[0]})", fmt="0")
    area = A(r, 2, True)
    r += 2
    S(r, 1, "B. Barycentric weights at sample points: w0 = E12(p) / area, w1 = E20(p) / area, w2 = E01(p) / area",
      border=False, bold=True)
    r += 1
    s.heads(r, 1, ["sample", "px", "py", "w0", "w1", "w2", "sum", "inside"], height=18)
    samples = [("centroid", 240, 210), ("near V1", 360, 160), ("edge V0-V2 mid", 160, 240), ("centre-left", 200, 180),
               ("outside", 420, 330)]
    srows = []

    def edge(a, b_, px, py):
        return f"(({X[b_]}-{X[a]})*({py}-{Y[a]})-({Y[b_]}-{Y[a]})*({px}-{X[a]}))"

    for name, px, py in samples:
        r += 1
        srows.append(r)
        S(r, 1, name, bold=True)
        S(r, 2, px, inp=True, fmt="0")
        S(r, 3, py, inp=True, fmt="0")
        p_x, p_y = A(r, 2), A(r, 3)
        S(r, 4, f"={edge(1, 2, p_x, p_y)}/{area}")
        S(r, 5, f"={edge(2, 0, p_x, p_y)}/{area}")
        S(r, 6, f"={edge(0, 1, p_x, p_y)}/{area}")
        S(r, 7, f"={A(r, 4)}+{A(r, 5)}+{A(r, 6)}")
        S(r, 8, f'=IF(AND({A(r, 4)}>=0,{A(r, 5)}>=0,{A(r, 6)}>=0),"yes","no")', center=True)
    r += 2
    S(r, 1, "C. Gouraud (interpolate vertex I) vs Phong (interpolate N, normalise, then light)", border=False, bold=True)
    r += 1
    s.heads(r, 1, ["sample", "I Gouraud", "N.x", "N.y", "N.z", "|N|", "I Phong", "Phong - Gouraud"], height=18)
    for sr in srows[:4]:
        r += 1
        S(r, 1, f"={A(sr, 1)}", bold=True)
        w = [A(sr, 4 + i) for i in range(3)]
        S(r, 2, "=" + "+".join(f"{w[i]}*{A(vrow[i], 9, True)}" for i in range(3)))
        for j in range(3):
            S(r, 3 + j, "=" + "+".join(f"{w[i]}*{A(vrow[i], 6 + j, True)}" for i in range(3)))
        S(r, 6, f"=SQRT({A(r, 3)}^2+{A(r, 4)}^2+{A(r, 5)}^2)")
        nz = f"({A(r, 5)}/{A(r, 6)})"
        S(r, 7, f"={kd}*MAX(0,{nz})+{ks}*MAX(0,{nz})^{nexp}")
        S(r, 8, f"={A(r, 7)}-{A(r, 2)}")
    r += 1
    s.note(r, 1, "The highlight sits inside the triangle (interpolated N close to H): Phong finds it, Gouraud misses it.")
    r += 2
    S(r, 1, "D. Texture coordinate u: affine vs perspective-correct (clip w per vertex)", border=False, bold=True)
    r += 1
    s.heads(r, 1, ["sample", "u affine", "u persp.", "error"], height=18)
    for sr in srows[:4]:
        r += 1
        S(r, 1, f"={A(sr, 1)}", bold=True)
        w = [A(sr, 4 + i) for i in range(3)]
        uu = [A(vrow[i], 5, True) for i in range(3)]
        ww = [A(vrow[i], 4, True) for i in range(3)]
        S(r, 2, "=" + "+".join(f"{w[i]}*{uu[i]}" for i in range(3)))
        S(r, 3, "=(" + "+".join(f"{w[i]}*{uu[i]}/{ww[i]}" for i in range(3)) + ")/(" +
          "+".join(f"{w[i]}/{ww[i]}" for i in range(3)) + ")")
        S(r, 4, f"={A(r, 2)}-{A(r, 3)}")
    r += 1
    s.note(r, 1, "Perspective-correct: interpolate u / w and 1 / w linearly in screen space, then divide.")
    left_end = r

    # ---- right side: sRGB blend
    c0 = 11
    r = 5
    S(r, c0, "E. Blending two sRGB colours: naive (in sRGB) vs correct (in linear light)", border=False, bold=True)
    r += 1
    s.heads(r, c0, ["colour", "R (0-255)", "G", "B", "R linear", "G linear", "B linear"], height=18)
    cols = [("A", (255, 64, 0)), ("B", (0, 128, 255))]
    crow = []
    for name, rgb in cols:
        r += 1
        crow.append(r)
        S(r, c0, name, bold=True, center=True)
        for j, v in enumerate(rgb):
            S(r, c0 + 1 + j, v, inp=True, fmt="0")
            c = f"({A(r, c0 + 1 + j)}/255)"
            S(r, c0 + 4 + j, f"=IF({c}<=0.04045,{c}/12.92,(({c}+0.055)/1.055)^2.4)")
    r += 1
    s.heads(r, c0, ["t", "naive R", "naive G", "naive B", "linear R", "linear G", "linear B", "Y naive", "Y linear"],
            height=18)
    for t in (0, 0.25, 0.5, 0.75, 1.0):
        r += 1
        S(r, c0, t, inp=True, fmt="0.00")
        tt = A(r, c0)
        for j in range(3):
            a, b_ = A(crow[0], c0 + 1 + j, True), A(crow[1], c0 + 1 + j, True)
            S(r, c0 + 1 + j, f"=(1-{tt})*{a}+{tt}*{b_}", fmt="0.0")
            la, lb = A(crow[0], c0 + 4 + j, True), A(crow[1], c0 + 4 + j, True)
            x = f"((1-{tt})*{la}+{tt}*{lb})"
            S(r, c0 + 4 + j, f"=255*IF({x}<=0.0031308,12.92*{x},1.055*{x}^(1/2.4)-0.055)", fmt="0.0")
        lin = lambda j, base: (f"IF({A(r, base + j)}/255<=0.04045,{A(r, base + j)}/255/12.92,"
                               f"(({A(r, base + j)}/255+0.055)/1.055)^2.4)")
        S(r, c0 + 7, "=" + "+".join(f"{wt}*{lin(j, c0 + 1)}" for j, wt in enumerate((0.2126, 0.7152, 0.0722))))
        S(r, c0 + 8, "=" + "+".join(f"{wt}*{lin(j, c0 + 4)}" for j, wt in enumerate((0.2126, 0.7152, 0.0722))))
    r += 1
    s.note(r, c0, "sRGB decode: c <= 0.04045 ? c/12.92 : ((c+0.055)/1.055)^2.4; encode = srgbEncode (ExplorerMath).")
    r += 1
    s.note(r, c0, "Y = relative luminance (Rec. 709). The naive mid-point is darker than the light-correct one.")

    # ---- sky over the day
    r += 2
    S(r, c0, "F. Sky zenith colour over the day (solar model formulas; sky colour from the export)", border=False,
      bold=True)
    r += 1
    s.heads(r, c0, ["hour", "e formula (deg)", "e export (deg)", "diff", "daylight w", "phase (no hysteresis)",
                    "zenith R", "zenith G", "zenith B", "Y zenith", "phase export"], height=30)
    grid = {round(g["hour"], 2): g for g in exp["v6"]["day_cycle"]["grid"]}
    hours = [0, 4, 5.5, 6, 6.5, 7, 9, 12, 15, 17, 17.5, 18, 18.5, 20, 22]
    phi, dec = 22.8, -4.6
    S(r - 1, c0 + 7, "phi, delta:", border=False, bold=True)
    S(r - 1, c0 + 8, phi, inp=True)
    S(r - 1, c0 + 9, dec, inp=True)
    PH, DE = A(r - 1, c0 + 8, True), A(r - 1, c0 + 9, True)
    for h in hours:
        g = grid[round(h, 2)]
        r += 1
        S(r, c0, h, inp=True, fmt="0.00")
        hh = A(r, c0)
        S(r, c0 + 1, f"=DEGREES(ASIN(SIN(RADIANS({PH}))*SIN(RADIANS({DE}))+COS(RADIANS({PH}))*COS(RADIANS({DE}))*"
                     f"COS(RADIANS(15*({hh}-12)))))")
        S(r, c0 + 2, g["elevation_deg"], inp=True)
        S(r, c0 + 3, f"={A(r, c0 + 1)}-{A(r, c0 + 2)}")
        S(r, c0 + 4, "=" + smooth(A(r, c0 + 1), -8, 6))
        e_ = A(r, c0 + 1)
        S(r, c0 + 5, f'=IF({e_}<-6,"NIGHT",IF({e_}>4,"DAY",IF({hh}<12,"DAWN","DUSK")))', center=True)
        for j in range(3):
            S(r, c0 + 6 + j, g["sky_zenith"][j], inp=True)
        S(r, c0 + 9, f"=0.2126*{A(r, c0 + 6)}+0.7152*{A(r, c0 + 7)}+0.0722*{A(r, c0 + 8)}")
        S(r, c0 + 10, g["phase"], inp=True, center=True)
    r += 1
    s.note(r, c0, "Day sky: Preetham analytic sky (math/SkyModel), blended with the night preset in log space "
                  "(mixLog, weight w).")
    r += 1
    s.note(r, c0, "The export applies 0.25 deg hysteresis to the phase, so DAWN / DUSK can last a little longer.")
    s.finish(last_col=21, last_row=max(left_end, r))


def sheet_ray_tracing(wb):
    s = Sheet(wb, "Ray_Tracing", "Ray_Tracing: intersections, reflection, refraction, recursion (worked example)",
              "Worked textbook example. The code: raytrace/Intersect (sphere / box / heightfield), Optics.h reflectDir, "
              "refractDir, fresnelDielectric; RayTracer.h whittedDepth = 5 (';' render modes).")
    s.widths({"A": 34, "B": 10, "C": 10, "D": 10, "E": 3, "F": 34, "G": 10, "H": 10, "I": 10, "J": 10})
    S = s.put
    r = 5
    s.heads(r, 1, ["1. Ray-sphere", "x", "y", "z"])
    ref = {}

    def vrow(lab, vals, inp=False, c0=1):
        nonlocal r
        r += 1
        S(r, c0, lab, bold=True)
        for j, v in enumerate(vals):
            S(r, c0 + 1 + j, v, inp=inp)
        ref[lab] = [A(r, c0 + 1 + j, True) for j in range(len(vals))]
        return ref[lab]

    o = vrow("ray origin o", [0, 2, -10], inp=True)
    draw = vrow("direction (unnormalised)", [0.1, -0.05, 1], inp=True)
    ln = f"SQRT({draw[0]}^2+{draw[1]}^2+{draw[2]}^2)"
    d = vrow("d = normalize", [f"={draw[j]}/{ln}" for j in range(3)])
    c = vrow("sphere centre c", [0, 1, 0], inp=True)
    rad = vrow("radius r", [1.5], inp=True)[0]
    oc = vrow("oc = o - c", [f"={o[j]}-{c[j]}" for j in range(3)])
    dot = lambda a, b_: "+".join(f"{a[j]}*{b_[j]}" for j in range(3))
    a_ = vrow("a = d . d", [f"={dot(d, d)}"])[0]
    b_ = vrow("b = 2 d . oc", [f"=2*({dot(d, oc)})"])[0]
    cc = vrow("c' = oc . oc - r^2", [f"={dot(oc, oc)}-{rad}^2"])[0]
    disc = vrow("discriminant b^2 - 4 a c'", [f"={b_}^2-4*{a_}*{cc}"])[0]
    t1 = vrow("t1 = (-b - sqrt(disc)) / 2a (near)", [f'=IF({disc}>=0,(-{b_}-SQRT({disc}))/(2*{a_}),"miss")'])[0]
    t2 = vrow("t2 = (-b + sqrt(disc)) / 2a (far)", [f'=IF({disc}>=0,(-{b_}+SQRT({disc}))/(2*{a_}),"miss")'])[0]
    p = vrow("hit p = o + t1 d", [f"={o[j]}+{t1}*{d[j]}" for j in range(3)])
    n = vrow("normal n = (p - c) / r", [f"=({p[j]}-{c[j]})/{rad}" for j in range(3)])
    dn = vrow("d . n", [f"={dot(d, n)}"])[0]
    rr_ = vrow("reflection r = d - 2 (d . n) n", [f"={d[j]}-2*{dn}*{n[j]}" for j in range(3)])
    vrow("check |r| (= 1)", [f"=SQRT({rr_[0]}^2+{rr_[1]}^2+{rr_[2]}^2)"])
    r += 2
    s.heads(r, 1, ["2. Ray-plane (ground y = 0)", "x", "y", "z"])
    o2 = vrow("ray origin o", [0, 5, 0], inp=True)
    d2r = vrow("direction (unnormalised)", [1, -1, 0.5], inp=True)
    ln2 = f"SQRT({d2r[0]}^2+{d2r[1]}^2+{d2r[2]}^2)"
    d2 = vrow("d = normalize", [f"={d2r[j]}/{ln2}" for j in range(3)])
    pn = vrow("plane normal n", [0, 1, 0], inp=True)
    p0 = vrow("point on plane p0", [0, 0, 0], inp=True)
    den = vrow("n . d", [f"={dot(pn, d2)}"])[0]
    tt = vrow("t = n . (p0 - o) / (n . d)",
              [f'=IF({den}<>0,(' + "+".join(f"{pn[j]}*({p0[j]}-{o2[j]})" for j in range(3)) + f')/{den},"parallel")'])[0]
    vrow("hit p = o + t d", [f"={o2[j]}+{tt}*{d2[j]}" for j in range(3)])
    rr2 = vrow("reflection r = d - 2 (d . n) n", [f"={d2[j]}-2*{den}*{pn[j]}" for j in range(3)])
    left_end = r

    # ---- right: refraction
    r = 5
    s.heads(r, 6, ["3. Snell refraction (air -> glass)", "x", "y", "z"])
    di = vrow("incident d (unnormalised)", [0.5, -1, 0], inp=True, c0=6)
    lni = f"SQRT({di[0]}^2+{di[1]}^2+{di[2]}^2)"
    du = vrow("d = normalize", [f"={di[j]}/{lni}" for j in range(3)], c0=6)
    nr = vrow("normal n (facing against d)", [0, 1, 0], inp=True, c0=6)
    n12 = vrow("n1, n2", [1.0, 1.5], inp=True, c0=6)
    eta = vrow("eta = n1 / n2", [f"={n12[0]}/{n12[1]}"], c0=6)[0]
    cosi = vrow("cos i = -d . n", [f"=-({dot(du, nr)})"], c0=6)[0]
    k = vrow("k = 1 - eta^2 (1 - cos^2 i)", [f"=1-{eta}^2*(1-{cosi}^2)"], c0=6)[0]
    vrow("t = eta d + (eta cos i - sqrt k) n",
         [f'=IF({k}<0,"TIR",{eta}*{du[j]}+({eta}*{cosi}-SQRT({k}))*{nr[j]})' for j in range(3)], c0=6)
    thi = vrow("incidence angle (deg)", [f"=DEGREES(ACOS({cosi}))"], c0=6)[0]
    tht = vrow("refraction angle (deg)", [f"=DEGREES(ASIN({eta}*SIN(RADIANS({thi}))))"], c0=6)[0]
    vrow("Snell check n1 sin i - n2 sin t", [f"={n12[0]}*SIN(RADIANS({thi}))-{n12[1]}*SIN(RADIANS({tht}))"], c0=6)
    vrow("Fresnel R (exact, fresnelDielectric)",
         [f"=0.5*((({eta}*{cosi}-SQRT({k}))/({eta}*{cosi}+SQRT({k})))^2+(({cosi}-{eta}*SQRT({k}))/({cosi}+{eta}*SQRT({k})))^2)"],
         c0=6)
    vrow("Schlick R0 = ((n1 - n2) / (n1 + n2))^2", [f"=(({n12[0]}-{n12[1]})/({n12[0]}+{n12[1]}))^2"], c0=6)
    r += 2
    S(r, 6, "4. Total internal reflection (glass -> air, eta = 1.5)", border=False, bold=True)
    r += 1
    s.heads(r, 6, ["incidence angle (deg)", "k = 1 - eta^2 sin^2 i", "TIR?", "Fresnel R"])
    r += 1
    S(r, 6, "critical angle asin(1/1.5)", bold=True)
    S(r, 7, f"=DEGREES(ASIN({n12[0]}/{n12[1]}))")
    for ang in (20, 30, 41, 42, 60):
        r += 1
        S(r, 6, ang, inp=True, fmt="0")
        e2 = f"({n12[1]}/{n12[0]})"
        S(r, 7, f"=1-{e2}^2*SIN(RADIANS({A(r, 6)}))^2")
        S(r, 8, f'=IF({A(r, 7)}<0,"yes","no")', center=True)
        ci_ = f"COS(RADIANS({A(r, 6)}))"
        S(r, 9, f"=IF({A(r, 7)}<0,1,0.5*((({e2}*{ci_}-SQRT({A(r, 7)}))/({e2}*{ci_}+SQRT({A(r, 7)})))^2+"
                f"(({ci_}-{e2}*SQRT({A(r, 7)}))/({ci_}+{e2}*SQRT({A(r, 7)})))^2))")
    r += 2
    S(r, 6, "5. Whitted recursion: weight of bounce k = kr^k", border=False, bold=True)
    r += 1
    S(r, 6, "reflectance kr, depth limit (whittedDepth)", bold=True)
    S(r, 7, 0.5, inp=True)
    S(r, 8, 5, inp=True, fmt="0")
    kr, dl = A(r, 7, True), A(r, 8, True)
    r += 1
    s.heads(r, 6, ["bounce k", "kr^k", "cumulative", "share of the infinite sum"])
    for kk in range(0, 7):
        r += 1
        S(r, 6, kk, inp=True, fmt="0")
        S(r, 7, f'=IF({A(r, 6)}<={dl},{kr}^{A(r, 6)},0)')
        S(r, 8, f"=SUM({A(r - kk, 7)}:{A(r, 7)})")
        S(r, 9, f"={A(r, 8)}*(1-{kr})")
    r += 1
    s.note(r, 6, "Bounces past the limit (k > 5) contribute 0; the lost share of the infinite sum is kr^6.")
    s.finish(last_col=10, last_row=max(left_end, r))


def sheet_thermal_kernels(wb):
    s = Sheet(wb, "Thermal_Kernels", "Thermal_Kernels: Gaussian and Sobel kernels on a 5x5 thermal patch",
              "math/ExplorerMath.cpp kernel(): Gaussian 5x5 = binomial [1 4 6 4 1] (x) [1 4 6 4 1] / 256 (sigma ~ 1), "
              "Sobel Gx / Gy, convolve() (correlation, kernel row 0 = row above, clamp-to-edge), applyGamma v^(1/g).")
    for col in range(1, 31):
        s.ws.column_dimensions[get_column_letter(col)].width = 7.2
    s.ws.column_dimensions["A"].width = 3
    S = s.put
    r = 5
    S(r, 2, "sigma", bold=True)
    S(r, 3, 1.0, inp=True)
    sig = A(r, 3, True)
    S(r, 5, "size 5 (offsets -2..2)", border=False)
    S(r, 10, "threshold T on |G|", bold=True)
    s.ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=12)
    S(r, 13, 2.0, inp=True)
    thr = A(r, 13, True)
    b = 7
    raw = [[f"=EXP(-(({x})^2+({y})^2)/(2*{sig}^2))" for x in range(-2, 3)] for y in range(-2, 3)]
    gr = s.grid(b, 2, "Gaussian exp(-(x^2+y^2)/(2 s^2))", raw)
    rr0 = f"${get_column_letter(gr[1])}${gr[0]}:${get_column_letter(gr[1] + 4)}${gr[0] + 4}"
    gn = s.grid(b, 8, "normalised (sum = 1)", [[f"={A(gr[0] + i, gr[1] + j)}/SUM({rr0})" for j in range(5)] for i in range(5)])
    bin_ = [1, 4, 6, 4, 1]
    S(b, 20, "binomial", border=False, bold=True)
    for j, v in enumerate(bin_):
        S(b + 1 + j, 20, v, inp=True, fmt="0")
    bc = [A(b + 1 + j, 20, True) for j in range(5)]
    gc = s.grid(b, 14, "code: b_y b_x / 256", [[f"={bc[i]}*{bc[j]}/256" for j in range(5)] for i in range(5)])
    S(b + 6, 2, "sum", bold=True)
    S(b + 6, 3, f"=SUM({rng(gn[0], gn[1], gn[0] + 4, gn[1] + 4)})")
    S(b + 6, 8, "max |normalised - code|", bold=True)
    s.ws.merge_cells(start_row=b + 6, start_column=8, end_row=b + 6, end_column=11)
    S(b + 6, 12, "=MAX(" + ",".join(f"ABS({A(gn[0] + i, gn[1] + j)}-{A(gc[0] + i, gc[1] + j)})" for i in range(5)
                                    for j in range(5)) + ")")
    S(b + 6, 14, "sum", bold=True)
    S(b + 6, 15, f"=SUM({rng(gc[0], gc[1], gc[0] + 4, gc[1] + 4)})")
    b += 8
    patch = [[0.20, 0.22, 0.25, 0.30, 0.32],
             [0.21, 0.24, 0.60, 0.65, 0.34],
             [0.22, 0.30, 0.85, 0.90, 0.40],
             [0.20, 0.28, 0.70, 0.75, 0.36],
             [0.19, 0.21, 0.26, 0.31, 0.30]]
    pg = s.grid(b, 2, "5x5 thermal patch (worked example, 0..1)", patch, inp=True, fmt="0.00")
    sx = s.grid(b, 8, "Sobel Gx", [[-1, 0, 1], [-2, 0, 2], [-1, 0, 1]], inp=True, fmt="0")
    sy = s.grid(b, 12, "Sobel Gy", [[-1, -2, -1], [0, 0, 0], [1, 2, 1]], inp=True, fmt="0")
    # products for the output pixel at patch (1,1) -> window rows 0..2, cols 0..2  (centre output: patch (2,2))
    oi, oj = 1, 1   # output index in the 3x3 valid result (centre pixel = patch row 2, col 2)
    # window for the valid output (oi, oj) covers patch rows oi..oi+2, cols oj..oj+2
    prod = [[f"={A(pg[0] + oi + i, pg[1] + oj + j)}*{A(sx[0] + i, sx[1] + j)}" for j in range(3)] for i in range(3)]
    pr = s.grid(b, 16, "Gx step: patch x kernel at the centre", prod)
    S(b + 4, 16, "sum", bold=True)
    S(b + 4, 17, f"=SUM({rng(pr[0], pr[1], pr[0] + 2, pr[1] + 2)})")
    S(b + 4, 18, "SUMPRODUCT:", border=False)
    S(b + 4, 20, f"=SUMPRODUCT({rng(pg[0] + oi, pg[1] + oj, pg[0] + oi + 2, pg[1] + oj + 2)},"
                 f"{rng(sx[0], sx[1], sx[0] + 2, sx[1] + 2)})")
    S(b, 21, "", border=False)
    S(b + 6, 2, "Gaussian 5x5 at the centre = SUMPRODUCT(patch, normalised kernel)", bold=True, border=False)
    S(b + 6, 12, f"=SUMPRODUCT({rng(pg[0], pg[1], pg[0] + 4, pg[1] + 4)},{rng(gn[0], gn[1], gn[0] + 4, gn[1] + 4)})")
    S(b + 6, 14, "code kernel:", bold=True, border=False)
    S(b + 6, 16, f"=SUMPRODUCT({rng(pg[0], pg[1], pg[0] + 4, pg[1] + 4)},{rng(gc[0], gc[1], gc[0] + 4, gc[1] + 4)})")
    S(b + 6, 18, "centre input:", bold=True, border=False)
    S(b + 6, 20, f"={A(pg[0] + 2, pg[1] + 2)}")
    b += 8

    def conv(k):
        return [[f"=SUMPRODUCT({rng(pg[0] + i, pg[1] + j, pg[0] + i + 2, pg[1] + j + 2)},"
                 f"{rng(k[0], k[1], k[0] + 2, k[1] + 2, )})" for j in range(3)] for i in range(3)]

    gx = s.grid(b, 2, "Gx (3x3 valid output)", conv(sx))
    gy = s.grid(b, 6, "Gy (3x3 valid output)", conv(sy))
    gm = s.grid(b, 10, "|G| = sqrt(Gx^2 + Gy^2)",
                [[f"=SQRT({A(gx[0] + i, gx[1] + j)}^2+{A(gy[0] + i, gy[1] + j)}^2)" for j in range(3)] for i in range(3)])
    s.grid(b, 14, "edge = |G| >= T", [[f"=IF({A(gm[0] + i, gm[1] + j)}>={thr},1,0)" for j in range(3)] for i in range(3)],
           fmt="0")
    s.note(b + 4, 2, "Valid region: output (i, j) uses patch rows i..i+2, cols j..j+2 (the code clamps at the border).")
    s.note(b + 5, 2, "The Frame Explorer shows Gx, Gy with bias 0.5 (zero = mid-grey); Sobel |G| uses both kernels.")
    # gamma table
    c0 = 23
    S(5, c0, "Gamma: out = v^(1/g)", border=False, bold=True)
    s.heads(6, c0, ["v", "g = 0.5", "g = 1.0", "g = 2.2"])
    for j, gv in enumerate((0.5, 1.0, 2.2)):
        S(7, c0 + 1 + j, gv, inp=True, fmt="0.0")
    S(7, c0, "g:", bold=True)
    for i in range(11):
        rr = 8 + i
        S(rr, c0, i / 10, inp=True, fmt="0.0")
        for j in range(3):
            S(rr, c0 + 1 + j, f"={A(rr, c0)}^(1/{A(7, c0 + 1 + j, True)})")
    s.note(20, c0, "g > 1 lifts dark tones (applyGamma).")
    s.finish(last_col=26)


def sheet_duty_cycle(wb, exp):
    s = Sheet(wb, "Duty_Cycle", "Duty_Cycle: 24-hour roster of the 8 on-duty soldiers (2 shifts, not 6-hour shifts)",
              "scene/Garrison.h kDayShiftStart 6.4, kDayShiftEnd 17.6, dayShift(); Garrison.cpp makePrograms() and "
              "shiftDelay = 1 + 2.5 rank s; rosters from ts --export-docs v6.garrison.soldiers.")
    s.widths({"A": 4, "B": 13})
    for col in range(3, 27):
        s.ws.column_dimensions[get_column_letter(col)].width = 5.6
    S = s.put
    r = 5
    S(r, 2, "day start (h)", bold=True)
    s.ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    S(r, 4, 6.4, inp=True, fmt="0.00")
    S(r, 6, "day end (h)", bold=True)
    s.ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
    S(r, 8, 17.6, inp=True, fmt="0.00")
    st, en = A(r, 4, True), A(r, 8, True)
    S(r, 10, "day shift", bold=True)
    s.ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=11)
    S(r, 12, f'=TEXT({st}/24,"hh:mm")&"-"&TEXT({en}/24,"hh:mm")', center=True)
    s.ws.merge_cells(start_row=r, start_column=12, end_row=r, end_column=14)
    S(r, 15, "hours", bold=True)
    s.ws.merge_cells(start_row=r, start_column=15, end_row=r, end_column=16)
    S(r, 17, f"={en}-{st}", fmt="0.00")
    S(r, 19, "night hours", bold=True)
    s.ws.merge_cells(start_row=r, start_column=19, end_row=r, end_column=20)
    S(r, 21, f"=24-({en}-{st})", fmt="0.00")
    ev = {e["event"]: e["hour"] for e in exp["v6"]["day_cycle"]["events"]}
    S(r, 23, "DayBreak, Dusk", bold=True)
    s.ws.merge_cells(start_row=r, start_column=23, end_row=r, end_column=24)
    S(r, 25, round(ev["DayBreak"], 2), inp=True, fmt="0.00")
    S(r, 26, round(ev["Dusk"], 2), inp=True, fmt="0.00")
    soldiers = exp["v6"]["garrison"]["soldiers"]
    # roster table (rows 8..)
    r = 7
    s.heads(r, 1, ["#", "role"])
    S(r, 3, "day roster (export)", head=True)
    s.ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=14)
    S(r, 15, "night roster (export)", head=True)
    s.ws.merge_cells(start_row=r, start_column=15, end_row=r, end_column=24)
    S(r, 25, "day", head=True)
    S(r, 26, "night", head=True)
    ros = []
    for i, sd in enumerate(soldiers):
        r += 1
        ros.append(r)
        S(r, 1, i + 1, fmt="0", center=True)
        S(r, 2, sd["name"].split()[-1], inp=True)
        S(r, 3, sd["roster_day"], inp=True, wrap=False)
        s.ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=14)
        S(r, 15, sd["roster_night"], inp=True)
        s.ws.merge_cells(start_row=r, start_column=15, end_row=r, end_column=24)
        for c, src in ((25, 3), (26, 15)):
            kind = f'LEFT({A(r, src)},FIND(":",{A(r, src)})-1)'
            S(r, c, f'=IF({kind}="off duty","off",{kind})', center=True)
        for c in range(3, 25):
            s.ws.cell(row=r, column=c).font = Font(name="Calibri", size=9)
    # 24-hour grid
    r += 2
    S(r, 1, "24-hour grid: state at the middle of each hour = IF(day start <= h + 0.5 < day end, day kind, night kind)",
      border=False, bold=True)
    r += 1
    S(r, 1, "#", head=True)
    S(r, 2, "role", head=True)
    hr = r
    for h in range(24):
        S(r, 3 + h, h, head=True, fmt="0")
    g0 = r + 1
    for i, rr in enumerate(ros):
        r += 1
        S(r, 1, i + 1, fmt="0", center=True)
        S(r, 2, f"={A(rr, 2)}")
        for h in range(24):
            hc = f"{get_column_letter(3 + h)}${hr}"
            S(r, 3 + h, f"=IF(AND({hc}+0.5>={st},{hc}+0.5<{en}),${get_column_letter(25)}${rr},${get_column_letter(26)}${rr})",
              center=True, size=9)
    g1 = r
    r += 1
    S(r, 1, "on duty (not off)", bold=True)
    s.ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    for h in range(24):
        col = get_column_letter(3 + h)
        S(r, 3 + h, f'=COUNTIF({col}{g0}:{col}{g1},"<>off")', fmt="0", center=True)
    r += 1
    S(r, 1, "shift", bold=True)
    s.ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    for h in range(24):
        col = get_column_letter(3 + h)
        S(r, 3 + h, f'=IF(AND({col}${hr}+0.5>={st},{col}${hr}+0.5<{en}),"DAY","NIGHT")', center=True, size=8)
    shift_r = r
    area = f"C{g0}:Z{g1}"
    for txt, color in (("off", "D9D9D9"), ("guard", "F8CBAD"), ("work", "BDD7EE"), ("walk", "C6EFCE"),
                       ("run", "FFE699")):
        s.ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=[f'"{txt}"'],
                                                         fill=PatternFill("solid", fgColor=color)))
    s.ws.conditional_formatting.add(f"C{shift_r}:Z{shift_r}", CellIsRule(operator="equal", formula=['"NIGHT"'],
                                    fill=PatternFill("solid", fgColor="8EA9DB")))
    s.ws.conditional_formatting.add(f"C{shift_r}:Z{shift_r}", CellIsRule(operator="equal", formula=['"DAY"'],
                                    fill=PatternFill("solid", fgColor="FFF2CC")))
    r += 1
    S(r, 1, "Colours: grey off duty, red guard, blue work, green walk (patrol), yellow run.", border=False, italic=True)
    # stagger table
    r += 2
    S(r, 1, "Staggered shift change: departure delay = 1 + 2.5 rank (s, simulation time; rank = seeded order)",
      border=False, bold=True)
    r += 1
    S(r, 1, "rank", head=True)
    s.ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    for k in range(8):
        S(r, 3 + 2 * k, k, inp=True, fmt="0", center=True)
        s.ws.merge_cells(start_row=r, start_column=3 + 2 * k, end_row=r, end_column=4 + 2 * k)
    r += 1
    S(r, 1, "delay (s)", bold=True)
    s.ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    for k in range(8):
        S(r, 3 + 2 * k, f"=1+2.5*{A(r - 1, 3 + 2 * k)}", fmt="0.0", center=True)
        s.ws.merge_cells(start_row=r, start_column=3 + 2 * k, end_row=r, end_column=4 + 2 * k)
    r += 1
    s.note(r, 1, "The day shift follows the sun (DayBreak about 06:27, Dusk 17:36 at latitude 22.8 deg). "
                 "Two shifts; the spec's 6-hour shifts are not in the code.")
    s.finish(last_col=26)


def sheet_performance(wb):
    s = Sheet(wb, "Performance", "Performance: benchmark import (report/data/bench.csv)",
              "report/data/bench.csv written by the benchmark run (ts --benchmark / --bench-shots); this sheet is "
              "rebuilt by tools/v10/build_cg_workbook.py.")
    s.widths({"A": 30, "B": 11, "C": 11, "D": 11, "E": 11, "F": 12, "G": 13, "H": 14, "I": 14})
    S = s.put
    r = 5
    S(r, 1, "IMPORT AREA: rows below are copied from bench.csv (yellow); derived columns are formulas.",
      border=False, bold=True)
    r += 1
    s.heads(r, 1, ["config", "avg_ms", "p95_ms", "avg_fps", "draw_calls", "triangles", "1000 / avg_ms (fps)",
                   "fps check (diff)", "triangles per draw"], height=30)
    h = r
    rows = []
    if BENCH_CSV.exists():
        with BENCH_CSV.open() as f:
            for rec in csv.DictReader(f):
                rows.append(rec)
    if not rows:
        r += 1
        S(r, 1, "TODO: no bench.csv yet", inp=True)
        for c in range(2, 10):
            S(r, c, "-", center=True)
        s.finish(last_col=9)
        return
    for rec in rows:
        r += 1

        def num(k):
            try:
                v = float(rec.get(k, ""))
                return int(v) if v.is_integer() and k in ("draw_calls", "triangles") else v
            except ValueError:
                return rec.get(k, "")

        S(r, 1, rec.get("config", ""), inp=True)
        for j, k in enumerate(("avg_ms", "p95_ms", "avg_fps", "draw_calls", "triangles")):
            v = num(k)
            S(r, 2 + j, v, inp=True, fmt="0" if k in ("draw_calls", "triangles") else "0.00")
        S(r, 7, f"=IF(B{r}>0,1000/B{r},0)", fmt="0.0")
        S(r, 8, f"=G{r}-D{r}", fmt="0.0")
        S(r, 9, f"=IF(E{r}>0,F{r}/E{r},0)", fmt="0")
    last = r
    ch = BarChart()
    ch.type = "bar"
    ch.style = 10
    ch.title = "Frame time per configuration (ms)"
    ch.y_axis.title = "ms"
    data = Reference(s.ws, min_col=2, max_col=3, min_row=h, max_row=last)
    cats = Reference(s.ws, min_col=1, min_row=h + 1, max_row=last)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.height = 9
    ch.width = 22
    s.ws.add_chart(ch, f"A{last + 3}")
    s.finish(last_col=9, last_row=last + 3 + 19)


def build(export_path):
    exp = load_export(export_path)
    wb = Workbook()
    wb.remove(wb.active)
    sheet_primitives(wb, exp)
    sheet_decomposition(wb, exp)
    # View_Projection must exist before Transform_Trace references it; reorder afterwards.
    tt_placeholder = wb.create_sheet("Transform_Trace_tmp")
    vp = sheet_view_projection(wb, exp)
    wb.remove(tt_placeholder)
    sheet_transform_trace(wb, exp, vp)
    sheet_clipping(wb)
    sheet_illumination(wb, exp, vp)
    sheet_color_interp(wb, exp)
    sheet_ray_tracing(wb)
    sheet_thermal_kernels(wb)
    sheet_duty_cycle(wb, exp)
    sheet_performance(wb)
    wb._sheets.sort(key=lambda ws: SHEETS.index(ws.title))
    wb.active = 0
    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX.relative_to(ROOT)} ({len(wb.sheetnames)} sheets)")


# --------------------------------------------------------------------------------------------- snaps + checks
RECALC_XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""

ERRORS = ("#VALUE!", "#REF!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:", "#ERR")


def soffice(args, profile):
    cmd = [str(SOFFICE), f"-env:UserInstallation=file://{profile}", "--headless", "--norestore"] + args
    env = dict(os.environ)
    if NSS_LIB.is_dir():  # the extracted AppImage needs libnss3 / libssl3 (apt-get download + dpkg-deb -x, no sudo)
        env["LD_LIBRARY_PATH"] = f"{NSS_LIB}:{env.get('LD_LIBRARY_PATH', '')}"
    subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300, env=env)


def snap():
    import fitz  # pymupdf
    from openpyxl import load_workbook
    from PIL import Image, ImageChops

    work = Path(tempfile.mkdtemp(prefix="cgwb_"))
    profile = work / "profile"
    (profile / "user").mkdir(parents=True)
    (profile / "user" / "registrymodifications.xcu").write_text(RECALC_XCU)
    pdf_dir, xl_dir = work / "pdf", work / "xlsx"
    soffice(["--convert-to", "pdf", "--outdir", str(pdf_dir), str(OUT_XLSX)], profile)
    soffice(["--convert-to", "xlsx", "--outdir", str(xl_dir), str(OUT_XLSX)], profile)

    # 1. errors in the LibreOffice-recalculated copy
    calc = load_workbook(xl_dir / OUT_XLSX.name, data_only=True)
    form = load_workbook(OUT_XLSX)
    problems, n_formulas = [], 0
    for ws in form.worksheets:
        cws = calc[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    n_formulas += 1
                    v = cws[cell.coordinate].value
                    if v is None or (isinstance(v, str) and v.startswith(ERRORS)):
                        problems.append(f"{ws.title}!{cell.coordinate}: {v!r}  {cell.value[:90]}")
    print(f"formula cells: {n_formulas}, errors / empty results: {len(problems)}")
    for p in problems[:40]:
        print("  ", p)
    tt = calc["Transform_Trace"]
    for row in tt.iter_rows():
        for cell in row:
            if cell.value in ("max |M - export|", "max |World - export|", "inside clip volume |x|,|y|,|z| <= w",
                              "eye-space distance -z_eye (m)"):
                print(f"  Transform_Trace {cell.value} = {tt.cell(row=cell.row, column=cell.column + 3).value}")

    # 2. PDF pages -> sheets (title text) -> PNG
    doc = fitz.open(pdf_dir / (OUT_XLSX.stem + ".pdf"))
    titles = {}
    for ws in form.worksheets:
        titles[ws.title] = str(ws["A1"].value).split(":")[0] + ":"
    owner, cur = [], None
    for page in doc:
        text = page.get_text()
        for name in SHEETS:
            if titles.get(name) and text.lstrip().startswith(titles[name]) or (
                    titles.get(name) and titles[name] in text[:200]):
                cur = name
                break
        owner.append(cur)
    SNAP_DIR.mkdir(parents=True, exist_ok=True)
    for old in SNAP_DIR.glob("*.png"):
        old.unlink()
    count = {}
    sizes = []
    for i, page in enumerate(doc):
        name = owner[i] or f"page{i + 1}"
        count[name] = count.get(name, 0) + 1
        fn = SNAP_DIR / (f"{name}.png" if count[name] == 1 else f"{name}_p{count[name]}.png")
        pix = page.get_pixmap(dpi=200)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        bg = Image.new("RGB", img.size, (255, 255, 255))
        bbox = ImageChops.difference(img, bg).convert("L").point(lambda x: 255 if x > 8 else 0).getbbox()
        if bbox:
            pad = 24
            bbox = (max(0, bbox[0] - pad), max(0, bbox[1] - pad), min(img.width, bbox[2] + pad),
                    min(img.height, bbox[3] + pad))
            img = img.crop(bbox)
        img.save(fn, optimize=True)
        sizes.append((fn.name, img.width, img.height, fn.stat().st_size))
    for n, w, h, b in sizes:
        print(f"  {n}: {w}x{h} px, {b // 1024} KiB")
    shutil.rmtree(work, ignore_errors=True)
    return problems


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--export", help="ts --export-docs JSON (default: run build/ts-v9)")
    ap.add_argument("--no-snap", action="store_true")
    ap.add_argument("--snap-only", action="store_true")
    a = ap.parse_args()
    if not a.snap_only:
        build(a.export)
    if a.no_snap:
        return 0
    try:
        import fitz  # noqa: F401
    except ImportError:
        return subprocess.run([str(UV), "run", "-q", "--with", "openpyxl", "--with", "pymupdf", "--with", "pillow",
                               "python", str(Path(__file__).resolve()), "--snap-only"]).returncode
    problems = snap()
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
