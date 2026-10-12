"""Register sheets of the workbook: Index, World_Coordinates, Layout, Algorithms, Film.

These make the workbook the one place where every coordinate and constant of the project can be looked up:
  Index              what each sheet holds, the report section that prints it, its source in the code, the checker that guards it
  World_Coordinates  every component of the scene (parts and groups) with its S, R, T and its place in the world
  Layout             sites of the base, the things that can be operated, the lights, the walk space
  Algorithms         every constant of the vision pipeline, the tracker, the squad, the flight model, the intruders' planner,
                     the mission and the garrison, with the source line it is read from
  Film               the shot list of the full-length film and the numbers it states
Called by build_cg_workbook.py (add(wb, exp, Sheet, ROOT))."""
import json
import re
import subprocess

from openpyxl.utils import get_column_letter

INDEX = [
    ("Index", "this list", "-", "tools/v10/wb_registers.py", "-"),
    ("Primitives", "the 16 generators: vertex and triangle counts as formulas of their parameters", "6", "src/primitives/*.cpp", "check_workbook.py, check_counts.py"),
    ("Decomposition", "parts of the UAV and of a soldier, counted by primitive", "7", "scene export: components", "check_workbook.py"),
    ("Designs", "the 62 designs: instances, parts per instance, primitives", "7", "scene export: designs (src/report/Designs.cpp)", "check_counts.py"),
    ("Transform_Trace", "one real vertex from the generator to its pixel: M = T R S, world, view, clip, divide, viewport", "8, 10-12", "scene export + src/math/Math3D.cpp", "check_workbook.py"),
    ("Vertex_Groups", "worked groups of vertices under a chain of transformations", "8, 9", "scene export", "check_workbook.py"),
    ("View_Projection", "the look-at and perspective matrices of the default camera, depth precision", "10-12", "src/scene/Camera.cpp, src/math/Math3D.cpp", "check_workbook.py"),
    ("Clipping", "Cohen-Sutherland and Liang-Barsky on a worked segment", "13", "worked example", "check_workbook.py"),
    ("Illumination", "sun position, air mass, light colour, the lighting equation on one surface", "16, 22", "src/scene/DayCycle.cpp, shaders", "check_workbook.py"),
    ("Color_Interp", "edge functions, barycentric and perspective-correct interpolation", "15, 21", "src/math/ExplorerMath.cpp", "check_workbook.py"),
    ("Ray_Tracing", "reflection, refraction and Fresnel on a worked ray", "20, 25", "src/raytrace/Optics.h", "check_workbook.py"),
    ("Thermal_Kernels", "the Gaussian and Sobel kernels applied to a patch", "26", "src/math/ExplorerMath.cpp", "check_workbook.py"),
    ("Tracking", "the Kalman filter of the multi-target tracker, step by step", "27", "src/vision/MultiTracker.cpp", "check_workbook.py"),
    ("Line_Follower", "sensor error, PD controller and wheel speeds of an intruder rover", "29", "src/ai/HideController.cpp", "check_workbook.py"),
    ("Explosion", "particle motion and the blackbody colour ramp", "28", "src/scene/Combat.cpp", "check_workbook.py"),
    ("Duty_Cycle", "the soldiers' shifts and rosters", "24, 27", "src/scene/Garrison.cpp", "check_workbook.py"),
    ("Performance", "the benchmark per preset", "31", "report/data/bench.csv", "check_workbook.py"),
    ("Report_Numbers", "every number printed in the report's text, with its source", "all", "code + scene export", "numbers_from_workbook.py --check, check_numbers.py"),
    ("Coordinates", "every component of every design in its design frame (the report's coordinate sheets)", "7", "scene export: world matrices", "check_workbook.py"),
    ("World_Coordinates", "every component of the scene in world coordinates, with S, R, T, colour and heat", "4, 7, 8", "scene export: components", "check_workbook.py"),
    ("Layout", "sites of the base, the things that can be operated, lights, walk space", "4, 5, 16, 24", "scene export: layout, controllables, lights, practical_lights, walkspace", "check_workbook.py"),
    ("Algorithms", "every constant of the vision, tracking, squad, flight, intruder, mission and garrison code", "26-29", "the header and source files named in each row", "check_workbook.py"),
    ("Film", "shot list of the full-length film and the numbers it states", "-", "tools/film/acts/*.py -> report/data/film.json", "check_workbook.py"),
]

CONST_FILES = ["src/vision/CvPipeline.h", "src/vision/MultiTracker.h", "src/vision/SensorModel.h", "src/ai/DroneController.h", "src/ai/Squad.h",
               "src/ai/HideController.h", "src/ai/HideController.cpp", "src/scene/Mission.h", "src/scene/Scene.h", "src/scene/Garrison.h",
               "src/scene/Combat.h", "src/scene/DayCycle.h", "src/scene/Gait.h"]


def code_constants(root):
    """[(file, line, name, value, comment)] of the constexpr numbers in the algorithm sources."""
    out = []
    for rel in CONST_FILES:
        p = root / rel
        if not p.exists():
            continue
        for n, ln in enumerate(p.read_text(errors="replace").splitlines(), 1):
            m = re.match(r"\s*(?:static\s+)?constexpr\s+(?:float|int|double|unsigned)\s+(.*?);\s*(?://\s*(.*))?$", ln)
            if not m:
                continue
            for item in m.group(1).split(","):
                q = re.match(r"\s*(\w+)\s*=\s*(-?[\d.]+)f?\s*$", item)
                if q:
                    v = float(q.group(2))
                    out.append((rel, n, q.group(1), int(v) if v == int(v) and "." not in q.group(2) else v, (m.group(2) or "").strip()))
    return out


def flatten(prefix, v, out):
    if isinstance(v, dict):
        for k, w in v.items():
            flatten(prefix + [str(k)], w, out)
    elif isinstance(v, list):
        if v and all(isinstance(x, (int, float)) and not isinstance(x, bool) for x in v) and len(v) <= 4:
            for k, w in zip("xyzw", v):
                out.append((".".join(prefix) + "." + k, w))
        elif len(v) <= 12 and all(not isinstance(x, (dict, list)) for x in v):
            out.append((".".join(prefix), ", ".join(str(x) for x in v)))
    elif isinstance(v, bool):
        out.append((".".join(prefix), "yes" if v else "no"))
    elif v is not None:
        out.append((".".join(prefix), v))


def add(wb, exp, Sheet, ROOT):
    names = [r[0] for r in INDEX]
    # ---------------------------------------------------------------------------------------------- Index
    s = Sheet(wb, "Index", "ThermalSense workbook: the register of every number and coordinate of the project",
              "generated by tools/v10/build_cg_workbook.py from the scene export of the program and from its source files")
    s.ws.cell(row=2, column=1).value = "How to read it: yellow cells are inputs taken from the code or from the scene export; white cells are formulas."
    s.widths({"A": 22, "B": 78, "C": 14, "D": 62, "E": 52})
    s.heads(5, 1, ["sheet", "holds", "report section", "source", "guarded by"], height=18)
    for i, row in enumerate(INDEX):
        for j, v in enumerate(row):
            s.put(6 + i, 1 + j, v, wrap=j in (1, 3, 4))
        s.ws.cell(row=6 + i, column=1).hyperlink = "#'%s'!A1" % row[0]
    r = 8 + len(INDEX)
    try:
        commit = subprocess.run(["git", "rev-parse", "--short", "HEAD"], cwd=ROOT, capture_output=True, text=True).stdout.strip()
    except Exception:
        commit = ""
    for k, (a, b) in enumerate((("scene export schema", exp["schema_version"]), ("drawn parts in the scene", exp["leaf_count"]),
                                ("components (parts and groups)", len(exp["components"])), ("designs", len(exp["designs"])),
                                ("source revision the workbook was built on", commit))):
        s.put(r + k, 1, a, bold=True)
        s.put(r + k, 2, b, inp=True)
    s.finish(last_col=5, fit=False, scale=70)

    # ---------------------------------------------------------------------------------------------- World_Coordinates
    s = Sheet(wb, "World_Coordinates", "World_Coordinates: every component of the scene (m, degrees), at t = 0",
              "scene export: components[]. S, R, T are the part's own scale, rotation (pitch x, yaw y, roll z) and translation in its "
              "parent's frame; M = T Ry Rx Rz S; world = parent world x M. Centre = last column of the world matrix; extent = box of the "
              "unit primitive under the world matrix.")
    cols = ["id", "object", "design", "component", "primitive", "parent", "depth", "S x", "S y", "S z", "R x", "R y", "R z", "T x", "T y", "T z",
            "world x", "world y", "world z", "extent x", "extent y", "extent z", "heat", "red", "green", "blue", "hidden at t = 0"]
    s.widths({"A": 11, "B": 22, "C": 24, "D": 30, "E": 16, "F": 11, **{get_column_letter(c): 9 for c in range(7, 28)}})
    first = 9
    last = first + len(exp["components"]) - 1
    s.put(5, 1, "drawn parts (not groups)", bold=True)
    s.put(5, 4, "=COUNTIF(E%d:E%d,\"<>Group\")" % (first, last), fmt="0")
    s.put(5, 5, "scene export says", border=False, italic=True)
    s.put(5, 6, exp["leaf_count"], inp=True)
    s.put(5, 7, "=IF(D5=F5,\"agree\",\"DIFFER\")")
    s.put(6, 1, "groups", bold=True)
    s.put(6, 4, "=COUNTIF(E%d:E%d,\"Group\")" % (first, last), fmt="0")
    s.put(7, 1, "components", bold=True)
    s.put(7, 4, "=D5+D6", fmt="0")
    s.heads(first - 1, 1, cols, height=18)
    for k, c in enumerate(exp["components"]):
        r = first + k
        m = c["world_matrix"]
        ext = [abs(m[4 * i]) + abs(m[4 * i + 1]) + abs(m[4 * i + 2]) for i in range(3)]
        vals = [c["id"], c["object"], c.get("design", ""), c["name"], c["primitive"], c.get("parent_id") or "", c["depth"]]
        for j, v in enumerate(vals):
            s.put(r, 1 + j, v)
        nums = list(c["S"]) + list(c["R_deg"]) + list(c["T"]) + [m[3], m[7], m[11]] + ext + [c.get("heat", 0)] + list(c.get("color", [0, 0, 0]))
        for j, v in enumerate(nums):
            s.put(r, 8 + j, round(float(v), 4), inp=True, fmt="0.000")
        s.put(r, 27, "yes" if c.get("hidden_at_export") else "")
    s.ws.freeze_panes = "E%d" % first
    s.finish(last_col=len(cols), fit=False, scale=60)
    s.ws.print_area = "A1:P%d" % (first + 40)    # the printed copy shows the head of the list; the sheet holds all of it

    # ---------------------------------------------------------------------------------------------- Layout
    s = Sheet(wb, "Layout", "Layout: sites, things that can be operated, lights and the walk space (m)",
              "scene export: layout (src/objects/BaseLayout.h), controllables, lights, practical_lights, walkspace")
    s.widths({"A": 34, "B": 26, **{get_column_letter(c): 11 for c in range(3, 14)}})
    r = 5
    s.put(r, 1, "A. Sites of the base", border=False, bold=True)
    s.heads(r + 1, 1, ["name", "value"])
    flat = []
    flatten([], exp["layout"], flat)
    for k, (a, b) in enumerate(flat):
        s.put(r + 2 + k, 1, a)
        s.put(r + 2 + k, 2, b, inp=True, fmt="0.000" if isinstance(b, float) else None)
    r += 4 + len(flat)
    s.put(r, 1, "B. Things that can be operated", border=False, bold=True)
    s.heads(r + 1, 1, ["name", "object", "x", "y", "z", "parts", "can be controlled", "key", "controls"])
    for k, c in enumerate(exp["controllables"]):
        s.put(r + 2 + k, 1, c["name"])
        s.put(r + 2 + k, 2, c["root_object"])
        for j, v in enumerate(c["position"]):
            s.put(r + 2 + k, 3 + j, round(float(v), 3), inp=True, fmt="0.00")
        s.put(r + 2 + k, 6, c["parts"], inp=True)
        s.put(r + 2 + k, 7, "yes" if c["possessable"] else "no")
        s.put(r + 2 + k, 8, str(c.get("key", "")))
        s.put(r + 2 + k, 9, c.get("controls", ""))
    n_ctl = len(exp["controllables"])
    s.put(r + 2 + n_ctl, 1, "that can be controlled", bold=True)
    s.put(r + 2 + n_ctl, 7, "=COUNTIF(G%d:G%d,\"yes\")" % (r + 2, r + 1 + n_ctl), fmt="0")
    r += 5 + n_ctl
    s.put(r, 1, "C. Scene lights", border=False, bold=True)
    s.heads(r + 1, 1, ["name", "value"])
    flat = []
    flatten([], exp["lights"], flat)
    for k, (a, b) in enumerate(flat):
        s.put(r + 2 + k, 1, a)
        s.put(r + 2 + k, 2, b, inp=True, fmt="0.000" if isinstance(b, float) else None)
    r += 4 + len(flat)
    s.put(r, 1, "D. Practical lights (lamps)", border=False, bold=True)
    s.heads(r + 1, 1, ["name", "object", "x", "y", "z", "red", "green", "blue", "intensity", "range (m)", "cone (deg)"])
    for k, p in enumerate(exp["practical_lights"]):
        s.put(r + 2 + k, 1, p["name"])
        s.put(r + 2 + k, 2, p["object"])
        for j, v in enumerate(list(p["position"]) + list(p["color"]) + [p["intensity"], p["range"], p.get("cone_deg") or 0]):
            s.put(r + 2 + k, 3 + j, round(float(v), 3), inp=True, fmt="0.00")
    n_pl = len(exp["practical_lights"])
    s.put(r + 2 + n_pl, 1, "lamps", bold=True)
    s.put(r + 2 + n_pl, 2, "=COUNTA(A%d:A%d)" % (r + 2, r + 1 + n_pl), fmt="0")
    r += 5 + n_pl
    s.put(r, 1, "E. Walk space: floors", border=False, bold=True)
    s.heads(r + 1, 1, ["#", "", "x0", "z0", "x1", "z1", "y0", "y1"])
    for k, f in enumerate(exp["walkspace"]["floors"]):
        s.put(r + 2 + k, 1, k + 1)
        for j, key in enumerate(("x0", "z0", "x1", "z1", "y0", "y1")):
            s.put(r + 2 + k, 3 + j, round(float(f[key]), 3), inp=True, fmt="0.00")
    r += 4 + len(exp["walkspace"]["floors"])
    s.put(r, 1, "F. Walk space: walls (boxes)", border=False, bold=True)
    s.heads(r + 1, 1, ["#", "", "min x", "min y", "min z", "max x", "max y", "max z"])
    for k, w in enumerate(exp["walkspace"]["walls"]):
        s.put(r + 2 + k, 1, k + 1)
        for j, v in enumerate(list(w["min"]) + list(w["max"])):
            s.put(r + 2 + k, 3 + j, round(float(v), 3), inp=True, fmt="0.00")
    s.finish(last_col=13, fit=False, scale=70)

    # ---------------------------------------------------------------------------------------------- Algorithms
    s = Sheet(wb, "Algorithms", "Algorithms: every constant of the vision, tracking, squad, flight, intruder, mission and garrison code",
              "A: constants read from the source files (file and line given). B: the values the running program reports in its scene export.")
    s.widths({"A": 34, "B": 8, "C": 30, "D": 14, "E": 96})
    s.put(5, 1, "A. Constants in the source", border=False, bold=True)
    s.heads(6, 1, ["file", "line", "name", "value", "comment in the source"])
    consts = code_constants(ROOT)
    for k, (f, n, name, v, cm) in enumerate(consts):
        r = 7 + k
        s.put(r, 1, f)
        s.put(r, 2, n)
        s.put(r, 3, name)
        s.put(r, 4, v, inp=True, fmt="0" if isinstance(v, int) else "0.0###")
        s.put(r, 5, cm)
    r = 9 + len(consts)
    s.put(r, 1, "B. Values reported by the program (scene export)", border=False, bold=True)
    s.heads(r + 1, 1, ["group", "", "name", "value"])
    flat = []
    for grp in (("v5", "squad"), ("v5", "flight_model"), ("v5", "tracker"), ("v5", "cv"), ("v5", "armour"), ("v5", "mission"), ("v6", "sensor"),
                ("v6", "day_cycle", "config"), ("combat", "constants"), ("terrain",)):
        node = exp
        for g in grp:
            node = node.get(g, {}) if isinstance(node, dict) else {}
        one = []
        flatten([], node, one)
        flat += [(".".join(grp), a, b) for a, b in one]
    for k, (g, a, b) in enumerate(flat):
        s.put(r + 2 + k, 1, g)
        s.put(r + 2 + k, 3, a)
        s.put(r + 2 + k, 4, b, inp=True, fmt="0.0###" if isinstance(b, float) else None)
    s.finish(last_col=5, fit=False, scale=70)

    # ---------------------------------------------------------------------------------------------- Film
    film_json = ROOT / "report" / "data" / "film.json"
    s = Sheet(wb, "Film", "Film: shot list of the full-length film and the numbers it states",
              "report/data/film.json, written by tools/film/build.py plan from tools/film/acts/*.py (narration lengths set the times)")
    s.widths({"A": 9, "B": 14, "C": 38, "D": 10, "E": 26, "F": 44, "G": 9})
    if film_json.exists():
        film = json.loads(film_json.read_text())
        s.put(5, 1, "length (s)", bold=True)
        s.put(5, 3, film["length"], inp=True, fmt="0.0")
        s.put(5, 4, "=C5/60", fmt="0.0")
        s.put(5, 5, "minutes", border=False, italic=True)
        s.put(6, 1, "beats", bold=True)
        s.put(6, 3, "=COUNTA(A%d:A%d)" % (10, 9 + len(film["beats"])), fmt="0")
        s.heads(9, 1, ["start (s)", "act", "chapter", "picture", "source", "title on screen", "length (s)"])
        for k, b in enumerate(film["beats"]):
            r = 10 + k
            s.put(r, 1, b["start"], inp=True, fmt="0.0")
            s.put(r, 2, b["act"])
            s.put(r, 3, b["chapter"])
            s.put(r, 4, b["kind"])
            s.put(r, 5, b["source"])
            s.put(r, 6, b["title"])
            s.put(r, 7, b["length"], inp=True, fmt="0.0")
        r = 12 + len(film["beats"])
        s.put(r, 1, "sum of the lengths", bold=True)
        s.put(r, 7, "=SUM(G10:G%d)" % (9 + len(film["beats"])), fmt="0.0")
        s.put(r + 2, 1, "Numbers the narration states, with their source", border=False, bold=True)
        s.heads(r + 3, 1, ["", "", "what", "value", "", "source"])
        said = [("generators (primitives)", len(exp["primitives"]), "scene export: primitives"), ("drawn parts", exp["leaf_count"], "scene export: leaf_count"),
                ("designs", len(exp["designs"]), "scene export: designs"),
                ("things that can be operated", sum(1 for c in exp["controllables"] if c["possessable"]), "scene export: controllables (possessable)"),
                ("lamps (practical lights)", len(exp["practical_lights"]), "scene export: practical_lights"),
                ("slices of the sphere", next(p for p in exp["primitives"] if p["name"] == "sphere")["params"]["slices"], "scene export: primitives"),
                ("corners of the sphere", next(p for p in exp["primitives"] if p["name"] == "sphere")["vertices"], "scene export: primitives"),
                ("triangles of the sphere", next(p for p in exp["primitives"] if p["name"] == "sphere")["triangles"], "scene export: primitives"),
                ("parts of a drone", next(o for o in exp["objects"] if o["object"] == "Autonomous UAV")["parts"], "scene export: objects"),
                ("parts of a tank", next(o for o in exp["objects"] if o["object"] == "Main Battle Tank")["parts"], "scene export: objects"),
                ("parts of a soldier", next(o for o in exp["objects"] if o["object"] == "Soldier 1")["parts"], "scene export: objects"),
                ("latitude (degrees north)", exp["v6"]["day_cycle"]["config"]["latitude_deg"], "scene export: v6.day_cycle.config")]
        want = {"kSquadLeadLock": "lead's lock time (s)", "kSquadWingLock": "wing's lock time (s)", "kSquadWingWaitFree": "lead waits for a free wing (s)",
                "kSquadSplitDist": "split distance (m)", "kSquadRejoinDist": "rejoin distance (m)", "kSquadSplitThreat": "urgency of a second intruder (m)",
                "kEvadeRadius": "a drone this close makes a rover hide (m)", "kEvadeTime": "time in cover (s)", "kGridCells": "cells per side of the rover's grid",
                "kCellSize": "size of a cell (m)", "kIrSensorCount": "line sensors of a rover", "kKp": "line follower: proportional gain",
                "kKd": "line follower: derivative gain", "kCvWidth": "thermal image width (pixels)", "kCvHeight": "thermal image height (pixels)",
                "kRoverHealth": "armour of a rover", "kRoundDamage": "damage of a round at its centre", "kKillScore": "score of a kill",
                "kCombinedBonus": "bonus of a combined kill", "kBreachDamage": "base integrity lost per breach (%)"}
        for f, n, name, v, _ in consts:
            if name in want:
                said.append((want.pop(name), v, "%s:%d %s" % (f, n, name)))
        for k, (a, b, c) in enumerate(said):
            s.put(r + 4 + k, 3, a)
            s.put(r + 4 + k, 4, b, inp=True, fmt="0.0###" if isinstance(b, float) else None)
            s.put(r + 4 + k, 6, c)
    else:
        s.put(5, 1, "report/data/film.json not found: run tools/film/build.py plan")
    s.finish(last_col=7, fit=False, scale=70)
    return names
