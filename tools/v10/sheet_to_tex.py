"""v10 M1: the key blocks of each workbook sheet as native LaTeX (booktabs) fragments.

Reads report/workbooks/cg_workbook.xlsx (formulas, fills: yellow = input from the code) and
report/workbooks/cg_workbook_values.xlsx (the LibreOffice-recalculated copy written by build_cg_workbook.py, so formula
cells hold cached values; openpyxl does not evaluate formulas). Writes report/sections/wb/<Sheet>_<block>.tex.
Input cells keep the workbook's convention as a light yellow cell shade (\\cellcolor{wbin}); formula cells are white.

Block kinds:
  ("table", name, range, opts)      opts: cols=[letters] (subset), rows=[row numbers] (subset), head=k header rows,
                                    fmt="4" (decimals for non-integers), wrap={letter: width} (p-columns)
  ("mats", name, [(label, range), ...])   one display line of 4x4 (or n x m) bmatrix matrices
Run: ~/.local/bin/uv run -q --with openpyxl python tools/v10/sheet_to_tex.py"""
import os
import re

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
WB = os.path.join(ROOT, "report", "workbooks", "cg_workbook.xlsx")
WBV = os.path.join(ROOT, "report", "workbooks", "cg_workbook_values.xlsx")
OUT = os.path.join(ROOT, "report", "sections", "wb")

L = list  # readability of the config below
BLOCKS = {
    "Primitives": [
        ("table", "counts", "A5:Q13", dict(cols=L("ACDEFGHLMP"), hw="1.3cm")),
        ("table", "param", "A26:J31", dict(cols=L("ABJ"), wrap={"B": "8.2cm", "J": "5.2cm"})),
    ],
    "Decomposition": [
        ("table", "totals", "A6:L40", dict(first="3.2cm", hw="1.2cm")),
        ("table", "uav", "A44:R70", dict(cols=L("BCDEFGHIJKLM"), wrap={"D": "2.6cm"}, first="2.6cm", hw="1.0cm", fmt=3)),
    ],
    "Transform_Trace": [
        ("table", "inputs", "A5:E10", dict(first="4.6cm")),
        ("mats", "srt", [("S", "B13:E16"), ("R_x", "G13:J16"), ("R_y", "L13:O16"), ("R_z", "Q13:T16")]),
        ("mats", "tm", [("T", "B19:E22"), ("M", "B25:E28")]),
        ("mats", "tw", [("W_{\\mathrm{par}}", "Q25:T28"), ("W", "B31:E34")]),
        ("table", "stages", "G37:O43", dict(first="5.6cm")),
        ("table", "checks", "Q37:T41", dict(first="5cm")),
    ],
    "View_Projection": [
        ("table", "basis", "A5:D25", dict(first="6.4cm")),
        ("table", "depth", "G5:M18", dict(hw="2.1cm")),
        ("mats", "vp", [("V", "B28:E31"), ("P", "B35:E38")]),
    ],
    "Clipping": [
        ("table", "cs", "A8:S12", dict(hw="0.9cm", wrap={"S": "3.2cm"})),
        ("table", "lb", "A15:S19", dict(hw="1.3cm")),
        ("table", "sh1", "A24:J29", dict(hw="1.0cm")),
        ("table", "sh2", "L24:U30", dict(hw="1.0cm")),
    ],
    "Illumination": [
        ("table", "geometry", "A5:F31", dict(first="6.2cm", hw="1.9cm")),
        ("table", "terms", "A34:F53", dict(first="5.2cm", hw="1.9cm")),
    ],
    "Color_Interp": [
        ("table", "verts", "A6:I12", dict(first="2.4cm", wrap={"B": "2.4cm"})),
        ("table", "srgb", "K9:S14", dict(hw="1.3cm")),
        ("table", "bary", "A15:H20", dict()),
        ("table", "gp", "A23:H27", dict(hw="1.5cm")),
        ("table", "uv", "A31:D35", dict()),
        ("table", "sky", "K19:U33", dict(hw="1.3cm")),
    ],
    "Ray_Tracing": [
        ("table", "sphere", "A5:D22", dict(first="5.2cm")),
        ("table", "snell", "F5:I18", dict(first="5.4cm")),
        ("table", "tir", "F21:I27", dict(hw="2.2cm")),
        ("table", "plane", "A24:D33", dict(first="5.2cm")),
        ("table", "whitted", "F31:I38", dict(hw="2.2cm")),
    ],
    "Thermal_Kernels": [
        ("mats", "gauss", [("G_{\\text{raw}}", "B8:F12"), ("G=G_{\\text{raw}}/\\textstyle\\sum", "H8:L12")]),
        ("mats", "patch", [("I", "B16:F20"), ("S_x", "H16:J18"), ("S_y", "L16:N18")]),
        ("mats", "out", [("G_x", "B24:D26"), ("G_y", "F24:H26"), ("\\lVert\\nabla I\\rVert", "J24:L26"), ("\\text{edge}", "N24:P26")]),
    ],
    "Tracking": [
        ("mats", "predict", [("\\mathbf{x}", "A9:A12"), ("F", "C9:F12"), ("P", "H9:K12"), ("Q", "M9:P12")]),
        ("mats", "predict2", [("\\mathbf{x}'=F\\mathbf{x}", "A15:A18"), ("P'=FPF^\\top+Q", "C15:F18")]),
        ("table", "meas", "A21:C22", dict(first="2.6cm")),
        ("mats", "update", [("S", "A24:B25"), ("S^{-1}", "D24:E25"), ("K", "G24:H27"), ("\\mathbf{x}_{\\text{new}}", "J24:J27")]),
        ("table", "d2", "A27:D27", dict(first="2.6cm", head=0)),
        ("table", "tri", "A31:D35", dict(first="2.2cm")),
        ("table", "tri2", "A38:D42", dict(first="2.6cm")),
        ("table", "tot", "M23:R25", dict(hw="1.3cm")),
    ],
    "Line_Follower": [
        ("table", "inputs", "B5:M6", dict(hw="1.1cm")),
        ("table", "sensors", "A9:D14", dict()),
        ("table", "pd", "A17:B24", dict(first="2.6cm", head=0)),
    ],
    "Explosion": [
        ("table", "bb", "A10:H21", dict(hw="1.1cm")),
        ("table", "life", "J10:S21", dict(hw="1.0cm")),
        ("table", "spark", "A27:F56", dict(rows=[27, 28, 32, 36, 40, 44, 48, 52, 54, 56], hw="1.1cm")),
        ("table", "debris", "J27:N33", dict(hw="1.3cm")),
    ],
    "Duty_Cycle": [
        ("table", "roster", "A7:Z15", dict(cols=L("ABYZ"))),
        ("table", "grid", "A18:Z28", dict(hw="0.5cm")),
    ],
    "Performance": [
        ("table", "bench", "A6:I17", dict(cols=L("ABGHI"), first="5.4cm", hw="2.0cm")),
    ],
}

TEX_ESC = {"&": "\\&", "%": "\\%", "#": "\\#", "_": "\\_", "$": "\\$", "{": "\\{", "}": "\\}", "~": "\\textasciitilde{}",
           "^": "\\textasciicircum{}", "\\": "\\textbackslash{}"}


def esc(s):
    s = "".join(TEX_ESC.get(ch, ch) for ch in s)
    return s.replace("<=", "$\\le$").replace(">=", "$\\ge$").replace("->", "$\\to$").replace("|", "$|$")


def fmt(v, decimals=4):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, int):
        return f"{v:,}".replace(",", "{,}") if abs(v) >= 10000 else str(v)
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e9:
            iv = int(v)
            return f"{iv:,}".replace(",", "{,}") if abs(iv) >= 10000 else str(iv)
        if v != 0 and (abs(v) < 1e-3 or abs(v) >= 1e6):
            m, e = f"{v:.3e}".split("e")
            return f"${m}\\times10^{{{int(e)}}}$"
        return f"{v:.{decimals}f}"
    return esc(str(v))


def is_input(cell):
    f = cell.fill
    return bool(f and f.fill_type and str(f.fgColor.rgb).upper().endswith("FFF2CC"))


def merged_lookup(ws):
    m = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                m[(r, c)] = (rng.min_row, rng.min_col, rng.max_col - rng.min_col + 1, (r, c) == (rng.min_row, rng.min_col))
    return m


def table(ws, wv, rng, opts):
    c0, r0, c1, r1 = range_boundaries(rng)
    cols = [column_index_from_string(x) for x in opts["cols"]] if opts.get("cols") else list(range(c0, c1 + 1))
    rows = opts.get("rows") or list(range(r0, r1 + 1))
    dec = int(opts.get("fmt", 4))
    head = opts.get("head", 1)
    merged = merged_lookup(ws)
    grid = []
    for r in rows:
        line = []
        for c in cols:
            mi = merged.get((r, c))
            if mi and not mi[3]:
                line.append(None)  # covered by a merge (its text is in the merge's first cell)
                continue
            line.append((wv.cell(row=r, column=c).value, is_input(ws.cell(row=r, column=c)), ws.cell(row=r, column=c).font.b))
        grid.append(line)
    # drop columns and rows that are empty everywhere
    keep_c = [j for j in range(len(cols)) if any(g[j] and g[j][0] not in (None, "") for g in grid)]
    keep_r = [i for i in range(len(rows)) if any(grid[i][j] and grid[i][j][0] not in (None, "") for j in keep_c)]
    out = []
    ncol = len(keep_c)
    wrap = opts.get("wrap", {})
    spec = []
    for jj, j in enumerate(keep_c):
        letter = get_column_letter(cols[j])
        if jj == 0 and opts.get("first"):
            spec.append(">{\\raggedright\\arraybackslash}p{" + opts["first"] + "}")
        elif letter in wrap:
            spec.append(">{\\raggedright\\arraybackslash}p{" + wrap[letter] + "}")
        else:
            spec.append("l" if jj == 0 else "r")
    hw = opts.get("hw", "1.6cm")
    out.append("\\begin{tabular}{@{}" + "".join(spec) + "@{}}\\toprule")
    for k, i in enumerate(keep_r):
        cells = []
        for j in keep_c:
            g = grid[i][j]
            if g is None:
                cells.append("")
                continue
            v, inp, bold = g
            t = fmt(v, dec)
            if k < head and t:
                t = "\\textbf{" + t + "}"
                if len(str(v)) > 9 and j != keep_c[0]:
                    t = "\\parbox[b]{" + hw + "}{\\raggedleft " + t + "}"
            if inp and t:
                t = "\\cellcolor{wbin}" + t
            cells.append(t)
        out.append(" & ".join(cells) + "\\\\" + ("\\midrule" if k == head - 1 else ""))
    out.append("\\bottomrule\\end{tabular}")
    return "\n".join(out)


def mats(ws, wv, items):
    parts = []
    for label, rng in items:
        c0, r0, c1, r1 = range_boundaries(rng)
        rows = []
        for r in range(r0, r1 + 1):
            vals = []
            for c in range(c0, c1 + 1):
                v = wv.cell(row=r, column=c).value
                s = fmt(v, 4)
                s = s.strip("$")
                vals.append(s)
            rows.append(" & ".join(vals))
        parts.append(f"{label}=\\begin{{bmatrix}}" + "\\\\".join(rows) + "\\end{bmatrix}")
    return "\\[\\scriptstyle " + ",\\quad ".join(parts) + "\\]"


def main():
    wf = load_workbook(WB)
    wv = load_workbook(WBV, data_only=True)
    os.makedirs(OUT, exist_ok=True)
    for old in os.listdir(OUT):
        if old.endswith('.tex'):
            os.remove(os.path.join(OUT, old))
    n = 0
    for sheet, blocks in BLOCKS.items():
        ws, vs = wf[sheet], wv[sheet]
        for b in blocks:
            kind, name = b[0], b[1]
            body = table(ws, vs, b[2], b[3]) if kind == "table" else mats(ws, vs, b[2])
            fn = os.path.join(OUT, f"{sheet}_{name}.tex")
            with open(fn, "w") as f:
                f.write(f"% generated by tools/v10/sheet_to_tex.py from cg_workbook.xlsx sheet {sheet}: do not edit\n")
                f.write(body + "\n")
            n += 1
    print("wrote", n, "fragments to", OUT)


if __name__ == "__main__":
    main()
